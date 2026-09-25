"""ICD-11 MMS catalog: import, stats, browsing and policy helpers.

Mirrors ``app/services/icd10_2019_2_service.py`` conventions for the ICD-11
MMS linearization catalog (``mas_icd11_mms``), seeded from the frozen WHO
Simple Tabulation export. See docs/policy/icd11-reference-catalog.md and
docs/icd-causegrp-mappings/migration-artifacts/icd11-mms-2026-01-base-2026-09-16/README.md
for the source file's column semantics.
"""

from __future__ import annotations

import csv
import json
from dataclasses import dataclass
from io import BytesIO
from pathlib import Path

import sqlalchemy as sa
import sqlalchemy.orm as so
from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Font, PatternFill

from app import db
from app.models import MasIcd11Mms
from app.services.icd10_2019_2_service import get_icd10_2019_2_coding_context
from app.services.icd_coding_value import extract_icd_code
from app.services.icd_search_vocabulary_service import (
    CLASSIFICATION_ICD11,
    FUZZY_MIN_QUERY_LEN,
    RESULT_TIER_EXPANDED,
    RESULT_TIER_FOCUSED,
    fuzzy_vocabulary_matches,
    merge_vocabulary_results,
    result_tier,
    spelling_like_clauses,
    spelling_variants,
    vocabulary_matches,
)

DEFAULT_ICD11_RELEASE = "2026-01"
DEFAULT_ICD11_MMS_EXPORT_PATH = Path(
    "docs/icd-causegrp-mappings/migration-artifacts/"
    "icd11-mms-2026-01-base-2026-09-16/SimpleTabulation-ICD-11-MMS-en.txt"
)
DEFAULT_ICD11_MMS_CSV_PATH = Path("resource/icd11_mms_2026_01_hierarchy.csv")
SOURCE_VERSION = "ICD-11-MMS-2026-01"

SEX_SELECTABLE_OPTIONS = ("both", "female", "male")
AGE_GROUP_SELECTABLE_OPTIONS = ("all", "neonate", "infant", "child", "adult")
# Same values the ICD-10 catalog carries; informational only (coding search
# reads is_coding_selectable/sex/age, never policy_status).
POLICY_STATUS_OPTIONS = ("unreviewed", "reviewed")
CODING_FILTER_OPTIONS = ("any", "active", "disabled")
POLICY_EDITABLE_CLASS_KINDS = frozenset({"category"})

_IMPORT_BATCH_SIZE = 1000
_CODING_MIN_QUERY_LEN = 2
_CODING_MAX_RESULTS = 30

# WHO's export header/row shape (see the frozen export's README.md).
_EXPECTED_HEADER_CELLS = 20
_EXPECTED_ROW_CELLS = 19
_COLUMNS = (
    "foundation_uri",
    "linearization_uri",
    "code",
    "block_id",
    "title",
    "class_kind",
    "depth_in_kind",
    "is_residual",
    "chapter_no",
    "browser_link",
    "is_leaf",
    "primary_tabulation",
    "grouping1",
    "grouping2",
    "grouping3",
    "grouping4",
    "grouping5",
    "coding_note",
    "parent",
)


@dataclass(frozen=True)
class Icd11MmsImportResult:
    inserted: int
    updated: int
    deactivated: int
    total_rows: int


@dataclass(frozen=True)
class Icd11MmsPolicyUpdate:
    is_coding_selectable: bool | None
    sex_selectable: str | None
    age_group_selectable: str | None
    restriction_note: str | None
    # None means "leave the row's policy_status as it is".
    policy_status: str | None = None

    def values(self) -> dict:
        values = {
            "is_coding_selectable": self.is_coding_selectable,
            "sex_selectable": self.sex_selectable,
            "age_group_selectable": self.age_group_selectable,
            "restriction_note": self.restriction_note,
        }
        if self.policy_status is not None:
            values["policy_status"] = self.policy_status
        return values


@dataclass(frozen=True)
class Icd11MmsPolicyImportResult:
    total_items: int
    updated_items: int
    reset_items: int
    skipped_items: list[dict[str, str]]


def _optional_text(value: str | None) -> str | None:
    cleaned = (value or "").strip()
    return cleaned or None


def _parse_who_bool(value: str | None) -> bool:
    return (value or "").strip().lower() == "true"


def _parse_optional_bool(value: str | None) -> bool | None:
    normalized = (value or "").strip().lower()
    if not normalized:
        return None
    if normalized == "true":
        return True
    if normalized == "false":
        return False
    return None


def _strip_title_depth_prefix(title: str) -> str:
    """Strip WHO's ``- `` depth-indent prefix from a title cell."""
    stripped = title or ""
    while stripped.startswith("- "):
        stripped = stripped[2:]
    return stripped.strip()


_EMBEDDED_LF_PLACEHOLDER = "\x00LF\x00"


def _parse_tab_row(line: str) -> list[str]:
    """Parse one CRLF-delimited row with a tab-delimited CSV reader.

    Coding notes may carry bare (unquoted) embedded line feeds; since the
    row was already isolated by splitting on literal CRLF, any remaining
    ``\\n`` in ``line`` is such embedded data, not a row terminator. It is
    swapped for a placeholder so python's csv reader (which otherwise raises
    on a newline inside an unquoted field) can parse the row, then restored
    in each resulting cell.
    """
    safe_line = line.replace("\n", _EMBEDDED_LF_PLACEHOLDER)
    cells = next(csv.reader([safe_line], delimiter="\t"))
    return [cell.replace(_EMBEDDED_LF_PLACEHOLDER, "\n") for cell in cells]


def _iter_export_rows(export_path: Path):
    """Yield each data row of the frozen export as a dict of `_COLUMNS`.

    Splits strictly on CRLF (not any bare line feed), since coding notes
    contain unquoted embedded LFs; each row is then parsed with a
    tab-delimited CSV reader because title cells are quoted.
    """
    with export_path.open("rb") as handle:
        raw = handle.read()
    text = raw.decode("utf-8-sig")
    lines = text.split("\r\n")
    if lines and lines[-1] == "":
        lines = lines[:-1]
    if not lines:
        raise ValueError(f"ICD-11 export is empty: {export_path}")

    header_cells = _parse_tab_row(lines[0])
    if len(header_cells) != _EXPECTED_HEADER_CELLS:
        raise ValueError(
            f"ICD-11 export header has {len(header_cells)} cells, "
            f"expected {_EXPECTED_HEADER_CELLS}: {export_path}"
        )

    for line_no, line in enumerate(lines[1:], start=2):
        cells = _parse_tab_row(line)
        if len(cells) != _EXPECTED_ROW_CELLS:
            raise ValueError(
                f"ICD-11 export row {line_no} has {len(cells)} cells, "
                f"expected {_EXPECTED_ROW_CELLS}: {export_path}"
            )
        yield dict(zip(_COLUMNS, cells))


def _build_foundation_to_linearization_map(export_path: Path) -> dict[str, str]:
    mapping: dict[str, str] = {}
    for row in _iter_export_rows(export_path):
        foundation_uri = _optional_text(row["foundation_uri"])
        if foundation_uri:
            mapping[foundation_uri] = row["linearization_uri"]
    return mapping


def import_icd11_mms_from_export(
    export_path: str | Path = DEFAULT_ICD11_MMS_EXPORT_PATH,
    *,
    release: str = DEFAULT_ICD11_RELEASE,
    apply_policy_columns: bool = False,
) -> Icd11MmsImportResult:
    """Stream-import the frozen WHO Simple Tabulation export.

    Idempotent upsert on ``(release, linearization_uri)``. Rows missing from
    the export are marked ``is_active=False`` and never deleted. Local
    policy columns are preserved on rerun unless ``apply_policy_columns``.
    """
    export_path = Path(export_path)
    if not export_path.exists():
        raise ValueError(f"ICD-11 export not found: {export_path}")

    foundation_to_linearization = _build_foundation_to_linearization_map(export_path)

    existing_rows = {
        row.linearization_uri: row
        for row in db.session.scalars(
            sa.select(MasIcd11Mms).where(MasIcd11Mms.release == release)
        ).all()
    }

    inserted = 0
    updated = 0
    seen_uris: set[str] = set()
    total_rows = 0
    pending = 0

    for sort_order, row in enumerate(_iter_export_rows(export_path), start=1):
        total_rows += 1
        linearization_uri = row["linearization_uri"].strip()
        seen_uris.add(linearization_uri)
        record = existing_rows.get(linearization_uri)
        is_new = record is None
        if is_new:
            record = MasIcd11Mms(release=release, linearization_uri=linearization_uri)
            db.session.add(record)
            existing_rows[linearization_uri] = record
            inserted += 1
        else:
            updated += 1

        parent_foundation_uri = _optional_text(row["parent"])

        record.foundation_uri = _optional_text(row["foundation_uri"])
        record.code = _optional_text(row["code"])
        record.block_id = _optional_text(row["block_id"])
        record.title = _strip_title_depth_prefix(row["title"])
        record.class_kind = row["class_kind"].strip()
        depth_raw = _optional_text(row["depth_in_kind"])
        record.depth_in_kind = int(depth_raw) if depth_raw is not None else None
        record.chapter_no = _optional_text(row["chapter_no"])
        record.is_residual = _parse_who_bool(row["is_residual"])
        record.is_leaf = _parse_who_bool(row["is_leaf"])
        primary_tabulation_raw = _optional_text(row["primary_tabulation"])
        record.primary_tabulation = (
            _parse_who_bool(primary_tabulation_raw) if primary_tabulation_raw is not None else None
        )
        record.coding_note = _optional_text(row["coding_note"])
        record.sort_order = sort_order
        record.parent_foundation_uri = parent_foundation_uri
        record.parent_linearization_uri = (
            foundation_to_linearization.get(parent_foundation_uri)
            if parent_foundation_uri
            else None
        )
        record.source_version = SOURCE_VERSION
        record.source_path = str(export_path)
        record.is_active = True

        if is_new or apply_policy_columns:
            record.is_coding_selectable = None
            record.sex_selectable = None
            record.age_group_selectable = None
            record.policy_status = "unreviewed"
            record.restriction_note = None

        pending += 1
        if pending >= _IMPORT_BATCH_SIZE:
            db.session.flush()
            pending = 0

    deactivated = 0
    for linearization_uri, record in existing_rows.items():
        if linearization_uri in seen_uris or not record.is_active:
            continue
        record.is_active = False
        deactivated += 1

    db.session.commit()
    return Icd11MmsImportResult(
        inserted=inserted,
        updated=updated,
        deactivated=deactivated,
        total_rows=total_rows,
    )


def generate_icd11_mms_seed_csv(
    export_path: str | Path = DEFAULT_ICD11_MMS_EXPORT_PATH,
    csv_path: str | Path = DEFAULT_ICD11_MMS_CSV_PATH,
) -> int:
    """Generate the checked-in seed CSV consumed by the seeding migration.

    Returns the number of rows written.
    """
    export_path = Path(export_path)
    csv_path = Path(csv_path)
    foundation_to_linearization = _build_foundation_to_linearization_map(export_path)
    csv_path.parent.mkdir(parents=True, exist_ok=True)

    fieldnames = [
        "linearization_uri",
        "foundation_uri",
        "code",
        "block_id",
        "title",
        "class_kind",
        "depth_in_kind",
        "chapter_no",
        "is_residual",
        "is_leaf",
        "primary_tabulation",
        "coding_note",
        "sort_order",
        "parent_foundation_uri",
        "parent_linearization_uri",
    ]
    row_count = 0
    with csv_path.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.DictWriter(handle, fieldnames=fieldnames)
        writer.writeheader()
        for sort_order, row in enumerate(_iter_export_rows(export_path), start=1):
            parent_foundation_uri = _optional_text(row["parent"])
            primary_tabulation_raw = _optional_text(row["primary_tabulation"])
            writer.writerow(
                {
                    "linearization_uri": row["linearization_uri"].strip(),
                    "foundation_uri": _optional_text(row["foundation_uri"]) or "",
                    "code": _optional_text(row["code"]) or "",
                    "block_id": _optional_text(row["block_id"]) or "",
                    "title": _strip_title_depth_prefix(row["title"]),
                    "class_kind": row["class_kind"].strip(),
                    "depth_in_kind": _optional_text(row["depth_in_kind"]) or "",
                    "chapter_no": _optional_text(row["chapter_no"]) or "",
                    "is_residual": "true" if _parse_who_bool(row["is_residual"]) else "false",
                    "is_leaf": "true" if _parse_who_bool(row["is_leaf"]) else "false",
                    "primary_tabulation": (
                        ("true" if _parse_who_bool(primary_tabulation_raw) else "false")
                        if primary_tabulation_raw is not None
                        else ""
                    ),
                    "coding_note": _optional_text(row["coding_note"]) or "",
                    "sort_order": sort_order,
                    "parent_foundation_uri": parent_foundation_uri or "",
                    "parent_linearization_uri": (
                        foundation_to_linearization.get(parent_foundation_uri, "")
                        if parent_foundation_uri
                        else ""
                    ),
                }
            )
            row_count += 1
    return row_count


def get_icd11_mms_stats(release: str = DEFAULT_ICD11_RELEASE) -> dict[str, int]:
    rows = db.session.execute(
        sa.select(
            sa.func.count().label("total_rows"),
            sa.func.count().filter(MasIcd11Mms.is_active.is_(True)).label("active_rows"),
            sa.func.count().filter(MasIcd11Mms.class_kind == "chapter").label("chapters"),
            sa.func.count().filter(MasIcd11Mms.class_kind == "block").label("blocks"),
            sa.func.count().filter(MasIcd11Mms.class_kind == "category").label("categories"),
            sa.func.count().filter(MasIcd11Mms.is_residual.is_(True)).label("residual_rows"),
            sa.func.count().filter(MasIcd11Mms.is_leaf.is_(True)).label("leaf_rows"),
        ).where(MasIcd11Mms.release == release)
    ).mappings().one()
    return {key: int(value or 0) for key, value in rows.items()}


def _serialize_row(row: MasIcd11Mms, *, child_count: int | None = None) -> dict:
    return {
        "id": str(row.id),
        "release": row.release,
        "linearization_uri": row.linearization_uri,
        "foundation_uri": row.foundation_uri,
        "code": row.code,
        "block_id": row.block_id,
        "title": row.title,
        "class_kind": row.class_kind,
        "depth_in_kind": row.depth_in_kind,
        "chapter_no": row.chapter_no,
        "is_residual": row.is_residual,
        "is_leaf": row.is_leaf,
        "primary_tabulation": row.primary_tabulation,
        "coding_note": row.coding_note,
        "sort_order": row.sort_order,
        "parent_foundation_uri": row.parent_foundation_uri,
        "parent_linearization_uri": row.parent_linearization_uri,
        "child_count": child_count,
        "is_coding_selectable": row.is_coding_selectable,
        "sex_selectable": row.sex_selectable,
        "age_group_selectable": row.age_group_selectable,
        "policy_status": row.policy_status,
        "restriction_note": row.restriction_note,
        "is_policy_editable": row.class_kind in POLICY_EDITABLE_CLASS_KINDS,
        "is_active": row.is_active,
    }


def _policy_filters_applied(*, coding_filter: str, sex_filter: str, age_filter: str) -> bool:
    return coding_filter != "any" or sex_filter != "any" or age_filter != "any"


def _policy_filter_clause(*, coding_filter: str, sex_filter: str, age_filter: str):
    clause = sa.true()
    if coding_filter == "active":
        clause = sa.and_(clause, MasIcd11Mms.is_coding_selectable.is_(True))
    elif coding_filter == "disabled":
        clause = sa.and_(
            clause,
            sa.or_(
                MasIcd11Mms.is_coding_selectable.is_(False),
                MasIcd11Mms.is_coding_selectable.is_(None),
            ),
        )
    if sex_filter != "any":
        clause = sa.and_(clause, MasIcd11Mms.sex_selectable == sex_filter)
    if age_filter != "any":
        clause = sa.and_(clause, MasIcd11Mms.age_group_selectable == age_filter)
    return clause


def _policy_filter_hits_cte(
    release: str, *, coding_filter: str, sex_filter: str, age_filter: str
):
    """Recursive CTE of categories matching the filters plus all their ancestors.

    ICD-11 nests to a variable depth, so unlike ICD-10's fixed chapter/block/
    three-character/detailed levels a node is kept when it or any descendant
    category matches. UNION (not UNION ALL) dedupes shared ancestors.
    """
    hits = (
        sa.select(
            MasIcd11Mms.linearization_uri.label("linearization_uri"),
            MasIcd11Mms.parent_linearization_uri.label("parent_linearization_uri"),
        )
        .where(
            MasIcd11Mms.release == release,
            MasIcd11Mms.is_active.is_(True),
            MasIcd11Mms.class_kind.in_(tuple(POLICY_EDITABLE_CLASS_KINDS)),
            _policy_filter_clause(
                coding_filter=coding_filter, sex_filter=sex_filter, age_filter=age_filter
            ),
        )
        .cte("icd11_filter_hits", recursive=True)
    )
    parent = so.aliased(MasIcd11Mms)
    return hits.union(
        sa.select(parent.linearization_uri, parent.parent_linearization_uri)
        .join(hits, parent.linearization_uri == hits.c.parent_linearization_uri)
        .where(parent.release == release, parent.is_active.is_(True))
    )


def _status_indicator_for_row(row: MasIcd11Mms, *, selectable_child_count: int) -> str | None:
    """Tree dot for a category, mirroring ICD-10's three-character rule.

    green: selectable itself; yellow: not selectable but a direct child is;
    red: neither. Chapters and blocks carry no dot, as in the ICD-10 panel.
    """
    if row.class_kind not in POLICY_EDITABLE_CLASS_KINDS:
        return None
    if row.is_coding_selectable:
        return "green"
    if selectable_child_count > 0:
        return "yellow"
    return "red"


def list_icd11_mms_children(
    parent_linearization_uri: str | None,
    release: str = DEFAULT_ICD11_RELEASE,
    *,
    coding_filter: str = "any",
    sex_filter: str = "any",
    age_filter: str = "any",
) -> list[dict]:
    """Direct children of one node (roots when ``parent_linearization_uri`` is None).

    With any policy filter set, only children that match or have a matching
    descendant category are returned, and ``child_count`` counts only such
    children. Filter values are expected to be validated by the caller.
    """
    selectable_child_count_sq = (
        sa.select(
            MasIcd11Mms.parent_linearization_uri.label("parent_linearization_uri"),
            sa.func.count().label("child_count"),
            sa.func.count()
            .filter(MasIcd11Mms.is_coding_selectable.is_(True))
            .label("selectable_child_count"),
        )
        .where(MasIcd11Mms.release == release, MasIcd11Mms.is_active.is_(True))
        .group_by(MasIcd11Mms.parent_linearization_uri)
        .subquery()
    )
    filters_applied = _policy_filters_applied(
        coding_filter=coding_filter, sex_filter=sex_filter, age_filter=age_filter
    )
    hits = None
    if filters_applied:
        hits = _policy_filter_hits_cte(
            release, coding_filter=coding_filter, sex_filter=sex_filter, age_filter=age_filter
        )
        filtered_child_count_sq = (
            sa.select(
                hits.c.parent_linearization_uri.label("parent_linearization_uri"),
                sa.func.count().label("child_count"),
            )
            .group_by(hits.c.parent_linearization_uri)
            .subquery()
        )
        child_count_col = filtered_child_count_sq.c.child_count
    else:
        child_count_col = selectable_child_count_sq.c.child_count

    query = (
        sa.select(
            MasIcd11Mms,
            child_count_col,
            selectable_child_count_sq.c.selectable_child_count,
        )
        .outerjoin(
            selectable_child_count_sq,
            selectable_child_count_sq.c.parent_linearization_uri
            == MasIcd11Mms.linearization_uri,
        )
        .where(MasIcd11Mms.release == release, MasIcd11Mms.is_active.is_(True))
        .order_by(MasIcd11Mms.sort_order, MasIcd11Mms.linearization_uri)
    )
    if filters_applied:
        # A row is in the hit set when it matches itself or has a child in it;
        # testing that directly avoids a semi-join that sorts every hit URI.
        query = query.outerjoin(
            filtered_child_count_sq,
            filtered_child_count_sq.c.parent_linearization_uri
            == MasIcd11Mms.linearization_uri,
        ).where(
            sa.or_(
                sa.and_(
                    MasIcd11Mms.class_kind.in_(tuple(POLICY_EDITABLE_CLASS_KINDS)),
                    _policy_filter_clause(
                        coding_filter=coding_filter,
                        sex_filter=sex_filter,
                        age_filter=age_filter,
                    ),
                ),
                filtered_child_count_sq.c.child_count > 0,
            )
        )
    if parent_linearization_uri is None:
        query = query.where(MasIcd11Mms.parent_linearization_uri.is_(None))
    else:
        query = query.where(
            MasIcd11Mms.parent_linearization_uri == parent_linearization_uri
        )

    payload = []
    for row, child_count, selectable_child_count in db.session.execute(query).all():
        item = _serialize_row(row, child_count=int(child_count or 0))
        item["status_indicator"] = _status_indicator_for_row(
            row, selectable_child_count=int(selectable_child_count or 0)
        )
        payload.append(item)
    return payload


def get_icd11_mms_node_details(
    linearization_uri: str, release: str = DEFAULT_ICD11_RELEASE
) -> dict | None:
    row = db.session.scalar(
        sa.select(MasIcd11Mms).where(
            MasIcd11Mms.release == release,
            MasIcd11Mms.linearization_uri == linearization_uri,
        )
    )
    if row is None or not row.is_active:
        return None

    child_count, selectable_child_count = db.session.execute(
        sa.select(
            sa.func.count(),
            sa.func.count().filter(MasIcd11Mms.is_coding_selectable.is_(True)),
        ).where(
            MasIcd11Mms.release == release,
            MasIcd11Mms.parent_linearization_uri == linearization_uri,
            MasIcd11Mms.is_active.is_(True),
        )
    ).one()

    ancestors: list[dict[str, str]] = []
    current_uri = row.parent_linearization_uri
    while current_uri:
        ancestor = db.session.scalar(
            sa.select(MasIcd11Mms).where(
                MasIcd11Mms.release == release,
                MasIcd11Mms.linearization_uri == current_uri,
            )
        )
        if ancestor is None or not ancestor.is_active:
            break
        ancestors.append(
            {
                "linearization_uri": ancestor.linearization_uri,
                "title": ancestor.title,
                "class_kind": ancestor.class_kind,
            }
        )
        current_uri = ancestor.parent_linearization_uri
    ancestors.reverse()

    payload = _serialize_row(row, child_count=int(child_count or 0))
    payload["status_indicator"] = _status_indicator_for_row(
        row, selectable_child_count=int(selectable_child_count or 0)
    )
    payload["ancestors"] = ancestors
    return payload


def get_icd11_mms_policy_options() -> dict[str, list[str]]:
    return {
        "sex_selectable": list(SEX_SELECTABLE_OPTIONS),
        "age_group_selectable": list(AGE_GROUP_SELECTABLE_OPTIONS),
        "policy_status": list(POLICY_STATUS_OPTIONS),
    }


def _validate_policy_update(
    *,
    is_coding_selectable,
    sex_selectable,
    age_group_selectable,
    restriction_note,
    policy_status=None,
) -> Icd11MmsPolicyUpdate:
    if is_coding_selectable not in (True, False, None):
        raise ValueError("is_coding_selectable must be true, false, or null.")
    if sex_selectable not in (*SEX_SELECTABLE_OPTIONS, None):
        raise ValueError("sex_selectable must be one of both, female, male, or null.")
    if age_group_selectable not in (*AGE_GROUP_SELECTABLE_OPTIONS, None):
        raise ValueError(
            "age_group_selectable must be one of all, neonate, infant, child, adult, or null."
        )
    if restriction_note is not None and not isinstance(restriction_note, str):
        raise ValueError("restriction_note must be a string or null.")
    if policy_status not in (*POLICY_STATUS_OPTIONS, None):
        raise ValueError("policy_status must be one of unreviewed, reviewed, or omitted.")
    return Icd11MmsPolicyUpdate(
        is_coding_selectable=is_coding_selectable,
        sex_selectable=sex_selectable,
        age_group_selectable=age_group_selectable,
        restriction_note=_optional_text(restriction_note or ""),
        policy_status=policy_status,
    )


def _apply_policy_values(row: MasIcd11Mms, values: dict, *, dry_run: bool) -> bool:
    """Set each changed policy field on ``row``; return whether any differed.

    With ``dry_run`` the row is only compared, never mutated, so a preview
    leaves the session clean.
    """
    changed = False
    for field, value in values.items():
        if getattr(row, field) == value:
            continue
        changed = True
        if not dry_run:
            setattr(row, field, value)
    return changed


def update_icd11_mms_policy(
    linearization_uri: str,
    *,
    release: str = DEFAULT_ICD11_RELEASE,
    is_coding_selectable,
    sex_selectable,
    age_group_selectable,
    restriction_note,
    policy_status=None,
) -> dict:
    """Set one category's policy fields; ``policy_status=None`` keeps it as is.

    Raises LookupError for an unknown/inactive entity and ValueError for a
    non-category entity or an invalid value.
    """
    row = db.session.scalar(
        sa.select(MasIcd11Mms).where(
            MasIcd11Mms.release == release,
            MasIcd11Mms.linearization_uri == linearization_uri,
        )
    )
    if row is None or not row.is_active:
        raise LookupError(f"ICD-11 entity not found: {linearization_uri}")
    if row.class_kind not in POLICY_EDITABLE_CLASS_KINDS:
        raise ValueError("Policy fields are only editable for ICD-11 category entities.")

    update = _validate_policy_update(
        is_coding_selectable=is_coding_selectable,
        sex_selectable=sex_selectable,
        age_group_selectable=age_group_selectable,
        restriction_note=restriction_note,
        policy_status=policy_status,
    )
    _apply_policy_values(row, update.values(), dry_run=False)
    db.session.commit()
    return get_icd11_mms_node_details(linearization_uri, release) or _serialize_row(row)


def export_icd11_mms_policy_json(release: str = DEFAULT_ICD11_RELEASE) -> dict:
    """Curated category rows in the policy JSON format read by the importer.

    Carries ``restriction_note`` and ``policy_status`` so an export imports
    back unchanged.
    """
    rows = db.session.scalars(
        sa.select(MasIcd11Mms)
        .where(
            MasIcd11Mms.release == release,
            MasIcd11Mms.is_active.is_(True),
            MasIcd11Mms.class_kind.in_(tuple(POLICY_EDITABLE_CLASS_KINDS)),
            sa.or_(
                MasIcd11Mms.is_coding_selectable.is_not(None),
                MasIcd11Mms.sex_selectable.is_not(None),
                MasIcd11Mms.age_group_selectable.is_not(None),
                MasIcd11Mms.restriction_note.is_not(None),
                MasIcd11Mms.policy_status != "unreviewed",
            ),
        )
        .order_by(MasIcd11Mms.sort_order, MasIcd11Mms.linearization_uri)
    ).all()

    items = [
        {
            "linearization_uri": row.linearization_uri,
            "code": row.code,
            "title": row.title,
            "class_kind": row.class_kind,
            "chapter_no": row.chapter_no,
            "is_coding_selectable": row.is_coding_selectable,
            "sex_selectable": row.sex_selectable,
            "age_group_selectable": row.age_group_selectable,
            "policy_status": row.policy_status,
            "restriction_note": row.restriction_note,
        }
        for row in rows
    ]
    return {
        "release": release,
        "source_version": SOURCE_VERSION,
        "row_count": len(items),
        "items": items,
    }


_XLSX_HEADERS = (
    "ICD-11 Code",
    "Title",
    "Class Kind",
    "Chapter",
    "Residual",
    "Leaf",
    "Coding Allowed",
    "Age Selectable",
    "Sex Selectable",
    "Policy Status",
    "Restriction Note",
    "Linearization URI",
)
_XLSX_COLUMN_WIDTHS = (12, 48, 12, 9, 9, 7, 15, 15, 15, 14, 40, 48)


def export_icd11_mms_policy_xlsx(release: str = DEFAULT_ICD11_RELEASE) -> bytes:
    """Every active category of ``release`` with its policy, as an xlsx workbook.

    Write-only workbook over plain column tuples: the release has ~35k
    categories, which a regular openpyxl sheet of ORM rows holds in memory
    several times over. Styling mirrors the ICD-10 export's header.
    """
    rows = db.session.execute(
        sa.select(
            MasIcd11Mms.code,
            MasIcd11Mms.title,
            MasIcd11Mms.class_kind,
            MasIcd11Mms.chapter_no,
            MasIcd11Mms.is_residual,
            MasIcd11Mms.is_leaf,
            MasIcd11Mms.is_coding_selectable,
            MasIcd11Mms.age_group_selectable,
            MasIcd11Mms.sex_selectable,
            MasIcd11Mms.policy_status,
            MasIcd11Mms.restriction_note,
            MasIcd11Mms.linearization_uri,
        )
        .where(
            MasIcd11Mms.release == release,
            MasIcd11Mms.is_active.is_(True),
            MasIcd11Mms.class_kind.in_(tuple(POLICY_EDITABLE_CLASS_KINDS)),
        )
        .order_by(MasIcd11Mms.sort_order, MasIcd11Mms.linearization_uri)
    )

    workbook = Workbook(write_only=True)
    sheet = workbook.create_sheet("ICD11 Policy")
    for index, width in enumerate(_XLSX_COLUMN_WIDTHS):
        sheet.column_dimensions[chr(ord("A") + index)].width = width
    sheet.freeze_panes = "A2"
    header_fill = PatternFill("solid", fgColor="D9EAF7")
    header_font = Font(bold=True)
    header = []
    for label in _XLSX_HEADERS:
        cell = WriteOnlyCell(sheet, value=label)
        cell.fill = header_fill
        cell.font = header_font
        header.append(cell)
    sheet.append(header)

    row_count = 0
    for (
        code,
        title,
        class_kind,
        chapter_no,
        is_residual,
        is_leaf,
        is_coding_selectable,
        age_group_selectable,
        sex_selectable,
        policy_status,
        restriction_note,
        linearization_uri,
    ) in rows:
        sheet.append(
            [
                code,
                title,
                class_kind,
                chapter_no,
                "Yes" if is_residual else "No",
                "Yes" if is_leaf else "No",
                "Yes" if is_coding_selectable else "No",
                age_group_selectable,
                sex_selectable,
                policy_status,
                restriction_note,
                linearization_uri,
            ]
        )
        row_count += 1
    sheet.auto_filter.ref = f"A1:L{row_count + 1}"

    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def import_icd11_mms_policy_json(
    payload: str | bytes | dict,
    release: str = DEFAULT_ICD11_RELEASE,
    *,
    dry_run: bool = False,
) -> Icd11MmsPolicyImportResult:
    """Apply a policy JSON (the ``export_icd11_mms_policy_json`` format).

    Listed categories take the file's values; every other active category is
    reset to not selectable with no sex/age/note. ``policy_status`` is set
    only for items that carry it and is left alone on reset rows. With
    ``dry_run`` the counts are computed and nothing is written.

    Raises ValueError for a malformed file; bad items are reported in
    ``skipped_items`` instead.
    """
    if isinstance(payload, (str, bytes)):
        try:
            data = json.loads(payload)
        except json.JSONDecodeError as exc:
            raise ValueError("Policy import file is not valid JSON.") from exc
    elif isinstance(payload, dict):
        data = payload
    else:
        raise ValueError("Policy import payload must be JSON.")

    if not isinstance(data, dict):
        raise ValueError("Policy import JSON must be an object with an items array.")
    items = data.get("items")
    if not isinstance(items, list):
        raise ValueError("Policy import JSON must include an items array.")

    seen_uris: set[str] = set()
    updates_by_uri: dict[str, Icd11MmsPolicyUpdate] = {}
    skipped_items: list[dict[str, str]] = []

    for item in items:
        if not isinstance(item, dict):
            raise ValueError("Each policy import item must be an object.")
        raw_uri = item.get("linearization_uri")
        linearization_uri = _optional_text(raw_uri) if isinstance(raw_uri, str) else None
        if not linearization_uri:
            raise ValueError("Each policy import item must include a linearization_uri.")
        if linearization_uri in seen_uris:
            skipped_items.append({"code": linearization_uri, "reason": "duplicate_code"})
            continue
        seen_uris.add(linearization_uri)
        try:
            updates_by_uri[linearization_uri] = _validate_policy_update(
                is_coding_selectable=item.get("is_coding_selectable"),
                sex_selectable=item.get("sex_selectable"),
                age_group_selectable=item.get("age_group_selectable"),
                restriction_note=item.get("restriction_note"),
                policy_status=item.get("policy_status"),
            )
        except ValueError:
            skipped_items.append(
                {"code": linearization_uri, "reason": "invalid_policy_values"}
            )

    rows = db.session.scalars(
        sa.select(MasIcd11Mms).where(
            MasIcd11Mms.release == release,
            MasIcd11Mms.is_active.is_(True),
            MasIcd11Mms.class_kind.in_(tuple(POLICY_EDITABLE_CLASS_KINDS)),
        )
    ).all()
    row_map = {row.linearization_uri: row for row in rows}

    updated_items = 0
    imported_uris: set[str] = set()
    for linearization_uri, update in updates_by_uri.items():
        row = row_map.get(linearization_uri)
        if row is None:
            skipped_items.append(
                {"code": linearization_uri, "reason": "unknown_or_non_editable_code"}
            )
            continue
        if _apply_policy_values(row, update.values(), dry_run=dry_run):
            updated_items += 1
        imported_uris.add(linearization_uri)

    reset_values = {
        "is_coding_selectable": False,
        "sex_selectable": None,
        "age_group_selectable": None,
        "restriction_note": None,
    }
    reset_items = 0
    for row in rows:
        if row.linearization_uri in imported_uris:
            continue
        if _apply_policy_values(row, reset_values, dry_run=dry_run):
            reset_items += 1

    if not dry_run:
        db.session.commit()
    return Icd11MmsPolicyImportResult(
        total_items=len(items),
        updated_items=updated_items,
        reset_items=reset_items,
        skipped_items=skipped_items,
    )


def _normalize_query(raw_query: str) -> str:
    return " ".join((raw_query or "").strip().lower().split())


def _coding_policy_clause(*, age_group: str | None, sex: str | None):
    clause = MasIcd11Mms.is_coding_selectable.is_(True)
    if age_group:
        clause = sa.and_(
            clause,
            sa.or_(
                MasIcd11Mms.age_group_selectable == "all",
                MasIcd11Mms.age_group_selectable == age_group,
            ),
        )
    if sex:
        clause = sa.and_(
            clause,
            sa.or_(
                MasIcd11Mms.sex_selectable == "both",
                MasIcd11Mms.sex_selectable == sex,
            ),
        )
    return clause


def validate_icd11_mms_coding_value_for_submission(
    va_sid: str,
    value: str | None,
    release: str = DEFAULT_ICD11_RELEASE,
) -> None:
    """Reject a COD value that is not a selectable ICD-11 code for this death.

    Mirrors ``validate_icd10_2019_2_coding_value_for_submission``: the code
    must be an active category of ``release`` that the coding policy allows
    for the submission's age and sex. Raises ``ValueError`` when it is not,
    ``LookupError`` when the submission does not exist.
    """
    code = extract_icd_code(value, "icd11")
    if code is None:
        raise ValueError("Select a valid ICD-11 code.")

    context = get_icd10_2019_2_coding_context(va_sid)
    if context is None:
        raise LookupError(f"Submission not found: {va_sid}")

    is_allowed = db.session.scalar(
        sa.select(
            sa.exists().where(
                MasIcd11Mms.release == release,
                MasIcd11Mms.code == code,
                MasIcd11Mms.is_active.is_(True),
                MasIcd11Mms.class_kind.in_(tuple(POLICY_EDITABLE_CLASS_KINDS)),
                _coding_policy_clause(age_group=context["age_group"], sex=context["sex"]),
            )
        )
    )
    if not is_allowed:
        raise ValueError(f"{code} is not selectable for this submission.")


def search_icd11_mms(
    query: str,
    va_sid: str | None = None,
    release: str = DEFAULT_ICD11_RELEASE,
    limit: int = _CODING_MAX_RESULTS,
    *,
    age_group: str | None = None,
    sex: str | None = None,
) -> list[dict[str, str | bool | None]]:
    """Search selectable ICD-11 categories by code or title.

    When ``va_sid`` is given, results are filtered by that submission's age
    and sex policy, mirroring ``search_icd10_2019_2_coding_choices`` (any
    ``age_group``/``sex`` passed in are ignored — the submission is
    authoritative). Without a ``va_sid``, passing ``age_group`` and/or ``sex``
    explicitly applies the same policy filter without a submission to fake
    (used by the coding-search help demo); omitting both searches unfiltered.
    Matches from the admin-managed search vocabulary (``mas_icd_search_terms``)
    are prepended and flagged; expansion goes through the same filters, so it
    never bypasses coding policy (docs/policy/icd-coding-search-vocabulary.md).
    """
    normalized_query = _normalize_query(query)
    if len(normalized_query) < _CODING_MIN_QUERY_LEN:
        return []

    apply_policy = va_sid is not None or age_group is not None or sex is not None
    if va_sid is not None:
        context = get_icd10_2019_2_coding_context(va_sid)
        if context is None:
            raise LookupError(f"Submission not found: {va_sid}")
        age_group = context["age_group"]
        sex = context["sex"]

    # Escape LIKE wildcards so '%' and '_' in a query match literally.
    lower_code = sa.func.lower(sa.func.coalesce(MasIcd11Mms.code, ""))
    lower_title = sa.func.lower(MasIcd11Mms.title)

    # Spelling fold (UK/US) and hyphen forms OR-ed in; every variant ranks
    # at the same band as the typed query. The catalogues carry mixed UK/US
    # spellings, so both directions matter here.
    match_clause, rank_expr = spelling_like_clauses(
        (lower_code, lower_title),
        spelling_variants(normalized_query),
        escape="\\",
    )

    filters = [
        MasIcd11Mms.release == release,
        MasIcd11Mms.is_active.is_(True),
        MasIcd11Mms.class_kind.in_(tuple(POLICY_EDITABLE_CLASS_KINDS)),
        match_clause,
    ]
    if apply_policy:
        filters.append(_coding_policy_clause(age_group=age_group, sex=sex))

    rows = db.session.scalars(
        sa.select(MasIcd11Mms)
        .where(*filters)
        .order_by(rank_expr, MasIcd11Mms.sort_order, MasIcd11Mms.linearization_uri)
        .limit(min(limit, _CODING_MAX_RESULTS))
    ).all()

    lexical_results = [
        {
            "icd_code": row.code,
            "icd_to_display": f"{row.code} {row.title}" if row.code else row.title,
            "title": row.title,
            "linearization_uri": row.linearization_uri,
            "class_kind": row.class_kind,
            "tier": result_tier(
                code=row.code, title=row.title, normalized_query=normalized_query
            ),
        }
        for row in rows
    ]
    # digitva-wqc: re-rank the already-fetched rows so a focused tier
    # (word-boundary/qualifier-stripped match, computed above) lists
    # before an expanded one; a stable sort keeps SQL's own ordering
    # within each tier. This never changes WHICH rows SQL fetched.
    lexical_results.sort(key=lambda result: result["tier"] != RESULT_TIER_FOCUSED)
    # Clinician shorthand (MI, CVA, Kochs) shares no substring with any ICD
    # title, so the lexical pass alone never reaches it. Vocabulary matches
    # rank first, on top of the capped lexical list: a vocabulary hit never
    # pushes a title match out (vocabulary links per term are a handful).
    merged = merge_vocabulary_results(
        lexical_results,
        _resolve_vocabulary_links_icd11(
            vocabulary_matches(query, classification=CLASSIFICATION_ICD11),
            release=release,
            apply_policy=apply_policy,
            age_group=age_group,
            sex=sex,
        ),
    )
    if merged or len(normalized_query) < FUZZY_MIN_QUERY_LEN:
        # Vocabulary hits ride on top of the default cap; a caller asking
        # for fewer rows still gets at most what it asked for.
        return merged if limit >= _CODING_MAX_RESULTS else merged[:limit]
    # digitva-1ht: typo-tolerant fallback, only when the normal search for
    # this query found nothing at all.
    return merge_vocabulary_results(
        _fuzzy_title_icd11_hits(
            normalized_query,
            release=release,
            apply_policy=apply_policy,
            age_group=age_group,
            sex=sex,
        ),
        _resolve_vocabulary_links_icd11(
            fuzzy_vocabulary_matches(query, classification=CLASSIFICATION_ICD11),
            release=release,
            apply_policy=apply_policy,
            age_group=age_group,
            sex=sex,
            fuzzy=True,
        ),
    )


def _resolve_vocabulary_links_icd11(
    links: list[dict],
    *,
    release: str,
    apply_policy: bool,
    age_group: str | None,
    sex: str | None,
    fuzzy: bool = False,
) -> list[dict]:
    """Vocabulary links resolved through the same catalogue filters the
    lexical search applies (release, active, category class kind and, when a
    submission is in play, its age and sex policy). A target the policy
    filters out contributes nothing — the vocabulary never bypasses coding
    policy. Shared by the exact vocabulary lookup and the fuzzy fallback
    (digitva-1ht)."""
    if not links:
        return []
    filters = [
        MasIcd11Mms.release == release,
        MasIcd11Mms.code.in_([link["icd_code"] for link in links]),
        MasIcd11Mms.is_active.is_(True),
        MasIcd11Mms.class_kind.in_(tuple(POLICY_EDITABLE_CLASS_KINDS)),
    ]
    if apply_policy:
        filters.append(_coding_policy_clause(age_group=age_group, sex=sex))
    rows_by_code = {
        row.code: row for row in db.session.scalars(sa.select(MasIcd11Mms).where(*filters))
    }
    results = []
    for link in links:
        row = rows_by_code.get(link["icd_code"])
        if row is None:
            continue
        result = {
            "icd_code": row.code,
            "icd_to_display": (
                f"{link['term']} — {row.code} {row.title}" if row.code else f"{link['term']} — {row.title}"
            ),
            "title": row.title,
            "linearization_uri": row.linearization_uri,
            "class_kind": row.class_kind,
            "vocabulary": True,
            "tier": RESULT_TIER_FOCUSED,
        }
        if fuzzy:
            result["fuzzy"] = True
        results.append(result)
    return results


def _fuzzy_title_icd11_hits(
    normalized_query: str,
    *,
    release: str,
    apply_policy: bool,
    age_group: str | None,
    sex: str | None,
) -> list[dict]:
    """pg_trgm word-similarity fallback over titles (digitva-1ht): runs only
    when the normal lexical + vocabulary search for this query is empty.
    Uses the ``<%`` word-similarity operator (indexable on the title side,
    ``ix_mas_icd11_mms_title_trgm``) so a catalogue-wide scan stays cheap;
    the default ``pg_trgm.word_similarity_threshold`` (0.6) is the cut."""
    lower_title = sa.func.lower(MasIcd11Mms.title)
    similarity = sa.func.word_similarity(normalized_query, lower_title)
    filters = [
        MasIcd11Mms.release == release,
        MasIcd11Mms.is_active.is_(True),
        MasIcd11Mms.class_kind.in_(tuple(POLICY_EDITABLE_CLASS_KINDS)),
        sa.literal(normalized_query).op("<%")(lower_title),
    ]
    if apply_policy:
        filters.append(_coding_policy_clause(age_group=age_group, sex=sex))
    rows = db.session.scalars(
        sa.select(MasIcd11Mms)
        .where(*filters)
        .order_by(similarity.desc())
        .limit(_CODING_MAX_RESULTS)
    ).all()
    return [
        {
            "icd_code": row.code,
            "icd_to_display": f"{row.code} {row.title}" if row.code else row.title,
            "title": row.title,
            "linearization_uri": row.linearization_uri,
            "class_kind": row.class_kind,
            "tier": RESULT_TIER_EXPANDED,
            "fuzzy": True,
        }
        for row in rows
    ]
