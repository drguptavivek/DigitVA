"""Health-system organization master data: levels, units, cadres, workers.

Plan: docs/planning/health-system-organization-model-plan.md
Policy: docs/policy/organization-model.md

All codes are unique within one project only. Unit paths are ltree values
built from unit codes, so subtree queries use ``path <@ :ancestor_path``.
Every write here is validated against the project's level definitions and
never deletes rows: deactivation is the only removal.
"""
from __future__ import annotations

import csv
import io
import logging
import re
import uuid
from dataclasses import dataclass, field
from decimal import Decimal, InvalidOperation

import sqlalchemy as sa
from openpyxl import Workbook, load_workbook

from app import db
from app.models import (
    MapOrgLevelCadre,
    MasCadre,
    MasOrgLevel,
    MasOrgUnit,
    MasOrgUnitWorker,
    VaProjectMaster,
    VaUsers,
)

log = logging.getLogger(__name__)

__all__ = [
    "OrganizationError",
    "ImportPlan",
    "EXPORT_SHEETS",
    "normalize_code",
    "normalize_level_code",
    "odk_field_name_for_level",
    "odk_choice_list_name_for_level",
    # levels
    "list_levels", "create_level", "update_level",
    # units
    "list_units", "get_unit_tree", "create_unit", "update_unit", "set_unit_active",
    "find_unit_by_code", "subtree_unit_ids",
    # cadres and level permissions
    "list_cadres", "create_cadre", "update_cadre",
    "list_level_cadres", "upsert_level_cadre", "get_level_cadre_permission",
    # workers
    "list_workers", "create_worker", "update_worker",
    # template, export, import
    "seed_default_organization",
    "export_organization_rows", "export_organization_xlsx", "export_organization_csv",
    "export_odk_choices_rows", "export_odk_choices_csv",
    "parse_organization_workbook", "import_organization",
    # serializers
    "serialize_level", "serialize_unit", "serialize_cadre", "serialize_level_cadre", "serialize_worker",
    # export dimensions
    "resolve_org_unit_export_labels",
]

# ltree labels accept [A-Za-z0-9_]; codes are stored upper-case.
CODE_RE = re.compile(r"^[A-Z0-9_]{1,32}$")
LEVEL_CODE_RE = re.compile(r"^[a-z][a-z0-9_]{0,31}$")

# Seed template: District > Taluka (optional) > CHC > PHC > Sub-centre > Village.
DEFAULT_LEVEL_TEMPLATE: tuple[tuple[str, str, int, bool], ...] = (
    ("district", "District", 1, False),
    ("taluka", "Taluka / Sub-divisional Hospital", 2, True),
    ("chc", "Community Health Centre", 3, False),
    ("phc", "PHC / UPHC / AAM-PHC", 4, False),
    ("subcentre", "Sub-centre / AAM-SHC", 5, False),
    ("village", "Village", 6, False),
)
DEFAULT_CADRE_TEMPLATE: tuple[tuple[str, str], ...] = (
    ("SMO", "Senior Medical Officer"),
    ("MO", "Medical Officer"),
    ("CHO", "Community Health Officer"),
    ("MPW", "Multipurpose Worker"),
    ("ANM", "Auxiliary Nurse Midwife"),
    ("ASHA", "Accredited Social Health Activist"),
)
# (level_code, cadre_code) -> (can_fill_va_form, can_code_va_form)
DEFAULT_LEVEL_CADRE_TEMPLATE: dict[tuple[str, str], tuple[bool, bool]] = {
    ("chc", "SMO"): (False, True),
    ("chc", "MO"): (False, True),
    ("phc", "MO"): (False, True),
    ("phc", "CHO"): (True, False),
    ("subcentre", "CHO"): (True, False),
    ("subcentre", "MPW"): (True, False),
    ("subcentre", "ANM"): (True, False),
    ("village", "ASHA"): (True, False),
}

EXPORT_SHEETS = ("levels", "units", "cadres", "level_cadres", "workers")


class OrganizationError(ValueError):
    """Validation failure reported to the caller (HTTP 400 in routes)."""


# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------


def normalize_code(raw: object, *, what: str = "Code") -> str:
    code = str(raw or "").strip().upper().replace("-", "_").replace(" ", "_")
    if not CODE_RE.match(code):
        raise OrganizationError(
            f"{what} must be 1-32 characters of letters, digits or underscore (got {raw!r})."
        )
    return code


def normalize_level_code(raw: object) -> str:
    code = str(raw or "").strip().lower().replace("-", "_").replace(" ", "_")
    if not LEVEL_CODE_RE.match(code):
        raise OrganizationError(
            "Level code must start with a letter and use lower-case letters, digits or "
            f"underscore, at most 32 characters (got {raw!r})."
        )
    return code


def odk_field_name_for_level(level_code: str) -> str:
    """The XLSForm ``survey`` field carrying this level's unit code."""
    return f"org_{level_code}_code"


def odk_choice_list_name_for_level(level_code: str) -> str:
    """The XLSForm ``choices`` list_name holding this level's units."""
    return f"org_{level_code}"


def _clean_text(raw: object, *, max_len: int | None = None, required: bool = False, what: str = "Value") -> str | None:
    value = str(raw).strip() if raw is not None else ""
    if not value:
        if required:
            raise OrganizationError(f"{what} is required.")
        return None
    if max_len is not None and len(value) > max_len:
        raise OrganizationError(f"{what} must be at most {max_len} characters.")
    return value


def _clean_coordinate(raw: object, *, what: str, limit: int) -> Decimal | None:
    if raw is None or str(raw).strip() == "":
        return None
    try:
        value = Decimal(str(raw).strip())
    except InvalidOperation as exc:
        raise OrganizationError(f"{what} must be a decimal number.") from exc
    if value < -limit or value > limit:
        raise OrganizationError(f"{what} must be between -{limit} and {limit}.")
    return value.quantize(Decimal("0.000001"))


def _clean_url(raw: object) -> str | None:
    value = _clean_text(raw, what="Google Maps URL")
    if value is None:
        return None
    if not (value.startswith("https://") or value.startswith("http://")):
        raise OrganizationError("Google Maps URL must start with http:// or https://.")
    return value


def _to_bool(raw: object) -> bool:
    if isinstance(raw, bool):
        return raw
    return str(raw or "").strip().lower() in {"1", "true", "yes", "y", "t"}


def _get_project(project_id: str) -> VaProjectMaster:
    project = db.session.get(VaProjectMaster, project_id)
    if project is None:
        raise OrganizationError(f"Project not found: {project_id}")
    return project


def _uuid(raw: object, *, what: str) -> uuid.UUID:
    if isinstance(raw, uuid.UUID):
        return raw
    try:
        return uuid.UUID(str(raw))
    except (TypeError, ValueError) as exc:
        raise OrganizationError(f"{what} is not a valid id.") from exc


# ---------------------------------------------------------------------------
# Serializers
# ---------------------------------------------------------------------------


def serialize_level(level: MasOrgLevel) -> dict:
    return {
        "org_level_id": str(level.org_level_id),
        "project_id": level.project_id,
        "level_code": level.level_code,
        "level_name": level.level_name,
        "depth": level.depth,
        "is_optional": level.is_optional,
        "is_active": level.is_active,
        "odk_field_name": odk_field_name_for_level(level.level_code),
        "odk_choice_list_name": odk_choice_list_name_for_level(level.level_code),
    }


def serialize_unit(unit: MasOrgUnit, *, level: MasOrgLevel | None = None, parent_code: str | None = None) -> dict:
    level = level or unit.level
    return {
        "org_unit_id": str(unit.org_unit_id),
        "project_id": unit.project_id,
        "org_level_id": str(unit.org_level_id),
        "level_code": level.level_code,
        "level_name": level.level_name,
        "depth": level.depth,
        "parent_org_unit_id": str(unit.parent_org_unit_id) if unit.parent_org_unit_id else None,
        "parent_code": parent_code,
        "unit_code": unit.unit_code,
        "unit_name": unit.unit_name,
        "path": str(unit.path),
        "address": unit.address,
        "phone": unit.phone,
        "latitude": str(unit.latitude) if unit.latitude is not None else None,
        "longitude": str(unit.longitude) if unit.longitude is not None else None,
        "google_maps_url": unit.google_maps_url,
        "remarks": unit.remarks,
        "is_active": unit.is_active,
    }


def serialize_cadre(cadre: MasCadre) -> dict:
    return {
        "cadre_id": str(cadre.cadre_id),
        "project_id": cadre.project_id,
        "cadre_code": cadre.cadre_code,
        "cadre_name": cadre.cadre_name,
        "is_active": cadre.is_active,
    }


def serialize_level_cadre(row: MapOrgLevelCadre, *, level: MasOrgLevel, cadre: MasCadre) -> dict:
    return {
        "level_cadre_id": str(row.level_cadre_id),
        "org_level_id": str(row.org_level_id),
        "level_code": level.level_code,
        "cadre_id": str(row.cadre_id),
        "cadre_code": cadre.cadre_code,
        "can_fill_va_form": row.can_fill_va_form,
        "can_code_va_form": row.can_code_va_form,
        "is_active": row.is_active,
    }


def serialize_worker(worker: MasOrgUnitWorker, *, unit: MasOrgUnit, cadre: MasCadre, user_email: str | None) -> dict:
    return {
        "worker_id": str(worker.worker_id),
        "project_id": worker.project_id,
        "org_unit_id": str(worker.org_unit_id),
        "unit_code": unit.unit_code,
        "unit_name": unit.unit_name,
        "cadre_id": str(worker.cadre_id),
        "cadre_code": cadre.cadre_code,
        "worker_code": worker.worker_code,
        "worker_name": worker.worker_name,
        "phone": worker.phone,
        "user_id": str(worker.user_id) if worker.user_id else None,
        "user_email": user_email,
        "remarks": worker.remarks,
        "is_active": worker.is_active,
    }


# ---------------------------------------------------------------------------
# Levels
# ---------------------------------------------------------------------------


def list_levels(project_id: str, *, include_inactive: bool = False) -> list[MasOrgLevel]:
    stmt = sa.select(MasOrgLevel).where(MasOrgLevel.project_id == project_id)
    if not include_inactive:
        stmt = stmt.where(MasOrgLevel.is_active.is_(True))
    return list(db.session.scalars(stmt.order_by(MasOrgLevel.depth)).all())


def _get_level(project_id: str, org_level_id: object) -> MasOrgLevel:
    level = db.session.get(MasOrgLevel, _uuid(org_level_id, what="Level id"))
    if level is None or level.project_id != project_id:
        raise OrganizationError("Level not found in this project.")
    return level


def create_level(
    project_id: str,
    *,
    level_code: object,
    level_name: object,
    depth: object,
    is_optional: bool = False,
) -> MasOrgLevel:
    _get_project(project_id)
    code = normalize_level_code(level_code)
    name = _clean_text(level_name, max_len=128, required=True, what="Level name")
    try:
        depth_value = int(depth)
    except (TypeError, ValueError) as exc:
        raise OrganizationError("Depth must be a positive integer.") from exc
    if depth_value < 1:
        raise OrganizationError("Depth must be a positive integer.")
    clash = db.session.scalar(
        sa.select(MasOrgLevel).where(
            MasOrgLevel.project_id == project_id,
            sa.or_(MasOrgLevel.level_code == code, MasOrgLevel.depth == depth_value),
        )
    )
    if clash is not None:
        raise OrganizationError(
            f"A level with code {clash.level_code!r} or depth {clash.depth} already exists."
        )
    level = MasOrgLevel(
        project_id=project_id,
        level_code=code,
        level_name=name,
        depth=depth_value,
        is_optional=bool(is_optional),
    )
    db.session.add(level)
    db.session.flush()
    return level


def update_level(project_id: str, org_level_id: object, **fields) -> MasOrgLevel:
    level = _get_level(project_id, org_level_id)
    if "level_name" in fields:
        level.level_name = _clean_text(fields["level_name"], max_len=128, required=True, what="Level name")
    if "is_optional" in fields:
        level.is_optional = bool(fields["is_optional"])
    if "level_code" in fields:
        new_code = normalize_level_code(fields["level_code"])
        if new_code != level.level_code:
            dup = db.session.scalar(
                sa.select(MasOrgLevel.org_level_id).where(
                    MasOrgLevel.project_id == project_id, MasOrgLevel.level_code == new_code
                )
            )
            if dup is not None:
                raise OrganizationError(f"Level code {new_code!r} already exists.")
            level.level_code = new_code
    if "depth" in fields and int(fields["depth"]) != level.depth:
        has_units = db.session.scalar(
            sa.select(sa.func.count()).select_from(MasOrgUnit).where(MasOrgUnit.org_level_id == level.org_level_id)
        )
        if has_units:
            raise OrganizationError("Depth cannot change while units exist at this level.")
        new_depth = int(fields["depth"])
        if new_depth < 1:
            raise OrganizationError("Depth must be a positive integer.")
        dup = db.session.scalar(
            sa.select(MasOrgLevel.org_level_id).where(
                MasOrgLevel.project_id == project_id, MasOrgLevel.depth == new_depth
            )
        )
        if dup is not None:
            raise OrganizationError(f"Depth {new_depth} is already used.")
        level.depth = new_depth
    if "is_active" in fields:
        active = bool(fields["is_active"])
        if not active:
            active_units = db.session.scalar(
                sa.select(sa.func.count()).select_from(MasOrgUnit).where(
                    MasOrgUnit.org_level_id == level.org_level_id, MasOrgUnit.is_active.is_(True)
                )
            )
            if active_units:
                raise OrganizationError("Deactivate or move the units at this level first.")
        level.is_active = active
    db.session.flush()
    return level


# ---------------------------------------------------------------------------
# Units
# ---------------------------------------------------------------------------


def list_units(project_id: str, *, include_inactive: bool = False) -> list[dict]:
    parent = _aliased_parent()
    stmt = (
        sa.select(MasOrgUnit, MasOrgLevel, parent.unit_code)
        .join(MasOrgLevel, MasOrgLevel.org_level_id == MasOrgUnit.org_level_id)
        .outerjoin(parent, parent.org_unit_id == MasOrgUnit.parent_org_unit_id)
        .where(MasOrgUnit.project_id == project_id)
    )
    if not include_inactive:
        stmt = stmt.where(MasOrgUnit.is_active.is_(True))
    rows = db.session.execute(stmt.order_by(MasOrgUnit.path)).all()
    return [serialize_unit(unit, level=level, parent_code=parent_code) for unit, level, parent_code in rows]


def resolve_org_unit_export_labels(org_unit_ids) -> dict[uuid.UUID, dict]:
    """Batched unit code/name/level-path lookup for exports.

    Takes any iterable of unit ids (``None`` entries are ignored) and returns
    ``{org_unit_id: {"unit_code", "unit_name", "level_path"}}`` for every id
    actually found — inactive units included, since history must stay
    readable. Exactly two queries regardless of how many rows the caller is
    exporting: one for the referenced units themselves, one for the ancestor
    codes named in their ``path`` (deduplicated across all of them), so a
    large export never turns into a per-row walk up the tree.
    """
    ids = {uid for uid in org_unit_ids if uid is not None}
    if not ids:
        return {}

    units = db.session.execute(
        sa.select(
            MasOrgUnit.org_unit_id,
            MasOrgUnit.project_id,
            MasOrgUnit.unit_code,
            MasOrgUnit.unit_name,
            MasOrgUnit.path,
        ).where(MasOrgUnit.org_unit_id.in_(ids))
    ).all()
    if not units:
        return {}

    project_ids = {row.project_id for row in units}
    ancestor_codes = {code for row in units for code in str(row.path).split(".")}

    name_by_project_code: dict[tuple[str, str], str] = {}
    if ancestor_codes:
        code_rows = db.session.execute(
            sa.select(MasOrgUnit.project_id, MasOrgUnit.unit_code, MasOrgUnit.unit_name).where(
                MasOrgUnit.project_id.in_(project_ids),
                MasOrgUnit.unit_code.in_(ancestor_codes),
            )
        ).all()
        name_by_project_code = {
            (row.project_id, row.unit_code): row.unit_name for row in code_rows
        }

    labels: dict[uuid.UUID, dict] = {}
    for row in units:
        codes = str(row.path).split(".")
        level_path = " / ".join(
            name_by_project_code.get((row.project_id, code), code) for code in codes
        )
        labels[row.org_unit_id] = {
            "unit_code": row.unit_code,
            "unit_name": row.unit_name,
            "level_path": level_path,
        }
    return labels


def _aliased_parent():
    from sqlalchemy.orm import aliased

    return aliased(MasOrgUnit)


def get_unit_tree(project_id: str, *, include_inactive: bool = False) -> list[dict]:
    """Nested tree (children under ``children``) in path order."""
    flat = list_units(project_id, include_inactive=include_inactive)
    by_id = {u["org_unit_id"]: {**u, "children": []} for u in flat}
    roots: list[dict] = []
    for node in by_id.values():
        parent_id = node["parent_org_unit_id"]
        if parent_id and parent_id in by_id:
            by_id[parent_id]["children"].append(node)
        else:
            roots.append(node)
    return roots


def _get_unit(project_id: str, org_unit_id: object) -> MasOrgUnit:
    unit = db.session.get(MasOrgUnit, _uuid(org_unit_id, what="Unit id"))
    if unit is None or unit.project_id != project_id:
        raise OrganizationError("Unit not found in this project.")
    return unit


def _validate_parent_for_level(project_id: str, level: MasOrgLevel, parent: MasOrgUnit | None) -> None:
    levels = list_levels(project_id, include_inactive=False)
    if not levels:
        raise OrganizationError("Define the project's levels before adding units.")
    top_depth = levels[0].depth
    if parent is None:
        if level.depth != top_depth:
            raise OrganizationError(
                f"Units at level {level.level_code!r} need a parent; only the top level "
                f"({levels[0].level_code!r}) may have none."
            )
        return
    if parent.project_id != project_id:
        raise OrganizationError("Parent unit belongs to another project.")
    parent_level = parent.level
    if parent_level.depth >= level.depth:
        raise OrganizationError(
            f"Parent {parent.unit_code!r} is at level {parent_level.level_code!r} (depth "
            f"{parent_level.depth}); a {level.level_code!r} unit needs a parent at a shallower level."
        )
    skipped = [lv for lv in levels if parent_level.depth < lv.depth < level.depth and not lv.is_optional]
    if skipped:
        names = ", ".join(repr(lv.level_code) for lv in skipped)
        raise OrganizationError(
            f"Level(s) {names} between {parent_level.level_code!r} and {level.level_code!r} are not "
            "optional; the parent must be at the next level up."
        )


def create_unit(
    project_id: str,
    *,
    org_level_id: object,
    unit_code: object,
    unit_name: object,
    parent_org_unit_id: object | None = None,
    address: object = None,
    phone: object = None,
    latitude: object = None,
    longitude: object = None,
    google_maps_url: object = None,
    remarks: object = None,
) -> MasOrgUnit:
    _get_project(project_id)
    level = _get_level(project_id, org_level_id)
    if not level.is_active:
        raise OrganizationError("Level is inactive.")
    code = normalize_code(unit_code, what="Unit code")
    name = _clean_text(unit_name, required=True, what="Unit name")
    parent = _get_unit(project_id, parent_org_unit_id) if parent_org_unit_id else None
    if parent is not None and not parent.is_active:
        raise OrganizationError("Parent unit is inactive.")
    _validate_parent_for_level(project_id, level, parent)
    dup = db.session.scalar(
        sa.select(MasOrgUnit.org_unit_id).where(
            MasOrgUnit.project_id == project_id, MasOrgUnit.unit_code == code
        )
    )
    if dup is not None:
        raise OrganizationError(f"Unit code {code!r} already exists in this project.")
    unit = MasOrgUnit(
        project_id=project_id,
        org_level_id=level.org_level_id,
        parent_org_unit_id=parent.org_unit_id if parent else None,
        unit_code=code,
        unit_name=name,
        path=f"{parent.path}.{code}" if parent else code,
        address=_clean_text(address, what="Address"),
        phone=_clean_text(phone, max_len=32, what="Phone"),
        latitude=_clean_coordinate(latitude, what="Latitude", limit=90),
        longitude=_clean_coordinate(longitude, what="Longitude", limit=180),
        google_maps_url=_clean_url(google_maps_url),
        remarks=_clean_text(remarks, what="Remarks"),
    )
    db.session.add(unit)
    db.session.flush()
    return unit


def _rewrite_subtree_paths(old_path: str, new_path: str) -> None:
    """Rewrite ``path`` for a unit and every descendant after a move or rename."""
    db.session.execute(
        sa.text(
            """
            UPDATE mas_org_unit
            SET path = CASE
                WHEN path = CAST(:old_path AS ltree) THEN CAST(:new_path AS ltree)
                ELSE CAST(:new_path AS ltree) || subpath(path, nlevel(CAST(:old_path AS ltree)))
            END
            WHERE path <@ CAST(:old_path AS ltree)
            """
        ),
        {"old_path": old_path, "new_path": new_path},
    )
    db.session.expire_all()


def update_unit(project_id: str, org_unit_id: object, **fields) -> MasOrgUnit:
    unit = _get_unit(project_id, org_unit_id)
    old_path = str(unit.path)
    level = unit.level
    parent = unit.parent

    if "org_level_id" in fields and fields["org_level_id"]:
        level = _get_level(project_id, fields["org_level_id"])
    if "parent_org_unit_id" in fields:
        raw_parent = fields["parent_org_unit_id"]
        parent = _get_unit(project_id, raw_parent) if raw_parent else None
        if parent is not None:
            if parent.org_unit_id == unit.org_unit_id or str(parent.path).startswith(old_path + "."):
                raise OrganizationError("A unit cannot be moved under itself or its own descendants.")
            if not parent.is_active:
                raise OrganizationError("Parent unit is inactive.")
    level_or_parent_changed = level.org_level_id != unit.org_level_id or (
        (parent.org_unit_id if parent else None) != unit.parent_org_unit_id
    )
    if level_or_parent_changed:
        _validate_parent_for_level(project_id, level, parent)
        if level.org_level_id != unit.org_level_id:
            # Children must still sit below the new level.
            child_depth = db.session.scalar(
                sa.select(sa.func.min(MasOrgLevel.depth))
                .select_from(MasOrgUnit)
                .join(MasOrgLevel, MasOrgLevel.org_level_id == MasOrgUnit.org_level_id)
                .where(MasOrgUnit.parent_org_unit_id == unit.org_unit_id)
            )
            if child_depth is not None and child_depth <= level.depth:
                raise OrganizationError("Children of this unit would end up at or above its new level.")
        unit.org_level_id = level.org_level_id
        unit.parent_org_unit_id = parent.org_unit_id if parent else None

    new_code = unit.unit_code
    if "unit_code" in fields and fields["unit_code"] is not None:
        new_code = normalize_code(fields["unit_code"], what="Unit code")
        if new_code != unit.unit_code:
            dup = db.session.scalar(
                sa.select(MasOrgUnit.org_unit_id).where(
                    MasOrgUnit.project_id == project_id, MasOrgUnit.unit_code == new_code
                )
            )
            if dup is not None:
                raise OrganizationError(f"Unit code {new_code!r} already exists in this project.")
            unit.unit_code = new_code

    if "unit_name" in fields:
        unit.unit_name = _clean_text(fields["unit_name"], required=True, what="Unit name")
    if "address" in fields:
        unit.address = _clean_text(fields["address"], what="Address")
    if "phone" in fields:
        unit.phone = _clean_text(fields["phone"], max_len=32, what="Phone")
    if "latitude" in fields:
        unit.latitude = _clean_coordinate(fields["latitude"], what="Latitude", limit=90)
    if "longitude" in fields:
        unit.longitude = _clean_coordinate(fields["longitude"], what="Longitude", limit=180)
    if "google_maps_url" in fields:
        unit.google_maps_url = _clean_url(fields["google_maps_url"])
    if "remarks" in fields:
        unit.remarks = _clean_text(fields["remarks"], what="Remarks")

    db.session.flush()
    new_path = f"{parent.path}.{new_code}" if parent else new_code
    if new_path != old_path:
        _rewrite_subtree_paths(old_path, new_path)
        unit = db.session.get(MasOrgUnit, unit.org_unit_id)

    if "is_active" in fields:
        set_unit_active(project_id, unit.org_unit_id, bool(fields["is_active"]))
        unit = db.session.get(MasOrgUnit, unit.org_unit_id)
    return unit


def set_unit_active(project_id: str, org_unit_id: object, active: bool) -> int:
    """Deactivate a unit with its whole subtree, or reactivate a single unit.

    Returns the number of units changed.
    """
    unit = _get_unit(project_id, org_unit_id)
    if active:
        if unit.parent is not None and not unit.parent.is_active:
            raise OrganizationError("Reactivate the parent unit first.")
        if unit.is_active:
            return 0
        unit.is_active = True
        db.session.flush()
        return 1
    result = db.session.execute(
        sa.text(
            """
            UPDATE mas_org_unit
            SET is_active = false, updated_at = now()
            WHERE project_id = :project_id
              AND path <@ CAST(:path AS ltree)
              AND is_active = true
            """
        ),
        {"project_id": project_id, "path": str(unit.path)},
    )
    db.session.expire_all()
    return result.rowcount or 0


def find_unit_by_code(project_id: str, unit_code: object) -> MasOrgUnit | None:
    try:
        code = normalize_code(unit_code, what="Unit code")
    except OrganizationError:
        return None
    return db.session.scalar(
        sa.select(MasOrgUnit).where(MasOrgUnit.project_id == project_id, MasOrgUnit.unit_code == code)
    )


def subtree_unit_ids(unit: MasOrgUnit, *, include_inactive: bool = False) -> list[uuid.UUID]:
    stmt = sa.select(MasOrgUnit.org_unit_id).where(
        MasOrgUnit.project_id == unit.project_id,
        sa.text("path <@ CAST(:ancestor AS ltree)").bindparams(ancestor=str(unit.path)),
    )
    if not include_inactive:
        stmt = stmt.where(MasOrgUnit.is_active.is_(True))
    return list(db.session.scalars(stmt).all())


# ---------------------------------------------------------------------------
# Cadres and level permissions
# ---------------------------------------------------------------------------


def list_cadres(project_id: str, *, include_inactive: bool = False) -> list[MasCadre]:
    stmt = sa.select(MasCadre).where(MasCadre.project_id == project_id)
    if not include_inactive:
        stmt = stmt.where(MasCadre.is_active.is_(True))
    return list(db.session.scalars(stmt.order_by(MasCadre.cadre_code)).all())


def _get_cadre(project_id: str, cadre_id: object) -> MasCadre:
    cadre = db.session.get(MasCadre, _uuid(cadre_id, what="Cadre id"))
    if cadre is None or cadre.project_id != project_id:
        raise OrganizationError("Cadre not found in this project.")
    return cadre


def create_cadre(project_id: str, *, cadre_code: object, cadre_name: object) -> MasCadre:
    _get_project(project_id)
    code = normalize_code(cadre_code, what="Cadre code")
    name = _clean_text(cadre_name, max_len=128, required=True, what="Cadre name")
    dup = db.session.scalar(
        sa.select(MasCadre.cadre_id).where(MasCadre.project_id == project_id, MasCadre.cadre_code == code)
    )
    if dup is not None:
        raise OrganizationError(f"Cadre code {code!r} already exists in this project.")
    cadre = MasCadre(project_id=project_id, cadre_code=code, cadre_name=name)
    db.session.add(cadre)
    db.session.flush()
    return cadre


def update_cadre(project_id: str, cadre_id: object, **fields) -> MasCadre:
    cadre = _get_cadre(project_id, cadre_id)
    if "cadre_name" in fields:
        cadre.cadre_name = _clean_text(fields["cadre_name"], max_len=128, required=True, what="Cadre name")
    if "cadre_code" in fields and fields["cadre_code"] is not None:
        new_code = normalize_code(fields["cadre_code"], what="Cadre code")
        if new_code != cadre.cadre_code:
            dup = db.session.scalar(
                sa.select(MasCadre.cadre_id).where(
                    MasCadre.project_id == project_id, MasCadre.cadre_code == new_code
                )
            )
            if dup is not None:
                raise OrganizationError(f"Cadre code {new_code!r} already exists in this project.")
            cadre.cadre_code = new_code
    if "is_active" in fields:
        cadre.is_active = bool(fields["is_active"])
    db.session.flush()
    return cadre


def list_level_cadres(project_id: str) -> list[dict]:
    rows = db.session.execute(
        sa.select(MapOrgLevelCadre, MasOrgLevel, MasCadre)
        .join(MasOrgLevel, MasOrgLevel.org_level_id == MapOrgLevelCadre.org_level_id)
        .join(MasCadre, MasCadre.cadre_id == MapOrgLevelCadre.cadre_id)
        .where(MasOrgLevel.project_id == project_id)
        .order_by(MasOrgLevel.depth, MasCadre.cadre_code)
    ).all()
    return [serialize_level_cadre(row, level=level, cadre=cadre) for row, level, cadre in rows]


def upsert_level_cadre(
    project_id: str,
    *,
    org_level_id: object,
    cadre_id: object,
    can_fill_va_form: bool,
    can_code_va_form: bool,
    is_active: bool = True,
) -> MapOrgLevelCadre:
    level = _get_level(project_id, org_level_id)
    cadre = _get_cadre(project_id, cadre_id)
    row = db.session.scalar(
        sa.select(MapOrgLevelCadre).where(
            MapOrgLevelCadre.org_level_id == level.org_level_id,
            MapOrgLevelCadre.cadre_id == cadre.cadre_id,
        )
    )
    if row is None:
        row = MapOrgLevelCadre(org_level_id=level.org_level_id, cadre_id=cadre.cadre_id)
        db.session.add(row)
    row.can_fill_va_form = bool(can_fill_va_form)
    row.can_code_va_form = bool(can_code_va_form)
    row.is_active = bool(is_active)
    db.session.flush()
    return row


def get_level_cadre_permission(org_level_id: uuid.UUID, cadre_id: uuid.UUID) -> MapOrgLevelCadre | None:
    return db.session.scalar(
        sa.select(MapOrgLevelCadre).where(
            MapOrgLevelCadre.org_level_id == org_level_id,
            MapOrgLevelCadre.cadre_id == cadre_id,
            MapOrgLevelCadre.is_active.is_(True),
        )
    )


# ---------------------------------------------------------------------------
# Workers
# ---------------------------------------------------------------------------


def list_workers(project_id: str, *, include_inactive: bool = False, org_unit_id: object | None = None) -> list[dict]:
    stmt = (
        sa.select(MasOrgUnitWorker, MasOrgUnit, MasCadre, VaUsers.email)
        .join(MasOrgUnit, MasOrgUnit.org_unit_id == MasOrgUnitWorker.org_unit_id)
        .join(MasCadre, MasCadre.cadre_id == MasOrgUnitWorker.cadre_id)
        .outerjoin(VaUsers, VaUsers.user_id == MasOrgUnitWorker.user_id)
        .where(MasOrgUnitWorker.project_id == project_id)
    )
    if org_unit_id:
        stmt = stmt.where(MasOrgUnitWorker.org_unit_id == _uuid(org_unit_id, what="Unit id"))
    if not include_inactive:
        stmt = stmt.where(MasOrgUnitWorker.is_active.is_(True))
    rows = db.session.execute(stmt.order_by(MasOrgUnit.path, MasOrgUnitWorker.worker_code)).all()
    return [
        serialize_worker(worker, unit=unit, cadre=cadre, user_email=email)
        for worker, unit, cadre, email in rows
    ]


def _get_worker(project_id: str, worker_id: object) -> MasOrgUnitWorker:
    worker = db.session.get(MasOrgUnitWorker, _uuid(worker_id, what="Worker id"))
    if worker is None or worker.project_id != project_id:
        raise OrganizationError("Worker not found in this project.")
    return worker


def _require_cadre_at_unit_level(unit: MasOrgUnit, cadre: MasCadre) -> None:
    if get_level_cadre_permission(unit.org_level_id, cadre.cadre_id) is None:
        raise OrganizationError(
            f"Cadre {cadre.cadre_code!r} is not defined at level {unit.level.level_code!r}; "
            "add it in the level permissions grid first."
        )


def _resolve_user(raw: object) -> VaUsers | None:
    value = str(raw or "").strip()
    if not value:
        return None
    user = None
    try:
        user = db.session.get(VaUsers, uuid.UUID(value))
    except ValueError:
        user = db.session.scalar(sa.select(VaUsers).where(sa.func.lower(VaUsers.email) == value.lower()))
    if user is None:
        raise OrganizationError(f"DigitVA user not found: {value}")
    return user


def create_worker(
    project_id: str,
    *,
    org_unit_id: object,
    cadre_id: object,
    worker_code: object,
    worker_name: object,
    phone: object = None,
    user: object = None,
    remarks: object = None,
) -> MasOrgUnitWorker:
    _get_project(project_id)
    unit = _get_unit(project_id, org_unit_id)
    cadre = _get_cadre(project_id, cadre_id)
    _require_cadre_at_unit_level(unit, cadre)
    code = normalize_code(worker_code, what="Worker code")
    dup = db.session.scalar(
        sa.select(MasOrgUnitWorker.worker_id).where(
            MasOrgUnitWorker.project_id == project_id, MasOrgUnitWorker.worker_code == code
        )
    )
    if dup is not None:
        raise OrganizationError(f"Worker code {code!r} already exists in this project.")
    resolved_user = _resolve_user(user)
    worker = MasOrgUnitWorker(
        project_id=project_id,
        org_unit_id=unit.org_unit_id,
        cadre_id=cadre.cadre_id,
        worker_code=code,
        worker_name=_clean_text(worker_name, required=True, what="Worker name"),
        phone=_clean_text(phone, max_len=32, what="Phone"),
        user_id=resolved_user.user_id if resolved_user else None,
        remarks=_clean_text(remarks, what="Remarks"),
    )
    db.session.add(worker)
    db.session.flush()
    return worker


def update_worker(project_id: str, worker_id: object, **fields) -> MasOrgUnitWorker:
    worker = _get_worker(project_id, worker_id)
    unit = worker.unit
    cadre = worker.cadre
    if "org_unit_id" in fields and fields["org_unit_id"]:
        unit = _get_unit(project_id, fields["org_unit_id"])
    if "cadre_id" in fields and fields["cadre_id"]:
        cadre = _get_cadre(project_id, fields["cadre_id"])
    if unit.org_unit_id != worker.org_unit_id or cadre.cadre_id != worker.cadre_id:
        _require_cadre_at_unit_level(unit, cadre)
        worker.org_unit_id = unit.org_unit_id
        worker.cadre_id = cadre.cadre_id
    if "worker_code" in fields and fields["worker_code"] is not None:
        new_code = normalize_code(fields["worker_code"], what="Worker code")
        if new_code != worker.worker_code:
            dup = db.session.scalar(
                sa.select(MasOrgUnitWorker.worker_id).where(
                    MasOrgUnitWorker.project_id == project_id, MasOrgUnitWorker.worker_code == new_code
                )
            )
            if dup is not None:
                raise OrganizationError(f"Worker code {new_code!r} already exists in this project.")
            worker.worker_code = new_code
    if "worker_name" in fields:
        worker.worker_name = _clean_text(fields["worker_name"], required=True, what="Worker name")
    if "phone" in fields:
        worker.phone = _clean_text(fields["phone"], max_len=32, what="Phone")
    if "user" in fields:
        resolved = _resolve_user(fields["user"])
        worker.user_id = resolved.user_id if resolved else None
    if "remarks" in fields:
        worker.remarks = _clean_text(fields["remarks"], what="Remarks")
    if "is_active" in fields:
        worker.is_active = bool(fields["is_active"])
    db.session.flush()
    return worker


# ---------------------------------------------------------------------------
# Seed template
# ---------------------------------------------------------------------------


def seed_default_organization(project_id: str, *, include_cadres: bool = True) -> dict[str, int]:
    """Idempotently add the template levels, cadres and permissions to a project.

    Existing codes are left untouched, so this is safe to rerun.
    """
    _get_project(project_id)
    counts = {"levels": 0, "cadres": 0, "level_cadres": 0}
    existing_levels = {lv.level_code: lv for lv in list_levels(project_id, include_inactive=True)}
    used_depths = {lv.depth for lv in existing_levels.values()}
    for code, name, depth, optional in DEFAULT_LEVEL_TEMPLATE:
        if code in existing_levels or depth in used_depths:
            continue
        existing_levels[code] = create_level(
            project_id, level_code=code, level_name=name, depth=depth, is_optional=optional
        )
        used_depths.add(depth)
        counts["levels"] += 1
    if not include_cadres:
        return counts
    existing_cadres = {c.cadre_code: c for c in list_cadres(project_id, include_inactive=True)}
    for code, name in DEFAULT_CADRE_TEMPLATE:
        if code in existing_cadres:
            continue
        existing_cadres[code] = create_cadre(project_id, cadre_code=code, cadre_name=name)
        counts["cadres"] += 1
    for (level_code, cadre_code), (can_fill, can_code) in DEFAULT_LEVEL_CADRE_TEMPLATE.items():
        level = existing_levels.get(level_code)
        cadre = existing_cadres.get(cadre_code)
        if level is None or cadre is None:
            continue
        exists = db.session.scalar(
            sa.select(MapOrgLevelCadre.level_cadre_id).where(
                MapOrgLevelCadre.org_level_id == level.org_level_id,
                MapOrgLevelCadre.cadre_id == cadre.cadre_id,
            )
        )
        if exists is not None:
            continue
        upsert_level_cadre(
            project_id,
            org_level_id=level.org_level_id,
            cadre_id=cadre.cadre_id,
            can_fill_va_form=can_fill,
            can_code_va_form=can_code,
        )
        counts["level_cadres"] += 1
    db.session.flush()
    return counts


# ---------------------------------------------------------------------------
# Export
# ---------------------------------------------------------------------------

_LEVEL_HEADERS = ("level_code", "level_name", "depth", "is_optional", "is_active", "odk_field_name")
_UNIT_HEADERS = (
    "unit_code", "unit_name", "level_code", "parent_code", "path", "address", "phone",
    "latitude", "longitude", "google_maps_url", "remarks", "is_active",
)
_CADRE_HEADERS = ("cadre_code", "cadre_name", "is_active")
_LEVEL_CADRE_HEADERS = ("level_code", "cadre_code", "can_fill_va_form", "can_code_va_form", "is_active")
_WORKER_HEADERS = (
    "worker_code", "worker_name", "unit_code", "cadre_code", "phone", "user_email", "remarks", "is_active",
)


def export_organization_rows(project_id: str) -> dict[str, list[dict]]:
    """All master data of one project as plain rows keyed by sheet name."""
    _get_project(project_id)
    levels = [serialize_level(lv) for lv in list_levels(project_id, include_inactive=True)]
    units = list_units(project_id, include_inactive=True)
    cadres = [serialize_cadre(c) for c in list_cadres(project_id, include_inactive=True)]
    level_cadres = list_level_cadres(project_id)
    workers = list_workers(project_id, include_inactive=True)
    return {
        "levels": [{k: lv[k] for k in _LEVEL_HEADERS} for lv in levels],
        "units": [{k: u[k] for k in _UNIT_HEADERS} for u in units],
        "cadres": [{k: c[k] for k in _CADRE_HEADERS} for c in cadres],
        "level_cadres": [{k: lc[k] for k in _LEVEL_CADRE_HEADERS} for lc in level_cadres],
        "workers": [{k: w[k] for k in _WORKER_HEADERS} for w in workers],
    }


def export_organization_xlsx(project_id: str) -> bytes:
    rows = export_organization_rows(project_id)
    headers = {
        "levels": _LEVEL_HEADERS,
        "units": _UNIT_HEADERS,
        "cadres": _CADRE_HEADERS,
        "level_cadres": _LEVEL_CADRE_HEADERS,
        "workers": _WORKER_HEADERS,
    }
    workbook = Workbook()
    first = True
    for sheet_name in EXPORT_SHEETS:
        sheet = workbook.active if first else workbook.create_sheet()
        first = False
        sheet.title = sheet_name
        sheet.append(list(headers[sheet_name]))
        for row in rows[sheet_name]:
            sheet.append([row.get(col) for col in headers[sheet_name]])
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def export_organization_csv(project_id: str, sheet: str) -> str:
    if sheet not in EXPORT_SHEETS:
        raise OrganizationError(f"Unknown sheet {sheet!r}; choose one of {', '.join(EXPORT_SHEETS)}.")
    rows = export_organization_rows(project_id)[sheet]
    headers = {
        "levels": _LEVEL_HEADERS,
        "units": _UNIT_HEADERS,
        "cadres": _CADRE_HEADERS,
        "level_cadres": _LEVEL_CADRE_HEADERS,
        "workers": _WORKER_HEADERS,
    }[sheet]
    buffer = io.StringIO()
    writer = csv.DictWriter(buffer, fieldnames=list(headers))
    writer.writeheader()
    for row in rows:
        writer.writerow({k: ("" if row.get(k) is None else row.get(k)) for k in headers})
    return buffer.getvalue()


def export_odk_choices_rows(project_id: str) -> list[dict]:
    """XLSForm ``choices`` rows: one list per level, filtered by parent code.

    ``list_name`` is ``org_<level_code>``; ``name`` is the unit code; ``label``
    the unit name; ``parent_code`` is the filter column for cascading selects.
    The matching survey fields are ``org_<level_code>_code``.
    """
    levels = {lv.org_level_id: lv for lv in list_levels(project_id)}
    rows: list[dict] = []
    for unit in list_units(project_id):
        level = levels.get(uuid.UUID(unit["org_level_id"]))
        if level is None:
            continue
        rows.append(
            {
                "list_name": f"org_{level.level_code}",
                "name": unit["unit_code"],
                "label": unit["unit_name"],
                "parent_code": unit["parent_code"] or "",
            }
        )
    return rows


def export_odk_choices_csv(project_id: str) -> str:
    buffer = io.StringIO()
    writer = csv.DictWriter(buffer, fieldnames=["list_name", "name", "label", "parent_code"])
    writer.writeheader()
    for row in export_odk_choices_rows(project_id):
        writer.writerow(row)
    return buffer.getvalue()


# ---------------------------------------------------------------------------
# Import (dry run first, then apply)
# ---------------------------------------------------------------------------


@dataclass
class ImportPlan:
    sheets: dict[str, list[dict]] = field(default_factory=dict)
    creates: dict[str, list[str]] = field(default_factory=dict)
    updates: dict[str, list[str]] = field(default_factory=dict)
    deactivates: dict[str, list[str]] = field(default_factory=dict)
    errors: list[str] = field(default_factory=list)
    applied: bool = False

    def as_dict(self) -> dict:
        return {
            "creates": self.creates,
            "updates": self.updates,
            "deactivates": self.deactivates,
            "errors": self.errors,
            "applied": self.applied,
            "counts": {
                sheet: {
                    "rows": len(self.sheets.get(sheet, [])),
                    "create": len(self.creates.get(sheet, [])),
                    "update": len(self.updates.get(sheet, [])),
                    "deactivate": len(self.deactivates.get(sheet, [])),
                }
                for sheet in EXPORT_SHEETS
            },
        }


def parse_organization_workbook(file_obj) -> dict[str, list[dict]]:
    """Read the export workbook back into rows keyed by sheet; missing sheets are skipped."""
    workbook = load_workbook(file_obj, read_only=True, data_only=True)
    sheets: dict[str, list[dict]] = {}
    for sheet_name in EXPORT_SHEETS:
        if sheet_name not in workbook.sheetnames:
            continue
        ws = workbook[sheet_name]
        rows_iter = ws.iter_rows(values_only=True)
        try:
            header = [str(h).strip() if h is not None else "" for h in next(rows_iter)]
        except StopIteration:
            sheets[sheet_name] = []
            continue
        rows = []
        for raw in rows_iter:
            if raw is None or all(v is None or str(v).strip() == "" for v in raw):
                continue
            rows.append({header[i]: raw[i] for i in range(min(len(header), len(raw))) if header[i]})
        sheets[sheet_name] = rows
    workbook.close()
    return sheets


def import_organization(project_id: str, sheets: dict[str, list[dict]], *, dry_run: bool = True, deactivate_missing: bool = False) -> ImportPlan:
    """Upsert master data by business codes. Never deletes.

    With ``deactivate_missing`` rows absent from a supplied sheet are
    deactivated. Runs inside a SAVEPOINT so a dry run or a failed import
    leaves nothing behind.
    """
    _get_project(project_id)
    plan = ImportPlan(sheets=sheets)
    for sheet in EXPORT_SHEETS:
        plan.creates[sheet] = []
        plan.updates[sheet] = []
        plan.deactivates[sheet] = []

    savepoint = db.session.begin_nested()
    try:
        _import_levels(project_id, sheets.get("levels"), plan, deactivate_missing)
        _import_cadres(project_id, sheets.get("cadres"), plan, deactivate_missing)
        _import_level_cadres(project_id, sheets.get("level_cadres"), plan, deactivate_missing)
        _import_units(project_id, sheets.get("units"), plan, deactivate_missing)
        _import_workers(project_id, sheets.get("workers"), plan, deactivate_missing)
    except OrganizationError as exc:
        plan.errors.append(str(exc))
    if plan.errors or dry_run:
        savepoint.rollback()
        return plan
    savepoint.commit()
    plan.applied = True
    return plan


def _import_levels(project_id, rows, plan, deactivate_missing):
    if rows is None:
        return
    existing = {lv.level_code: lv for lv in list_levels(project_id, include_inactive=True)}
    seen = set()
    for i, row in enumerate(rows, start=2):
        try:
            code = normalize_level_code(row.get("level_code"))
            seen.add(code)
            if code in existing:
                lv = existing[code]
                update_level(
                    project_id,
                    lv.org_level_id,
                    level_name=row.get("level_name") or lv.level_name,
                    is_optional=_to_bool(row.get("is_optional")),
                    is_active=_to_bool(row.get("is_active", True)) if row.get("is_active") is not None else lv.is_active,
                    **({"depth": row.get("depth")} if row.get("depth") not in (None, "") else {}),
                )
                plan.updates["levels"].append(code)
            else:
                existing[code] = create_level(
                    project_id,
                    level_code=code,
                    level_name=row.get("level_name"),
                    depth=row.get("depth"),
                    is_optional=_to_bool(row.get("is_optional")),
                )
                plan.creates["levels"].append(code)
        except OrganizationError as exc:
            raise OrganizationError(f"levels row {i}: {exc}") from exc
    if deactivate_missing:
        for code, lv in existing.items():
            if code not in seen and lv.is_active:
                update_level(project_id, lv.org_level_id, is_active=False)
                plan.deactivates["levels"].append(code)


def _import_cadres(project_id, rows, plan, deactivate_missing):
    if rows is None:
        return
    existing = {c.cadre_code: c for c in list_cadres(project_id, include_inactive=True)}
    seen = set()
    for i, row in enumerate(rows, start=2):
        try:
            code = normalize_code(row.get("cadre_code"), what="Cadre code")
            seen.add(code)
            if code in existing:
                c = existing[code]
                update_cadre(
                    project_id,
                    c.cadre_id,
                    cadre_name=row.get("cadre_name") or c.cadre_name,
                    is_active=_to_bool(row.get("is_active", True)) if row.get("is_active") is not None else c.is_active,
                )
                plan.updates["cadres"].append(code)
            else:
                existing[code] = create_cadre(project_id, cadre_code=code, cadre_name=row.get("cadre_name"))
                plan.creates["cadres"].append(code)
        except OrganizationError as exc:
            raise OrganizationError(f"cadres row {i}: {exc}") from exc
    if deactivate_missing:
        for code, c in existing.items():
            if code not in seen and c.is_active:
                update_cadre(project_id, c.cadre_id, is_active=False)
                plan.deactivates["cadres"].append(code)


def _import_level_cadres(project_id, rows, plan, deactivate_missing):
    if rows is None:
        return
    levels = {lv.level_code: lv for lv in list_levels(project_id, include_inactive=True)}
    cadres = {c.cadre_code: c for c in list_cadres(project_id, include_inactive=True)}
    existing = {(lc["level_code"], lc["cadre_code"]): lc for lc in list_level_cadres(project_id)}
    seen = set()
    for i, row in enumerate(rows, start=2):
        try:
            level_code = normalize_level_code(row.get("level_code"))
            cadre_code = normalize_code(row.get("cadre_code"), what="Cadre code")
            level = levels.get(level_code)
            cadre = cadres.get(cadre_code)
            if level is None:
                raise OrganizationError(f"unknown level {level_code!r}")
            if cadre is None:
                raise OrganizationError(f"unknown cadre {cadre_code!r}")
            key = (level_code, cadre_code)
            seen.add(key)
            upsert_level_cadre(
                project_id,
                org_level_id=level.org_level_id,
                cadre_id=cadre.cadre_id,
                can_fill_va_form=_to_bool(row.get("can_fill_va_form")),
                can_code_va_form=_to_bool(row.get("can_code_va_form")),
                is_active=_to_bool(row.get("is_active", True)) if row.get("is_active") is not None else True,
            )
            (plan.updates if key in existing else plan.creates)["level_cadres"].append(f"{level_code}/{cadre_code}")
        except OrganizationError as exc:
            raise OrganizationError(f"level_cadres row {i}: {exc}") from exc
    if deactivate_missing:
        for key, lc in existing.items():
            if key not in seen and lc["is_active"]:
                upsert_level_cadre(
                    project_id,
                    org_level_id=lc["org_level_id"],
                    cadre_id=lc["cadre_id"],
                    can_fill_va_form=lc["can_fill_va_form"],
                    can_code_va_form=lc["can_code_va_form"],
                    is_active=False,
                )
                plan.deactivates["level_cadres"].append(f"{key[0]}/{key[1]}")


def _import_units(project_id, rows, plan, deactivate_missing):
    if rows is None:
        return
    levels = {lv.level_code: lv for lv in list_levels(project_id, include_inactive=True)}
    existing = {u["unit_code"]: u for u in list_units(project_id, include_inactive=True)}
    seen = set()
    # Parents must exist before children: sort by level depth, keep file order otherwise.
    indexed = []
    for i, row in enumerate(rows, start=2):
        level_code = str(row.get("level_code") or "").strip().lower()
        depth = levels[level_code].depth if level_code in levels else 10**6
        indexed.append((depth, i, row))
    indexed.sort(key=lambda t: (t[0], t[1]))
    for _, i, row in indexed:
        try:
            code = normalize_code(row.get("unit_code"), what="Unit code")
            level = levels.get(normalize_level_code(row.get("level_code")))
            if level is None:
                raise OrganizationError(f"unknown level {row.get('level_code')!r}")
            parent_code = str(row.get("parent_code") or "").strip().upper()
            parent_id = existing[parent_code]["org_unit_id"] if parent_code in existing else None
            if parent_code and parent_id is None:
                raise OrganizationError(f"unknown parent unit {parent_code!r}")
            seen.add(code)
            fields = dict(
                unit_name=row.get("unit_name"),
                address=row.get("address"),
                phone=row.get("phone"),
                latitude=row.get("latitude"),
                longitude=row.get("longitude"),
                google_maps_url=row.get("google_maps_url"),
                remarks=row.get("remarks"),
            )
            if code in existing:
                update_unit(
                    project_id,
                    existing[code]["org_unit_id"],
                    org_level_id=level.org_level_id,
                    parent_org_unit_id=parent_id,
                    is_active=_to_bool(row.get("is_active", True)) if row.get("is_active") is not None else existing[code]["is_active"],
                    **fields,
                )
                plan.updates["units"].append(code)
            else:
                unit = create_unit(
                    project_id,
                    org_level_id=level.org_level_id,
                    unit_code=code,
                    parent_org_unit_id=parent_id,
                    **fields,
                )
                existing[code] = serialize_unit(unit, level=level, parent_code=parent_code or None)
                plan.creates["units"].append(code)
        except OrganizationError as exc:
            raise OrganizationError(f"units row {i}: {exc}") from exc
    if deactivate_missing:
        for code, u in existing.items():
            if code not in seen and u["is_active"]:
                current = db.session.get(MasOrgUnit, uuid.UUID(u["org_unit_id"]))
                if current is not None and current.is_active:
                    set_unit_active(project_id, current.org_unit_id, False)
                    plan.deactivates["units"].append(code)


def _import_workers(project_id, rows, plan, deactivate_missing):
    if rows is None:
        return
    cadres = {c.cadre_code: c for c in list_cadres(project_id, include_inactive=True)}
    units_by_code = {u["unit_code"]: u["org_unit_id"] for u in list_units(project_id, include_inactive=True)}
    existing = {w["worker_code"]: w for w in list_workers(project_id, include_inactive=True)}
    seen = set()
    for i, row in enumerate(rows, start=2):
        try:
            code = normalize_code(row.get("worker_code"), what="Worker code")
            unit_id = units_by_code.get(str(row.get("unit_code") or "").strip().upper())
            if unit_id is None:
                raise OrganizationError(f"unknown unit {row.get('unit_code')!r}")
            cadre = cadres.get(normalize_code(row.get("cadre_code"), what="Cadre code"))
            if cadre is None:
                raise OrganizationError(f"unknown cadre {row.get('cadre_code')!r}")
            seen.add(code)
            fields = dict(
                worker_name=row.get("worker_name"),
                phone=row.get("phone"),
                user=row.get("user_email"),
                remarks=row.get("remarks"),
            )
            if code in existing:
                update_worker(
                    project_id,
                    existing[code]["worker_id"],
                    org_unit_id=unit_id,
                    cadre_id=cadre.cadre_id,
                    is_active=_to_bool(row.get("is_active", True)) if row.get("is_active") is not None else existing[code]["is_active"],
                    **fields,
                )
                plan.updates["workers"].append(code)
            else:
                create_worker(
                    project_id,
                    org_unit_id=unit_id,
                    cadre_id=cadre.cadre_id,
                    worker_code=code,
                    **fields,
                )
                plan.creates["workers"].append(code)
        except OrganizationError as exc:
            raise OrganizationError(f"workers row {i}: {exc}") from exc
    if deactivate_missing:
        for code, w in existing.items():
            if code not in seen and w["is_active"]:
                update_worker(project_id, w["worker_id"], is_active=False)
                plan.deactivates["workers"].append(code)
