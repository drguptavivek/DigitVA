"""Validate and apply project user grants from a bounded table upload."""

import csv
import io
import logging
import re
import secrets
from types import SimpleNamespace

import sqlalchemy as sa
from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation

from app import db
from app.models import (
    MapOrgLevelCadre,
    MasCadre,
    MasOrgLevel,
    MasOrgUnit,
    VaAccessRoles,
    VaAccessScopeTypes,
    VaProjectMaster,
    VaStatuses,
    VaUserAccessGrants,
    VaUsers,
)
from app.models.mas_languages import MasLanguages
from app.models.va_selectives import ROLE_LABELS
from app.services import authz
from app.services.mentor_institute_service import check_mentor_grant, member_user_ids
from app.services.org_grant_service import CADRE_FLAG_BY_ROLE, ROLES_ALLOWING_ORG_UNIT
from app.services.organization_service import _spreadsheet_safe
from app.services.tabular_import_service import TabularImportError, parse_table
from app.services.user_account_service import (
    PHONE_CANONICAL,
    canonical_mobile,
    mask_mobile,
    record_account_created,
)

HEADERS = ("email", "name", "role", "org_unit_code", "cadre_code", "language_codes", "phone")
MAX_BYTES = 1024 * 1024
MAX_ROWS = 1000
EMAIL_RE = re.compile(r"^[^\s@,;]+@[^\s@,;]+\.[^\s@,;]+$")
NOT_PERMITTED = "you are not permitted to grant this role here"
log = logging.getLogger(__name__)


class ProjectUserImportError(ValueError):
    """A table or one of its rows cannot be safely imported."""


def template_csv():
    """Return the blank, header-only CSV accepted by this importer."""
    output = io.StringIO()
    csv.writer(output).writerow(HEADERS)
    return output.getvalue()


def _defined_name(workbook, name, sheet, start_row, end_row, column):
    """Add a workbook name for a reference column, including empty lists."""
    workbook.defined_names.add(DefinedName(
        name=name,
        attr_text=f"'{sheet.title}'!${get_column_letter(column)}${start_row}:"
                  f"${get_column_letter(column)}${max(start_row, end_row)}",
    ))


def _grantable_role_values(actor):
    """Return the role values accepted by the importer for this route actor."""
    is_admin = bool(actor is not None and actor.is_admin())
    return [
        role.value for role in VaAccessRoles
        if role != VaAccessRoles.admin and (is_admin or role != VaAccessRoles.project_pi)
    ]


def template_xlsx(project_id, actor=None):
    """Return a project-aware, blank XLSX template with safe dependent lists.

    The first sheet remains a header-only import surface. Reference data is
    read in bounded batch queries and is hidden from ordinary users; the
    server importer remains authoritative when a workbook is uploaded.
    """
    levels = list(db.session.scalars(sa.select(MasOrgLevel).where(
        MasOrgLevel.project_id == project_id,
        MasOrgLevel.is_active.is_(True),
    ).order_by(MasOrgLevel.depth)).all())
    level_keys = {level.org_level_id: f"L{index}" for index, level in enumerate(levels, 1)}
    units = list(db.session.scalars(sa.select(MasOrgUnit).where(
        MasOrgUnit.project_id == project_id,
        MasOrgUnit.is_active.is_(True),
        MasOrgUnit.org_level_id.in_(list(level_keys)),
    ).order_by(MasOrgUnit.unit_code)).all()) if level_keys else []
    cadres = list(db.session.scalars(sa.select(MasCadre).where(
        MasCadre.project_id == project_id,
        MasCadre.is_active.is_(True),
    ).order_by(MasCadre.cadre_code)).all())
    permitted = list(db.session.execute(sa.select(
        MapOrgLevelCadre.org_level_id,
        MapOrgLevelCadre.cadre_id,
    ).join(MasOrgLevel, MasOrgLevel.org_level_id == MapOrgLevelCadre.org_level_id)
      .join(MasCadre, MasCadre.cadre_id == MapOrgLevelCadre.cadre_id)
      .where(
          MasOrgLevel.project_id == project_id,
          MasOrgLevel.is_active.is_(True),
          MasCadre.project_id == project_id,
          MasCadre.is_active.is_(True),
          MapOrgLevelCadre.is_active.is_(True),
      )).all())
    cadre_ids_by_level = {}
    for level_id, cadre_id in permitted:
        cadre_ids_by_level.setdefault(level_id, set()).add(cadre_id)
    cadres_by_id = {cadre.cadre_id: cadre for cadre in cadres}
    cadres_by_level = {
        level_key: sorted(
            (cadres_by_id[cadre_id] for cadre_id in cadre_ids_by_level.get(level_id, set())),
            key=lambda cadre: cadre.cadre_code,
        )
        for level_id, level_key in level_keys.items()
    }
    active_languages = list(db.session.execute(sa.select(
        MasLanguages.language_code, MasLanguages.language_name,
    ).where(MasLanguages.is_active.is_(True)).order_by(MasLanguages.language_code)).all())

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "project_users"
    instructions = workbook.create_sheet("Instructions")
    reference = workbook.create_sheet("Reference")
    reference.sheet_state = "hidden"

    header_font = Font(name="Arial", bold=True, color="FFFFFF")
    body_font = Font(name="Arial")
    header_fill = PatternFill("solid", fgColor="1F4E78")
    for column, header in enumerate(HEADERS, 1):
        cell = sheet.cell(1, column, header)
        cell.font = header_font
        cell.fill = header_fill
        cell.comment = Comment({
            "email": "Required for an existing or new account.",
            "name": "Required only when creating a new account.",
            "role": "Leave blank to use cadre defaults when unit and cadre are supplied.",
            "org_unit_code": "Leave blank for a whole-project grant.",
            "cadre_code": "Choose a cadre available at the selected unit level.",
            "language_codes": "Use active language codes separated by semicolons, for example en;hi.",
            "phone": "Optional mobile number; required when email is blank.",
        }[header], "DigitVA")
        sheet.column_dimensions[get_column_letter(column)].width = {
            "email": 32, "name": 24, "role": 24, "org_unit_code": 20,
            "cadre_code": 18, "language_codes": 24, "phone": 18,
        }[header]
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = "A1:G1001"
    for row in sheet.iter_rows(min_row=1, max_row=1001, min_col=1, max_col=len(HEADERS)):
        for cell in row:
            cell.font = header_font if cell.row == 1 else body_font
    for row in range(2, 1002):
        sheet.cell(row, 4).number_format = "@"
        sheet.cell(row, 5).number_format = "@"

    # Reference layout: unit lookup (A:D), all cadres (F:G), roles (I),
    # languages (K:L), and one safe-name cadre list per active level (N+).
    reference.append(["unit_key", "org_unit_code", "level_key", "unit_name", "", "cadre_code", "cadre_name", "", "role", "", "language_code", "language_name"])
    for index, unit in enumerate(units, 1):
        reference.cell(index + 1, 1, f"U{index}")
        reference.cell(index + 1, 2, str(unit.unit_code))
        reference.cell(index + 1, 3, level_keys[unit.org_level_id])
        reference.cell(index + 1, 4, _spreadsheet_safe(unit.unit_name))
    for index, cadre in enumerate(cadres, 1):
        reference.cell(index + 1, 6, str(cadre.cadre_code))
        reference.cell(index + 1, 7, _spreadsheet_safe(cadre.cadre_name))
    for index, role in enumerate(_grantable_role_values(actor), 1):
        reference.cell(index + 1, 9, _spreadsheet_safe(role))
    for index, (code, name) in enumerate(active_languages, 1):
        reference.cell(index + 1, 11, _spreadsheet_safe(code))
        reference.cell(index + 1, 12, _spreadsheet_safe(name))

    reference.cell(1, 14, "empty")
    for column_offset, (level_key, level_cadres) in enumerate(cadres_by_level.items(), 15):
        reference.cell(1, column_offset, level_key)
        for row, cadre in enumerate(level_cadres, 2):
            reference.cell(row, column_offset, str(cadre.cadre_code))
    reference_rows = max(
        2,
        len(units) + 1,
        len(cadres) + 1,
        len(_grantable_role_values(actor)) + 1,
        len(active_languages) + 1,
        *(len(level_cadres) + 1 for level_cadres in cadres_by_level.values()),
    )
    reference_columns = max(12, 14 + len(cadres_by_level))
    for row in reference.iter_rows(min_row=1, max_row=reference_rows, min_col=1, max_col=reference_columns):
        for cell in row:
            cell.font = body_font
    reference.sheet_state = "hidden"

    # Every list has a valid range even when the project has no active rows.
    reference.cell(2, 2).value = reference.cell(2, 2).value or None
    reference.cell(2, 3).value = reference.cell(2, 3).value or None
    reference.cell(2, 9).value = reference.cell(2, 9).value or None
    reference.cell(2, 11).value = reference.cell(2, 11).value or None
    _defined_name(workbook, "UnitCodes", reference, 2, max(2, len(units) + 1), 2)
    _defined_name(workbook, "UnitLevelKeys", reference, 2, max(2, len(units) + 1), 3)
    _defined_name(workbook, "GrantableRoles", reference, 2, max(2, len(_grantable_role_values(actor)) + 1), 9)
    _defined_name(workbook, "Languages", reference, 2, max(2, len(active_languages) + 1), 11)
    _defined_name(workbook, "Cadres_empty", reference, 2, 2, 14)
    for column_offset, level_key in enumerate(cadres_by_level, 15):
        _defined_name(workbook, f"Cadres_{level_key}", reference, 2, len(cadres_by_level[level_key]) + 1, column_offset)

    role_validation = DataValidation(type="list", formula1="=GrantableRoles", allow_blank=True)
    role_validation.errorStyle = "stop"
    role_validation.errorTitle = "Choose a listed role"
    role_validation.error = "Select a role from the list, or leave it blank for cadre defaults."
    role_validation.promptTitle = "Role"
    role_validation.prompt = "Choose a permitted role or leave blank for cadre defaults."
    role_validation.showErrorMessage = True
    role_validation.showInputMessage = True
    unit_validation = DataValidation(type="list", formula1="=UnitCodes", allow_blank=True)
    unit_validation.errorStyle = "stop"
    unit_validation.errorTitle = "Choose an active unit"
    unit_validation.error = "Select an active project unit or leave blank for a whole-project grant."
    unit_validation.promptTitle = "Organization unit"
    unit_validation.prompt = "Leave blank for a whole-project grant."
    unit_validation.showErrorMessage = True
    unit_validation.showInputMessage = True
    cadre_validation = DataValidation(
        type="list",
        formula1=(
            '=INDIRECT(IFERROR("Cadres_"&INDEX(UnitLevelKeys,'
            'MATCH($D2,UnitCodes,0)),"Cadres_empty"))'
        ),
        allow_blank=True,
    )
    cadre_validation.errorStyle = "stop"
    cadre_validation.errorTitle = "Choose a cadre for this unit"
    cadre_validation.error = "Select a cadre available at the selected unit level."
    cadre_validation.promptTitle = "Cadre"
    cadre_validation.prompt = "Choose a cadre after selecting an organization unit."
    cadre_validation.showErrorMessage = True
    cadre_validation.showInputMessage = True
    for validation, column in ((role_validation, "C"), (unit_validation, "D"), (cadre_validation, "E")):
        sheet.add_data_validation(validation)
        validation.add(f"{column}2:{column}1001")

    instructions.sheet_view.showGridLines = False
    instruction_rows = [
        ("DigitVA project users import", "Fill only the project_users sheet, then upload the saved workbook."),
        ("Required columns", "Provide an email or valid mobile phone. New accounts also need name and one active language code."),
        ("Whole-project grants", "Leave org_unit_code and cadre_code blank. A role is required for a whole-project grant."),
        ("Unit grants", "Select an active unit first; the cadre list then shows cadres configured at that unit level."),
        ("Default roles", "Leave role blank only when both unit and cadre are supplied; the server expands stored cadre defaults."),
        ("Languages", "Enter one or more active codes separated by semicolons, for example en;hi. Excel cannot safely multi-select here."),
        ("Validation", "Dropdowns are convenience checks. The server rechecks actor permissions, active units/cadres and all values on upload."),
        ("Example", "new@example.org | Asha Worker | interviewer | D01 | ANM | en;hi | 9876543210"),
    ]
    for row, (label, text) in enumerate(instruction_rows, 1):
        instructions.cell(row, 1, label)
        instructions.cell(row, 2, _spreadsheet_safe(text))
        instructions.cell(row, 1).font = Font(name="Arial", bold=True)
        instructions.cell(row, 2).font = body_font
    role_row = len(instruction_rows) + 1
    instructions.cell(role_row, 1, "Role legend")
    instructions.cell(role_row, 2, _spreadsheet_safe("; ".join(
        f"{role} ({ROLE_LABELS.get(role, role)})" for role in _grantable_role_values(actor)
    )))
    instructions.cell(role_row, 1).font = Font(name="Arial", bold=True)
    instructions.cell(role_row, 2).font = body_font
    instructions.column_dimensions["A"].width = 24
    instructions.column_dimensions["B"].width = 110
    instructions.freeze_panes = "A2"
    instructions.sheet_view.showGridLines = False

    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()


def parse_upload(stream, filename):
    """Read a bounded CSV or first-sheet XLSX into user rows."""
    try:
        rows = parse_table(stream, filename, HEADERS, max_bytes=MAX_BYTES,
                           max_rows=MAX_ROWS, require_all=True)
    except TabularImportError as exc:
        raise ProjectUserImportError(str(exc)) from exc
    if not rows:
        raise ProjectUserImportError("The file has no user rows.")
    for row in rows:
        for header in HEADERS:
            row[header] = str(row[header]) if row[header] is not None else ""
    return rows


def parse_csv(stream):
    """Retain the existing CSV API for callers outside the upload route."""
    return parse_upload(stream, "users.csv")


def _expand_default_roles(rows, units, cadres, permissions):
    """Give a row that names no role the default roles of its cadre at its level.

    Owner decision 2026-10-06 (digitva-vjt): a row with a blank ``role``, an
    organization unit and a cadre stands for one grant per ``default_roles``
    entry of that unit level's grid row. A row that names a role is left
    alone, and so is a blank-role row the defaults cannot serve (no unit or
    cadre, an unknown one, a grid row with no defaults): it fails the usual
    ``role must be one of`` check. Each expanded row is validated like an
    explicit one, so a default the importer may not grant fails that row.
    """
    for row in rows:
        unit = units.get(row["org_unit_code"].upper()) if row["org_unit_code"] else None
        cadre = cadres.get(row["cadre_code"].upper()) if row["cadre_code"] else None
        permission = permissions.get((unit.org_level_id, cadre.cadre_id)) if unit and cadre else None
        if row["role"].strip() or permission is None or not permission.default_roles:
            yield row
            continue
        for role in permission.default_roles:
            yield {**row, "role": role}


def prepare(project_id, rows, *, actor):
    """Resolve all users, roles and scopes without writing; return an apply plan.

    *actor* is the importing user: every row's grant must pass
    ``authz.can_grant`` for them (an admin writes anything; a project PI any
    role but admin and project_pi in their project). Raises
    ProjectUserImportError with row numbers. Existing user profiles are
    ignored; profile fields are used only when an admin creates a new account.

    A row with no email names its person by mobile number (docs/policy/
    account-onboarding-and-passwords.md section 2): an existing account whose
    sign-in number it is, or -- admin only -- a new mobile-only account,
    which needs a valid number no other account holds.
    """
    actor_grants = authz.resolve_grants(actor)
    is_admin = actor_grants.is_admin
    project = db.session.get(VaProjectMaster, project_id)
    if not project or project.project_status != VaStatuses.active:
        raise ProjectUserImportError("Active project not found.")
    if project.project_structure_mode != "organization":
        raise ProjectUserImportError("This project is not in organization mode.")

    emails = {(row["email"] or "").strip().lower() for row in rows} - {""}
    users = {user.email.lower(): user for user in db.session.scalars(
        sa.select(VaUsers).where(sa.func.lower(VaUsers.email).in_(emails))
    )}
    # Rows without an email: their people by sign-in number, one query.
    row_mobiles = {canonical_mobile(row["phone"]) for row in rows
                   if not (row["email"] or "").strip()} - {None}
    users_by_mobile = {user.mobile_login: user for user in db.session.scalars(
        sa.select(VaUsers).where(VaUsers.mobile_login.in_(row_mobiles))
    )} if row_mobiles else {}
    known_users = list(users.values()) + list(users_by_mobile.values())
    requested_units = {row["org_unit_code"].upper() for row in rows if row["org_unit_code"]}
    units = {unit.unit_code.upper(): unit for unit in db.session.scalars(
        sa.select(MasOrgUnit).where(MasOrgUnit.project_id == project_id,
                                    MasOrgUnit.unit_code.in_(requested_units))
    )}
    requested_cadres = {row["cadre_code"].upper() for row in rows if row["cadre_code"]}
    cadres = {cadre.cadre_code.upper(): cadre for cadre in db.session.scalars(
        sa.select(MasCadre).where(MasCadre.project_id == project_id,
                                  MasCadre.cadre_code.in_(requested_cadres))
    )}
    permissions = {(permission.org_level_id, permission.cadre_id): permission
                   for permission in db.session.scalars(
                       sa.select(MapOrgLevelCadre).where(
                           MapOrgLevelCadre.org_level_id.in_(
                               [unit.org_level_id for unit in units.values()]
                           ),
                           MapOrgLevelCadre.cadre_id.in_(
                               [cadre.cadre_id for cadre in cadres.values()]
                           ),
                       )
                   )}
    active_languages = set(db.session.scalars(
        sa.select(MasLanguages.language_code).where(MasLanguages.is_active.is_(True))
    ))
    existing_grants = {(grant.user_id, grant.role, grant.scope_type,
                        grant.org_unit_id if grant.scope_type == VaAccessScopeTypes.org_unit else grant.project_id): grant
                       for grant in db.session.scalars(sa.select(VaUserAccessGrants).where(
                           VaUserAccessGrants.user_id.in_([user.user_id for user in known_users]),
                           sa.or_(VaUserAccessGrants.project_id == project_id,
                                  VaUserAccessGrants.org_unit_id.in_([unit.org_unit_id for unit in units.values()])),
                       ))}
    mentor_members = member_user_ids(user.user_id for user in known_users)
    # docs/policy/mobile-sign-in.md section 1: a new account's phone that is a
    # valid mobile number becomes its sign-in number, so it must be unique:
    # one query for every number in the file, plus repeats within the file.
    file_mobiles = {canonical_mobile(row["phone"]) for row in rows} - {None}
    taken_mobiles = set()
    if file_mobiles:
        for held in db.session.execute(
            sa.select(VaUsers.mobile_login, PHONE_CANONICAL).where(sa.or_(
                VaUsers.mobile_login.in_(file_mobiles), PHONE_CANONICAL.in_(file_mobiles)
            ))
        ):
            taken_mobiles.update(held)
    mobile_rows = {}
    # can_grant per distinct (role, unit): rows repeat targets, and each
    # decision locates its target with one query.
    permitted = {}
    plan = []
    seen = set()
    new_profiles = {}
    errors = []
    failed_rows = set()
    for row in _expand_default_roles(rows, units, cadres, permissions):
        number = row["_line_number"]
        try:
            email = row["email"].strip().lower()
            mobile = canonical_mobile(row["phone"])
            if email:
                if len(email) > 128 or not EMAIL_RE.fullmatch(email):
                    raise ProjectUserImportError("invalid email")
                identity = email
            elif mobile is None:
                raise ProjectUserImportError(
                    "a row without an email needs a valid 10-digit mobile number")
            else:
                identity = f"mobile:{mobile}"
            allowed_roles = [role for role in VaAccessRoles
                             if role != VaAccessRoles.admin and
                             (is_admin or role != VaAccessRoles.project_pi)]
            if row["role"] not in {role.value for role in allowed_roles}:
                raise ProjectUserImportError("role must be one of: " + ", ".join(role.value for role in allowed_roles))
            role = VaAccessRoles(row["role"])
            if role == VaAccessRoles.admin or (role == VaAccessRoles.project_pi and not is_admin):
                raise ProjectUserImportError("role is not grantable")
            unit_code = row["org_unit_code"].upper()
            unit = units.get(unit_code) if unit_code else None
            # Authorize before validating the unit and cadre
            # (docs/policy/access-control-model.md): a non-admin gets one
            # refusal for a unit that is missing, inactive or out of scope,
            # so the import cannot probe which unit codes exist.
            if unit_code and (unit is None or not unit.is_active):
                if not is_admin:
                    raise ProjectUserImportError(NOT_PERMITTED)
                raise ProjectUserImportError("unit code is unknown or inactive")
            target_key = (role, unit.org_unit_id if unit else None)
            if target_key not in permitted:
                permitted[target_key] = bool(authz.can_grant(actor, authz.GrantTarget(
                    role=role,
                    scope_type=VaAccessScopeTypes.org_unit if unit else VaAccessScopeTypes.project,
                    project_id=None if unit else project_id,
                    org_unit_id=unit.org_unit_id if unit else None,
                ), _grants=actor_grants))
            if not permitted[target_key]:
                raise ProjectUserImportError(NOT_PERMITTED)
            # site_pi in an organization project is the In-charge, held at a unit.
            if not unit and role in (
                VaAccessRoles.interview_supervisor, VaAccessRoles.death_reporter, VaAccessRoles.site_pi
            ):
                raise ProjectUserImportError(f"{role.value} requires an organization unit")
            cadre_code = row["cadre_code"].upper()
            if cadre_code and not unit:
                raise ProjectUserImportError("cadre requires an organization unit")
            cadre = cadres.get(cadre_code) if cadre_code else None
            if cadre_code and cadre is None:
                raise ProjectUserImportError("cadre code is unknown")
            if unit:
                if role not in ROLES_ALLOWING_ORG_UNIT:
                    raise ProjectUserImportError("role cannot use organization unit scope")
                if cadre and not cadre.is_active:
                    raise ProjectUserImportError("cadre is inactive")
                permission = permissions.get((unit.org_level_id, cadre.cadre_id)) if cadre else None
                if cadre and (permission is None or not permission.is_active):
                    raise ProjectUserImportError("cadre is not active at this unit's level")
                if role in CADRE_FLAG_BY_ROLE:
                    flag, permits = CADRE_FLAG_BY_ROLE[role]
                    if permission is None or not getattr(permission, flag):
                        raise ProjectUserImportError(
                            f"{role.value} needs a cadre permitted to {permits} at this unit's level"
                        )
            key = (identity, role, unit_code)
            if key in seen:
                raise ProjectUserImportError("duplicate person, role and scope")
            seen.add(key)
            user = users.get(email) if email else users_by_mobile.get(mobile)
            if user and user.user_id in mentor_members:
                check_mentor_grant(user.user_id, role, unit)
            if not is_admin and (not user or user.user_status != VaStatuses.active):
                raise ProjectUserImportError("account is unavailable for this project import")
            if user and user.user_status != VaStatuses.active:
                raise ProjectUserImportError("existing user is inactive")
            languages = [code.strip() for code in row["language_codes"].split(";") if code.strip()]
            if not user:
                if not is_admin:
                    raise ProjectUserImportError("project PI cannot create accounts")
                if not row["name"] or len(row["name"]) > 128:
                    raise ProjectUserImportError("new user needs a name of at most 128 characters")
                if not languages or any(code not in active_languages for code in languages):
                    raise ProjectUserImportError("new user needs active language codes")
                if len(row["phone"]) > 15:
                    raise ProjectUserImportError("phone is too long")
                if mobile and (
                    mobile in taken_mobiles or mobile_rows.setdefault(mobile, identity) != identity
                ):
                    raise ProjectUserImportError("mobile number is already used by another account")
                profile = (row["name"], tuple(languages), row["phone"])
                if identity in new_profiles and new_profiles[identity] != profile:
                    raise ProjectUserImportError("new user profile differs from an earlier row")
                new_profiles[identity] = profile
            scope = VaAccessScopeTypes.org_unit if unit else VaAccessScopeTypes.project
            grant_key = (user.user_id, role, scope, unit.org_unit_id if unit else project_id) if user else None
            grant = existing_grants.get(grant_key)
            plan.append({"row": number, "email": email or None, "identity": identity,
                         # What the preview shows: never a whole mobile number.
                         "display": email or mask_mobile(mobile),
                         "name": row["name"], "phone": row["phone"],
                         "mobile": None if user else mobile,
                         "languages": languages, "role": role, "unit": unit, "cadre": cadre,
                         "user": user, "grant": grant,
                         "action": "create_user" if not user else ("reactivate" if grant and grant.grant_status != VaStatuses.active else "update_cadre" if grant and unit and cadre and grant.cadre_id != cadre.cadre_id else "retain" if grant else "grant")})
        except (ValueError, KeyError) as exc:
            # An expanded row can fail once per default role: report its original row once.
            if number not in failed_rows:
                failed_rows.add(number)
                errors.append(f"Row {number}: {exc}")
    if errors:
        raise ProjectUserImportError("; ".join(errors))
    return plan


def apply(project_id, plan, *, actor_user_id):
    """Write a validated plan in the caller's transaction.

    Returns ``(invitations, audit, sign_in_codes)``: new email accounts to
    send verification to, grant changes to log, and one first sign-in code
    per new mobile-only account (``row``, masked ``mobile``, ``name``,
    ``sign_in_code``) for the caller to show once and never log.
    """
    from app.services import mobile_sign_in_service

    new_users = {}
    changed_grants = []
    sign_in_codes = []
    for item in plan:
        user = item["user"] or new_users.get(item["identity"])
        if user is None:
            user = VaUsers(email=item["email"], name=item["name"], phone=item["phone"] or None,
                           mobile_login=item["mobile"],
                           user_status=VaStatuses.active, vacode_language=item["languages"],
                           permission={}, landing_page="coder", pw_reset_t_and_c=False,
                           email_verified=False, other={"created_by_user_id": str(actor_user_id)})
            user.set_password(secrets.token_urlsafe(32))
            db.session.add(user)
            db.session.flush()
            record_account_created(user, via="project_import", actor_user_id=actor_user_id)
            new_users[item["identity"]] = user
            if user.email is None:
                sign_in_codes.append({
                    "row": item["row"], "mobile": mask_mobile(user.mobile_login),
                    "name": user.name,
                    "sign_in_code": mobile_sign_in_service.issue_code(
                        user, actor_user_id=actor_user_id),
                })
        # Organization worker creation uses the same transaction and needs
        # the account created for an expanded default-role row.
        item["user"] = user
        grant = item["grant"]
        if grant is None:
            unit = item["unit"]
            grant = VaUserAccessGrants(
                user_id=user.user_id, role=item["role"],
                scope_type=VaAccessScopeTypes.org_unit if unit else VaAccessScopeTypes.project,
                project_id=None if unit else project_id,
                org_unit_id=unit.org_unit_id if unit else None,
                cadre_id=item["cadre"].cadre_id if item["cadre"] else None,
                grant_status=VaStatuses.active, created_by_user_id=actor_user_id,
            )
            db.session.add(grant)
            changed_grants.append((grant, "grant_created"))
        elif grant.grant_status != VaStatuses.active or item["action"] == "update_cadre":
            grant.grant_status = VaStatuses.active
            if item["cadre"] is not None or item["action"] == "reactivate":
                grant.cadre_id = item["cadre"].cadre_id if item["cadre"] else None
            changed_grants.append((grant, "grant_updated" if item["action"] == "update_cadre" else "grant_reactivated"))
    db.session.flush()
    # Snapshot values before commit expires ORM state; invitation helpers need
    # only name, email and ID, and audit logging needs grant scalar fields.
    invitations = [SimpleNamespace(user_id=user.user_id, email=user.email, name=user.name)
                   for user in new_users.values() if user.email]
    audit = [(dict(user_id=grant.user_id, grant_id=grant.grant_id, role=grant.role.value,
                   scope_type=grant.scope_type.value, org_unit_id=grant.org_unit_id,
                   cadre_id=grant.cadre_id), action) for grant, action in changed_grants]
    return invitations, audit, sign_in_codes


def finalize_apply(
    project_id,
    *,
    actor_user_id,
    actor_role,
    changed_grants,
    invitations,
    request_ip=None,
):
    """Run shared post-commit grant audit, cache invalidation and invites."""
    from app.logging.va_logger import log_grant_action

    for grant, action in changed_grants:
        authz.invalidate(grant["user_id"])
        log_grant_action(
            action=action,
            actor_user_id=actor_user_id,
            actor_role=actor_role,
            target_user_id=grant["user_id"],
            grant_id=grant["grant_id"],
            role=grant["role"],
            scope_type=grant["scope_type"],
            project_id=project_id,
            org_unit_id=grant["org_unit_id"],
            cadre_id=grant["cadre_id"],
            request_ip=request_ip,
        )

    invite_warnings = []
    invitations_queued = 0
    if invitations:
        from app.services.email_service import is_mail_configured, send_verification_email
        from app.services.token_service import generate_token

        if not is_mail_configured():
            invite_warnings.append(
                "Email delivery is not configured; invitations were not queued. "
                "Resend from Users after configuring email."
            )
        else:
            for user in invitations:
                try:
                    if send_verification_email(
                        user,
                        generate_token(user.user_id, "email_verify"),
                        actor_user_id=actor_user_id,
                        commit=False,
                    ):
                        invitations_queued += 1
                    else:
                        invite_warnings.append(
                            f"Invitation for {user.email} was skipped; resend from Users."
                        )
                except Exception:
                    log.exception(
                        "project users invitation failed | project=%s user_id=%s",
                        project_id,
                        user.user_id,
                    )
                    invite_warnings.append(
                        f"Invitation for {user.email} could not be queued; resend from Users."
                    )
            db.session.commit()
    return invitations_queued, invite_warnings
