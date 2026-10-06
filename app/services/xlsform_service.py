"""The project's ODK XLSForm: a project output (digitva-aek).

Policy: docs/policy/va-form-project-configuration.md, "The ODK form is a
project output".

A project admin or the project PI downloads one ``.xlsx`` (``survey``,
``choices``, ``settings``) for the WHO VA 2022 questionnaire carrying exactly
the extensions the project has switched on and the languages it serves, so ODK
and the web form collect the same field names and relevance. It is built from:

* the WHO reference workbook kept in the repository
  (``REFERENCE_WORKBOOK``, English columns only: every other language comes
  from the translation tables, below);
* one row spec per extension (``resource/xlsform_extensions/<extension>.json``,
  generated from the web form's own definition by
  ``tooling/who-va-2022/build-odk-extension-rows.mjs``; the DORIS support rows
  are read from ``vendor/who-va-2022/src/generated/odk-doris-support-rows.json``
  where their own generator writes them). A spec is a list of *blocks*: where
  the block goes in the WHO form (``odk.after`` / ``odk.before`` a row, or
  ``odk.afterGroupEnd`` a group), its ``survey`` and ``choices`` rows, and
  ``change`` cells that replace WHO cells (DORIS only);
* the project: its welcome note (``intake_screen``), organization levels and
  units (``geography``, the same choices rows the ODK choices CSV exports),
  narration languages, sites and display languages.

An extension that is enabled but has no spec is refused (422, naming it)
rather than emitted partially.

Translations are the approved **and** active locales of the instrument, taken
from ``export_translations`` (accepted rows only; machine drafts never), and
written as ``label::<Language> (<code>)`` / ``hint::`` / ``guidance_hint::`` /
``constraint_message::`` columns. A project that limits its languages
(``web_intake_available_locales``) gets only those.

The reference, the specs and nothing else static are parsed once per process
and copied per build.
"""
from __future__ import annotations

import copy
import io
import json
import logging
import re
from dataclasses import dataclass, field
from datetime import UTC, datetime
from functools import lru_cache
from pathlib import Path
from typing import Any

import sqlalchemy as sa
from openpyxl import Workbook, load_workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE

from app import db
from app.models import (
    MapProjectSiteOdk,
    MasInstrumentLocales,
    MasOrgUnit,
    VaProjectMaster,
    VaProjectSites,
    VaSiteMaster,
    VaStatuses,
)
from app.models.mas_instrument_locales import LIFECYCLE_APPROVED

log = logging.getLogger(__name__)

REPO_ROOT = Path(__file__).resolve().parents[2]

#: The WHO reference form the web instrument is rebuilt from (decided
#: 2026-09-20, ``digitva-13x``): V2.0, English columns only. Its other
#: languages are ignored; translations come from the database. This is the
#: vendored copy, byte-identical to ``docs/kb/WHO_VA_2022_Docs/`` (a test pins
#: that). ``.dockerignore`` leaves ``docs/`` out of the image, so ``vendor/`` is
#: the copy this module can rely on by itself; it is not the only runtime
#: dependency on ``docs/kb``, though: ``export_translations`` reads the V1.1
#: reference workbook there (``instrument_translation_service.REFERENCE_WORKBOOK``),
#: and works because docker-compose bind-mounts the working tree at ``/app``.
REFERENCE_WORKBOOK = REPO_ROOT / "vendor/who-va-2022/2022whova_xls_form_for_odk_multilingual.xlsx"
SPEC_DIR = REPO_ROOT / "resource/xlsform_extensions"
#: Written by ``tooling/who-va-2022/build-odk-doris-rows.mjs`` (same block schema).
DORIS_SPEC = REPO_ROOT / "vendor/who-va-2022/src/generated/odk-doris-support-rows.json"

INSTRUMENT_CODE = "WHO_2022_VA"
BASE_LOCALE = "en"
BASE_LANGUAGE_NAME = "English"

#: Every extension name a project can carry, in the order their blocks are
#: applied (blocks anchored at the same place stay in this order).
EXTENSION_ORDER = (
    "intake_screen",
    "digitva_core",
    "geography",
    "abha",
    "narration_language",
    "social_autopsy",
    "death_summary",
    "medical_records",
    "doris_support_whova_2022",
)

#: Extensions whose rows depend on the project, not on a spec file.
_PROJECT_EXTENSIONS = frozenset({"geography", "intake_screen"})

#: The deliberate departures of the web instrument from the reference (the two
#: ``digitva-13x`` relevance/age-group fixes,
#: the curated ``Id10477``-``Id10479`` lists, name regexes, messages, the
#: ``nmh`` section move, the ``consented`` label). One list, read by
#: ``tooling/who-va-2022/build-instrument-from-xlsform.py`` for the web
#: instrument and applied here to the ODK form, so the two cannot drift.
DEVIATIONS_FILE = REPO_ROOT / "resource/who_va_2022_deviations.json"

#: A project's org-unit choices are bounded so one download cannot build an
#: unbounded workbook.
MAX_ORG_CHOICES = 50_000

#: Translatable survey fields; choices carry ``label`` only.
_TRANSLATED_SURVEY_FIELDS = ("label", "hint", "guidance_hint", "constraint_message")

_SURVEY_COLUMNS = (
    "order", "type", "name", "relevant", "agegroup", "label", "hint", "guidance_hint",
    "required", "notes", "appearance", "calculation", "default", "constraint",
    "constraint_message", "read_only", "parameters", "choice_filter",
)
_LANGUAGE_HEADER = re.compile(r"^(label|hint|guidance_hint|constraint_message)::\s*(.+?)\s*\(([^()]+)\)\s*$")
_FORM_ID = re.compile(r"^[A-Za-z0-9_.-]{1,100}$")


class XlsFormError(RuntimeError):
    """The form cannot be built; ``status_code`` is the HTTP answer for it."""

    status_code = 422

    def __init__(self, message: str, *, status_code: int | None = None, forms: list[str] | None = None):
        super().__init__(message)
        if status_code is not None:
            self.status_code = status_code
        self.forms = forms or []


@dataclass
class _Block:
    extension: str
    block_id: str
    anchor: tuple[str, str] | None
    survey: list[dict]
    choices: list[dict]
    change: list[dict]


@dataclass
class XlsForm:
    """A composed form: the rows, before they are written to a workbook."""

    form_id: str
    version: str
    title: str
    languages: list[tuple[str, str]]
    survey: list[dict]
    choices: list[dict]
    extensions: list[str] = field(default_factory=list)
    style: str = "pages"

    @property
    def filename(self) -> str:
        return f"{self.form_id}_{self.version}.xlsx"


# ---------------------------------------------------------------------------
# Static inputs: parsed once per process
# ---------------------------------------------------------------------------


def _blank(value: Any) -> Any:
    """None for an empty or whitespace-only cell; text is stripped (as the web
    instrument builder strips every cell, so both read the same strings)."""
    if isinstance(value, str):
        return value.strip() or None
    return value


def read_sheet(path: Path, sheet: str) -> list[dict]:
    """One sheet as English-only logical rows.

    A ``label::English (en)`` / ``label`` column becomes ``label``; every other
    language column is dropped (translations come from the database). Blank
    cells are absent. Used for the reference and for a deployed workbook.
    """
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        if sheet not in workbook.sheetnames:
            raise XlsFormError(f"{path.name} has no {sheet!r} sheet.")
        iterator = workbook[sheet].iter_rows(values_only=True)
        header = next(iterator, None)
        if header is None:
            return []
        keys: list[str | None] = []
        for cell in header:
            text = str(cell).strip() if cell is not None else ""
            match = _LANGUAGE_HEADER.match(text)
            if match:
                keys.append(match.group(1) if match.group(3) == BASE_LOCALE else None)
            elif "::" in text:
                keys.append(None)
            else:
                keys.append(text or None)
        rows = []
        for values in iterator:
            row = {k: _blank(v) for k, v in zip(keys, values) if k and _blank(v) is not None}
            if row:
                rows.append(row)
        return rows
    finally:
        workbook.close()


def _text_row(row: dict) -> dict:
    """Spec cells are text; a number from a workbook stays as it was."""
    return {k: v for k, v in row.items() if _blank(v) is not None}


@lru_cache(maxsize=1)
def _reference() -> tuple[tuple[dict, ...], tuple[dict, ...], dict]:
    """``(survey rows, choices rows, settings)`` of the WHO reference."""
    if not REFERENCE_WORKBOOK.exists():
        raise XlsFormError(f"The WHO reference workbook is missing: {REFERENCE_WORKBOOK.name}.", status_code=500)
    survey = read_sheet(REFERENCE_WORKBOOK, "survey")
    choices = read_sheet(REFERENCE_WORKBOOK, "choices")
    settings = (read_sheet(REFERENCE_WORKBOOK, "settings") or [{}])[0]
    return tuple(survey), tuple(choices), settings


@lru_cache(maxsize=16)
def _spec_blocks(extension: str) -> tuple[_Block, ...] | None:
    """The blocks of one extension's spec file, or None when it has none."""
    path = DORIS_SPEC if extension == "doris_support_whova_2022" else SPEC_DIR / f"{extension}.json"
    if not path.exists():
        return None
    try:
        doc = json.loads(path.read_text(encoding="utf-8"))
    except (OSError, ValueError) as exc:
        raise XlsFormError(f"The {extension} row spec could not be read: {exc}", status_code=500) from exc
    blocks = []
    for raw in doc.get("blocks", ()):
        odk = raw.get("odk") or {}
        anchor = next(((kind, odk[kind]) for kind in ("after", "before", "afterGroupEnd") if odk.get(kind)), None)
        blocks.append(
            _Block(
                extension=extension,
                block_id=raw["id"],
                anchor=anchor,
                survey=[_text_row(r) for r in raw.get("survey", ())],
                choices=[_text_row(r) for r in raw.get("choices", ())],
                change=list(raw.get("change", ())),
            )
        )
    return tuple(blocks)


# ---------------------------------------------------------------------------
# Project inputs
# ---------------------------------------------------------------------------


def mapped_form_ids(project_id: str) -> list[str]:
    """The project's ODK form ids, one per mapped form, in name order."""
    return list(
        db.session.scalars(
            sa.select(MapProjectSiteOdk.odk_form_id)
            .where(MapProjectSiteOdk.project_id == project_id)
            .distinct()
            .order_by(MapProjectSiteOdk.odk_form_id)
        ).all()
    )


def resolve_form_id(project_id: str, requested: str | None) -> str:
    """The ODK ``form_id`` this project's form keeps (Central sync stays unbroken).

    The project's own mapping decides it, never the request: ``requested`` only
    chooses among the mapped ids. One mapped form -> it. None mapped (a
    web-only project) -> ``<PROJECT_ID>_WHOVA2022``. Several (one per site,
    e.g. ICMR01) -> the caller must say which; 409 lists them.
    """
    forms = mapped_form_ids(project_id)
    requested = (requested or "").strip()
    if requested:
        if not _FORM_ID.match(requested) or requested not in forms:
            raise XlsFormError(
                "That is not one of this project's ODK forms.", status_code=404, forms=forms
            )
        return requested
    if len(forms) == 1:
        return forms[0]
    if not forms:
        return f"{project_id}_WHOVA2022"
    raise XlsFormError(
        "This project has several ODK forms; choose one with form_id=: " + ", ".join(forms) + ".",
        status_code=409,
        forms=forms,
    )


def _display_locales(project: VaProjectMaster) -> list[tuple[str, str]]:
    """Approved and active locales the project serves, except the base ``en``.

    ``web_intake_available_locales`` NULL means every one; a stored list keeps
    its order and drops what the instrument lacks. ``in_review`` and ``draft``
    locales are never included (the web form's "English alongside" mode has no
    ODK equivalent).
    """
    rows = db.session.execute(
        sa.select(MasInstrumentLocales.locale_code, MasInstrumentLocales.language_name).where(
            MasInstrumentLocales.instrument_code == INSTRUMENT_CODE,
            MasInstrumentLocales.is_active.is_(True),
            MasInstrumentLocales.lifecycle_state == LIFECYCLE_APPROVED,
            MasInstrumentLocales.locale_code != BASE_LOCALE,
        )
    ).all()
    names = {row.locale_code: row.language_name for row in rows}
    stored = project.web_intake_available_locales
    codes = sorted(names) if stored is None else [c for c in stored if c in names]
    return [(code, names[code]) for code in dict.fromkeys(codes)]


def _site_choices(project_id: str, form_id: str) -> list[tuple[str, str]]:
    """``(site_id, site_name)`` for the ``Site`` list: the form's sites when it
    is mapped (a deployed form offers just its own site), else the project's."""
    mapped = db.session.execute(
        sa.select(VaSiteMaster.site_id, VaSiteMaster.site_name)
        .join(MapProjectSiteOdk, MapProjectSiteOdk.site_id == VaSiteMaster.site_id)
        .where(MapProjectSiteOdk.project_id == project_id, MapProjectSiteOdk.odk_form_id == form_id)
        .distinct()
        .order_by(VaSiteMaster.site_id)
    ).all()
    if mapped:
        return [(r.site_id, r.site_name) for r in mapped]
    rows = db.session.execute(
        sa.select(VaSiteMaster.site_id, VaSiteMaster.site_name)
        .join(VaProjectSites, VaProjectSites.site_id == VaSiteMaster.site_id)
        .where(
            VaProjectSites.project_id == project_id,
            VaProjectSites.project_site_status == VaStatuses.active,
        )
        .order_by(VaSiteMaster.site_id)
    ).all()
    return [(r.site_id, r.site_name) for r in rows]


def _geography_block(project_id: str) -> _Block:
    """One cascading ``org_<level>_code`` select per level, filled from the
    same rows the ODK choices CSV exports (the field names routing reads)."""
    from app.services import organization_service as org
    from app.services.org_unit_routing_service import expected_odk_fields

    levels = expected_odk_fields(project_id)
    rows: list[dict] = [
        {"type": "begin group", "name": "org_units", "label": "Organization unit", "appearance": "field-list", "agegroup": "ALL"}
    ]
    previous = None
    for level in levels:
        row = {
            "type": f"select_one {level['choice_list_name']}",
            "name": level["field_name"],
            "label": level["level_name"],
            "agegroup": "ALL",
        }
        if not level["is_optional"]:
            row["required"] = "yes"
        if previous:
            row["choice_filter"] = "parent_code=${" + previous + "}"
        rows.append(row)
        previous = level["field_name"]
    rows.append({"type": "end group"})

    # Count before loading: the cap must bound the read, not only the workbook.
    units = db.session.scalar(
        sa.select(sa.func.count()).select_from(MasOrgUnit).where(
            MasOrgUnit.project_id == project_id, MasOrgUnit.is_active.is_(True)
        )
    )
    if units > MAX_ORG_CHOICES:
        raise XlsFormError(
            f"The organization tree has {units} active units; the ODK form is limited to {MAX_ORG_CHOICES}."
        )
    choices = org.export_odk_choices_rows(project_id)
    return _Block("geography", "geography", ("afterGroupEnd", "Interviewer"), rows, choices, [])


def _intake_screen_block(note: str) -> _Block:
    rows = [
        {"type": "begin group", "name": "begin_screen", "appearance": "field-list", "agegroup": "ALL"},
        {"type": "note", "name": "intake_note", "label": note, "agegroup": "ALL"},
        {"type": "end group"},
    ]
    return _Block("intake_screen", "intake_screen", ("before", "Interviewer"), rows, [], [])


def _project_blocks(project: VaProjectMaster, extensions: list[str]) -> list[_Block]:
    """Every enabled extension's blocks, in ``EXTENSION_ORDER``.

    Raises 422 naming an enabled extension that has neither a spec nor a
    project provider: a form must never be emitted partially.
    """
    from app.services.web_intake_service import resolve_intake_note

    unknown = [e for e in extensions if e not in EXTENSION_ORDER]
    if unknown:
        raise XlsFormError(f"No ODK row spec exists for extension {', '.join(sorted(unknown))}.")
    blocks: list[_Block] = []
    for extension in EXTENSION_ORDER:
        if extension not in extensions:
            continue
        if extension == "geography":
            blocks.append(_geography_block(project.project_id))
        elif extension == "intake_screen":
            blocks.append(_intake_screen_block(resolve_intake_note(project)))
        else:
            spec = _spec_blocks(extension)
            if spec is None:
                raise XlsFormError(f"No ODK row spec exists for extension {extension}.")
            blocks.extend(spec)
    return blocks


# ---------------------------------------------------------------------------
# Composition
# ---------------------------------------------------------------------------


@lru_cache(maxsize=1)
def _deviations() -> dict:
    try:
        return json.loads(DEVIATIONS_FILE.read_text(encoding="utf-8"))
    except (OSError, ValueError) as exc:
        raise XlsFormError(f"The deviations file could not be read: {exc}", status_code=500) from exc


_QUESTION_DEVIATION_KEYS = {"constraint", "constraintMessage", "relevant", "ageGroup", "sectionPath", "choices", "validation"}


def _apply_deviations(survey: list[dict], choices: list[dict]) -> list[dict]:
    """Apply ``resource/who_va_2022_deviations.json`` to the WHO rows; returns
    the choices. Entries are in the instrument's own shape: ``constraint`` /
    ``relevant`` ``{"source": ...}`` (``null`` drops it), ``constraintMessage``
    ``{"en": ...}`` (``{}`` drops it), ``ageGroup``, ``sectionPath`` (the row
    moves to the top of that group) and ``choices`` (that question's list is
    replaced; the list must be its alone). ``validation`` is derived and
    ignored. The ``language`` question's choices are the one exception: the
    web serves English only, ODK offers the project's locales
    (``compose_project_form``), which meets the same "no placeholder" rule. A
    key this does not understand fails loudly rather than being dropped.
    """
    deviations = _deviations()
    rows = {row["name"]: row for row in survey if row.get("name")}
    for name, change in deviations["questions"].items():
        row = rows.get(name)
        if row is None:
            raise XlsFormError(f"The deviations name {name}, which the WHO form lacks.", status_code=500)
        unknown = set(change) - _QUESTION_DEVIATION_KEYS
        if unknown:
            raise XlsFormError(f"Deviation {name} has keys the ODK form cannot apply: {sorted(unknown)}.", status_code=500)
        for key in ("constraint", "relevant"):
            if key in change:
                _set_cell(row, key, (change[key] or {}).get("source"))
        if "constraintMessage" in change:
            _set_cell(row, "constraint_message", change["constraintMessage"].get("en"))
        if "ageGroup" in change:
            row["agegroup"] = change["ageGroup"]
        if "choices" in change and name != "language":
            choices = _replace_list(survey, choices, row, change["choices"])
        if "sectionPath" in change:
            _move_into_group(survey, row, change["sectionPath"][-1])
    for name, change in deviations["sections"].items():
        row = rows.get(name)
        if row is None or set(change) - {"label"}:
            raise XlsFormError(f"Section deviation {name} cannot be applied.", status_code=500)
        row["label"] = change["label"]["en"]
    return choices


def _set_cell(row: dict, column: str, value: str | None) -> None:
    """Set a cell, or clear it for None / empty."""
    if value:
        row[column] = value
    else:
        row.pop(column, None)


def _replace_list(survey: list[dict], choices: list[dict], row: dict, new: list[dict]) -> list[dict]:
    """The choices with this question's list replaced, in the same place."""
    parts = str(row["type"]).split()
    list_name = parts[1]
    sharing = [r["name"] for r in survey if r is not row and str(r.get("type", "")).split()[1:2] == [list_name]]
    if sharing:
        raise XlsFormError(f"List {list_name} is shared with {sharing}; a deviation cannot replace it.", status_code=500)
    rows = [{"list_name": list_name, "name": c["value"], "label": c["label"]["en"]} for c in new]
    out: list[dict] = []
    placed = False
    for choice in choices:
        if choice["list_name"] != list_name:
            out.append(choice)
        elif not placed:
            out.extend(rows)
            placed = True
    return out if placed else [*out, *rows]


def _move_into_group(survey: list[dict], row: dict, group: str) -> None:
    """Move ``row`` to just after ``begin group <group>``."""
    survey[:] = [other for other in survey if other is not row]
    for index, other in enumerate(survey):
        if other.get("name") == group and str(other.get("type", "")).startswith("begin group"):
            survey.insert(index + 1, row)
            return
    raise XlsFormError(f"Deviation moves {row.get('name')} into {group}, which the WHO form lacks.", status_code=500)


def _apply_changes(survey: list[dict], blocks: list[_Block]) -> None:
    """Replace the WHO cells a block's ``change`` names; refuse a missing row."""
    by_name = {row["name"]: row for row in survey if row.get("name")}
    for block in blocks:
        for change in block.change:
            row = by_name.get(change["name"])
            if row is None:
                raise XlsFormError(f"{block.block_id} changes {change['name']}, which the WHO form lacks.")
            for column, value in change["cells"].items():
                if _blank(value) is None:
                    row.pop(column, None)
                else:
                    row[column] = value


def _splice(survey: list[dict], blocks: list[_Block]) -> list[dict]:
    """The WHO rows with every block inserted at its anchor, in one pass.

    Anchors are matched on row ``name``; ``afterGroupEnd`` is the ``end group``
    that closes the named group. An anchor the form lacks fails loudly.
    """
    inserts: dict[tuple[str, str], list[dict]] = {}
    for block in blocks:
        if block.anchor and block.survey:
            inserts.setdefault(block.anchor, []).extend(copy.deepcopy(block.survey))
    used: set[tuple[str, str]] = set()
    out: list[dict] = []
    groups: list[str] = []

    def emit(key: tuple[str, str]) -> None:
        if key in inserts:
            used.add(key)
            out.extend(inserts[key])

    for row in survey:
        name = row.get("name")
        kind = str(row.get("type", ""))
        if name:
            emit(("before", name))
        out.append(row)
        if kind.startswith("begin group") or kind.startswith("begin_group"):
            groups.append(name)
        elif kind.startswith("end group") or kind.startswith("end_group"):
            emit(("afterGroupEnd", groups.pop()))
        elif name:
            # Specs anchor `after` on questions only: after a group row it
            # would land before the group's contents.
            emit(("after", name))
    missing = sorted(set(inserts) - used)
    if missing:
        raise XlsFormError(
            "The WHO form has no anchor for: " + ", ".join(f"{kind} {name}" for kind, name in missing) + "."
        )
    return out


def _merge_choices(base: list[dict], blocks: list[_Block], skip_lists: set[str]) -> list[dict]:
    """The reference choices plus every block's lists.

    A list already in the reference is not emitted twice; its values must match
    (a divergent list under one name would change what the web form accepts).
    """
    out = [row for row in base if row.get("list_name") not in skip_lists]
    existing: dict[str, list[str]] = {}
    for row in out:
        existing.setdefault(row["list_name"], []).append(str(row["name"]))
    # Two extensions may carry the same list (death_summary and medical_records
    # both use YES_NO_REF); the first copy of each value wins.
    pending: dict[str, dict[str, dict]] = {}
    for block in blocks:
        for row in block.choices:
            pending.setdefault(row["list_name"], {}).setdefault(str(row["name"]), row)
    for list_name, by_value in pending.items():
        rows = list(by_value.values())
        names = list(by_value)
        if list_name in existing:
            if sorted(existing[list_name]) != sorted(names):
                raise XlsFormError(f"Choice list {list_name} differs from the WHO list of that name.")
            continue
        out.extend(copy.deepcopy(rows))
    return out


def _require_choices(survey: list[dict], choices: list[dict]) -> None:
    """Every select needs a choice list with choices: an empty one (a project
    with no sites, an organization level with no units) would make a form ODK
    rejects, so say which list instead."""
    have = {row["list_name"] for row in choices}
    for row in survey:
        parts = str(row.get("type", "")).split()
        if parts and parts[0] in ("select_one", "select_multiple") and len(parts) > 1 and parts[1] not in have:
            raise XlsFormError(
                f"{row.get('name')} uses the choice list {parts[1]}, which has no choices for this project "
                "(add its sites or organization units first)."
            )


def _translations(languages: list[tuple[str, str]]) -> dict[str, dict]:
    from app.services.instrument_translation_service import export_translations

    return {code: export_translations(INSTRUMENT_CODE, code) for code, _name in languages}


def compose_project_form(
    project: VaProjectMaster,
    *,
    form_id: str | None = None,
    now: datetime | None = None,
) -> XlsForm:
    """The project's form as rows. Raises :class:`XlsFormError` when it cannot
    be built faithfully (non-WHO instrument, an extension without a spec, no
    unique ODK form)."""
    # Lazy: the routes package imports services, not the other way round.
    from app.routes.api.organization import (
        _active_languages,
        _resolve_narration_languages,
        project_instrument_and_extensions,
    )

    instrument_code, extensions = project_instrument_and_extensions(project)
    if instrument_code != INSTRUMENT_CODE:
        raise XlsFormError(
            f"Only the {INSTRUMENT_CODE} questionnaire has an ODK form export; this project uses {instrument_code}.",
            status_code=409,
        )
    chosen_form_id = resolve_form_id(project.project_id, form_id)
    blocks = _project_blocks(project, extensions)

    reference_survey, reference_choices, settings = _reference()
    survey = copy.deepcopy(list(reference_survey))
    reference_choices = _apply_deviations(survey, copy.deepcopy(list(reference_choices)))
    _apply_changes(survey, blocks)
    survey = _splice(survey, blocks)

    languages = [(BASE_LOCALE, BASE_LANGUAGE_NAME), *_display_locales(project)]
    # The WHO `language` list is the interview language: this project's display
    # languages replace the reference's placeholders ("Language 2"). Narration
    # languages are another axis and get their own list.
    project_lists = [
        [{"list_name": "language", "name": code, "label": name} for code, name in languages]
    ]
    if "narration_language" in extensions:
        narration = _resolve_narration_languages(project, _active_languages())
        project_lists.append(
            [{"list_name": "narr_language", "name": n["code"], "label": n["label"]} for n in narration]
        )
    if "digitva_core" in extensions:
        project_lists.append(
            [{"list_name": "site", "name": site_id, "label": name}
             for site_id, name in _site_choices(project.project_id, chosen_form_id)]
        )
    project_list_names = {rows[0]["list_name"] for rows in project_lists if rows}
    choices = _merge_choices(reference_choices, blocks, skip_lists={"language"} | project_list_names)
    choices = [*choices, *(row for rows in project_lists for row in rows)]

    _require_choices(survey, choices)
    translations = _translations(languages[1:])
    stamp = (now or datetime.now(UTC)).strftime("%Y%m%d%H%M%S")
    title = f"({project.project_name}) 2022 WHO Verbal Autopsy Instrument"
    form = XlsForm(
        form_id=chosen_form_id,
        version=stamp,
        title=title,
        languages=languages,
        survey=survey,
        choices=choices,
        extensions=list(extensions),
        style=str(settings.get("style") or "pages"),
    )
    _attach_translations(form, translations)
    return form


def _attach_translations(form: XlsForm, translations: dict[str, dict]) -> None:
    """Add ``<field>::<code>`` keys to rows; the choice labels of project lists
    (languages, sites, narration, units) read the same in every language."""
    for row in form.survey:
        name = row.get("name")
        if not name:
            continue
        for code, payload in translations.items():
            for field_name, text in (payload["questions"].get(name) or {}).items():
                if field_name in _TRANSLATED_SURVEY_FIELDS and text:
                    row[f"{field_name}::{code}"] = text
    for row in form.choices:
        label = row.get("label")
        for code, payload in translations.items():
            text = (payload["choices"].get(f"{row['list_name']}/{row['name']}") or {}).get("label")
            if text:
                row[f"label::{code}"] = text
            elif label and (row["list_name"] in {"language", "site", "narr_language"} or row["list_name"].startswith("org_")):
                row[f"label::{code}"] = label


# ---------------------------------------------------------------------------
# Workbook
# ---------------------------------------------------------------------------


def _cell(sheet, row: int, column: int, value: Any) -> None:
    if isinstance(value, str):
        # Control characters (a pasted unit name, a stray \x07) make openpyxl
        # raise; they are never meaningful in a form.
        value = ILLEGAL_CHARACTERS_RE.sub("", value)
    cell = sheet.cell(row=row, column=column, value=value)
    if isinstance(value, str):
        # Strings stay strings: an expression or a unit name starting with `=`
        # must not become a spreadsheet formula.
        cell.data_type = "s"


def _write_sheet(sheet, columns: list[str], rows: list[dict]) -> None:
    for index, name in enumerate(columns, start=1):
        _cell(sheet, 1, index, name)
    for r, row in enumerate(rows, start=2):
        for c, name in enumerate(columns, start=1):
            value = row.get(name)
            if value is not None:
                _cell(sheet, r, c, value)


def _columns(
    languages: list[tuple[str, str]], rows: list[dict], fixed: tuple[str, ...], translated: tuple[str, ...]
) -> list[str]:
    """Header order: the fixed columns with each translated field expanded to
    its language columns (English first), then any other column a row used."""
    columns: list[str] = []
    for name in fixed:
        if name in translated:
            columns.extend(f"{name}::{lang} ({code})" for code, lang in languages)
        else:
            columns.append(name)
    return columns + sorted({key for row in rows for key in row} - set(columns))


def _localized(rows: list[dict], languages: list[tuple[str, str]], translated: tuple[str, ...]) -> list[dict]:
    """Rename ``label`` to ``label::English (en)`` and ``label::hi`` to
    ``label::Hindi (hi)`` so a row is keyed by its header."""
    names = dict(languages)
    out = []
    for row in rows:
        item = {}
        for key, value in row.items():
            head, _, code = key.partition("::")
            if head in translated:
                item[f"{head}::{names.get(code or BASE_LOCALE, code)} ({code or BASE_LOCALE})"] = value
            else:
                item[key] = value
        out.append(item)
    return out


def to_workbook_bytes(form: XlsForm) -> bytes:
    """Write the form as an ``.xlsx`` with ``survey``, ``choices``, ``settings``."""
    workbook = Workbook()
    survey_sheet = workbook.active
    survey_sheet.title = "survey"
    survey_rows = _localized(form.survey, form.languages, _TRANSLATED_SURVEY_FIELDS)
    _write_sheet(
        survey_sheet,
        _columns(form.languages, survey_rows, _SURVEY_COLUMNS, _TRANSLATED_SURVEY_FIELDS),
        survey_rows,
    )
    choice_rows = _localized(form.choices, form.languages, ("label",))
    _write_sheet(
        workbook.create_sheet("choices"),
        _columns(form.languages, choice_rows, ("list_name", "name", "label"), ("label",)),
        choice_rows,
    )
    _write_sheet(
        workbook.create_sheet("settings"),
        ["form_title", "version", "form_id", "style", "default_language"],
        [
            {
                "form_title": form.title,
                "version": form.version,
                "form_id": form.form_id,
                "style": form.style,
                "default_language": f"{BASE_LANGUAGE_NAME} ({BASE_LOCALE})",
            }
        ],
    )
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def build_project_xlsform(
    project: VaProjectMaster, *, form_id: str | None = None, now: datetime | None = None
) -> tuple[XlsForm, bytes]:
    """The composed form and its workbook bytes, for the download."""
    form = compose_project_form(project, form_id=form_id, now=now)
    return form, to_workbook_bytes(form)


# ---------------------------------------------------------------------------
# Diff against a deployed workbook
# ---------------------------------------------------------------------------

#: What "changed" means for a row both forms have: what decides collection.
_DIFF_COLUMNS = ("type", "relevant", "constraint", "calculation", "choice_filter")


def _squash(value: Any) -> str:
    return " ".join(str(value).split()) if value is not None else ""


def diff_against_workbook(form: XlsForm, workbook: Path) -> dict:
    """Rows a deployed workbook has that the generated form lacks, and the
    reverse (site-local rows to review; nothing is merged).

    Survey rows are matched on ``name`` (groups included); choices on
    ``(list_name, name)``. A row both have is listed under ``changed`` when its
    type, relevance, constraint, calculation or choice filter differs.
    """
    deployed_survey = {r["name"]: r for r in read_sheet(workbook, "survey") if r.get("name")}
    deployed_choices = {
        (str(r["list_name"]), str(r["name"])) for r in read_sheet(workbook, "choices") if r.get("list_name") and "name" in r
    }
    generated_survey = {r["name"]: r for r in form.survey if r.get("name")}
    generated_choices = {(str(r["list_name"]), str(r["name"])) for r in form.choices}
    changed = []
    for name in sorted(deployed_survey.keys() & generated_survey.keys()):
        columns = [
            c for c in _DIFF_COLUMNS
            if _squash(deployed_survey[name].get(c)) != _squash(generated_survey[name].get(c))
        ]
        if columns:
            changed.append({"name": name, "columns": columns})
    return {
        "form_id": form.form_id,
        "workbook": workbook.name,
        "survey_only_in_workbook": [n for n in deployed_survey if n not in generated_survey],
        "survey_only_in_generated": [n for n in generated_survey if n not in deployed_survey],
        "survey_changed": changed,
        "choices_only_in_workbook": sorted(deployed_choices - generated_choices),
        "choices_only_in_generated": sorted(generated_choices - deployed_choices),
    }
