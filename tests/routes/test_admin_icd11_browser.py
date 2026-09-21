import io
import json
from datetime import UTC, datetime
from io import BytesIO

import sqlalchemy as sa
from openpyxl import load_workbook

from app import db
from app.models import MasIcd11Mms
from tests.base import BaseTestCase

_POLICY_URL = "/admin/api/icd11/mms/node/policy"
_IMPORT_URL = "/admin/api/icd11/mms/policy-import"


class TestAdminIcd11Browser(BaseTestCase):
    def setUp(self):
        super().setUp()
        db.session.execute(sa.delete(MasIcd11Mms))
        db.session.flush()

        now = datetime.now(UTC)
        db.session.add_all(
            [
                MasIcd11Mms(
                    release="2026-01",
                    linearization_uri="lin:chapter1",
                    foundation_uri="f:chapter1",
                    code=None,
                    title="Certain infectious or parasitic diseases",
                    class_kind="chapter",
                    depth_in_kind=1,
                    chapter_no="01",
                    is_residual=False,
                    is_leaf=False,
                    sort_order=1,
                    parent_foundation_uri=None,
                    parent_linearization_uri=None,
                    source_version="ICD-11-MMS-2026-01",
                    source_path="test",
                    is_active=True,
                    created_at=now,
                    updated_at=now,
                ),
                MasIcd11Mms(
                    release="2026-01",
                    linearization_uri="lin:1A00",
                    foundation_uri="f:1A00",
                    code="1A00",
                    title="Cholera",
                    class_kind="category",
                    depth_in_kind=1,
                    chapter_no="01",
                    is_residual=False,
                    is_leaf=True,
                    sort_order=2,
                    parent_foundation_uri="f:chapter1",
                    parent_linearization_uri="lin:chapter1",
                    is_coding_selectable=True,
                    sex_selectable="both",
                    age_group_selectable="all",
                    source_version="ICD-11-MMS-2026-01",
                    source_path="test",
                    is_active=True,
                    created_at=now,
                    updated_at=now,
                ),
                MasIcd11Mms(
                    release="2026-01",
                    linearization_uri="lin:1A00.0",
                    foundation_uri="f:1A00.0",
                    code="1A00.0",
                    title="Cholera due to Vibrio cholerae O1",
                    class_kind="category",
                    depth_in_kind=2,
                    chapter_no="01",
                    is_residual=False,
                    is_leaf=True,
                    sort_order=3,
                    parent_foundation_uri="f:1A00",
                    parent_linearization_uri="lin:1A00",
                    is_coding_selectable=None,
                    policy_status="unreviewed",
                    source_version="ICD-11-MMS-2026-01",
                    source_path="test",
                    is_active=True,
                    created_at=now,
                    updated_at=now,
                ),
            ]
        )
        db.session.flush()

    def test_admin_panel_renders_for_admin(self):
        self._login(str(self.base_admin_user.user_id))

        response = self.client.get("/admin/panels/icd11-browser")

        self.assertEqual(response.status_code, 200)
        body = response.get_data(as_text=True)
        self.assertIn("ICD-11 MMS Browser", body)

    def test_admin_panel_denied_for_project_pi(self):
        self._login(str(self.base_project_pi_user.user_id))

        response = self.client.get("/admin/panels/icd11-browser")

        self.assertEqual(response.status_code, 403)

    def test_admin_children_api_returns_root_chapter(self):
        self._login(str(self.base_admin_user.user_id))

        response = self.client.get("/admin/api/icd11/mms/children")

        self.assertEqual(response.status_code, 200)
        payload = response.get_json()
        self.assertEqual(len(payload["children"]), 1)
        self.assertEqual(payload["children"][0]["linearization_uri"], "lin:chapter1")
        self.assertEqual(payload["children"][0]["child_count"], 1)

    def test_admin_children_api_returns_category_under_chapter(self):
        self._login(str(self.base_admin_user.user_id))

        response = self.client.get(
            "/admin/api/icd11/mms/children",
            query_string={"parent_linearization_uri": "lin:chapter1"},
        )

        self.assertEqual(response.status_code, 200)
        codes = [row["code"] for row in response.get_json()["children"]]
        self.assertEqual(codes, ["1A00"])

    def test_admin_node_api_returns_details_with_ancestors(self):
        self._login(str(self.base_admin_user.user_id))

        response = self.client.get(
            "/admin/api/icd11/mms/node", query_string={"linearization_uri": "lin:1A00"}
        )

        self.assertEqual(response.status_code, 200)
        payload = response.get_json()
        self.assertEqual(payload["code"], "1A00")
        self.assertEqual(len(payload["ancestors"]), 1)
        self.assertEqual(payload["ancestors"][0]["linearization_uri"], "lin:chapter1")

    def test_admin_node_api_returns_404_for_unknown_node(self):
        self._login(str(self.base_admin_user.user_id))

        response = self.client.get(
            "/admin/api/icd11/mms/node", query_string={"linearization_uri": "lin:missing"}
        )

        self.assertEqual(response.status_code, 404)

    def test_admin_search_api_finds_selectable_category(self):
        self._login(str(self.base_admin_user.user_id))

        response = self.client.get("/admin/api/icd11/mms/search", query_string={"q": "cholera"})

        self.assertEqual(response.status_code, 200)
        results = response.get_json()["results"]
        self.assertTrue(any(row["icd_code"] == "1A00" for row in results))

    def test_admin_api_denied_for_anonymous(self):
        response = self.client.get("/admin/api/icd11/mms/children")

        self.assertIn(response.status_code, (302, 401, 403))

    # -- Parity with the ICD-10 browser: filters, dots, policy editing --------

    def _row(self, uri):
        db.session.expire_all()
        return db.session.scalar(
            sa.select(MasIcd11Mms).where(MasIcd11Mms.linearization_uri == uri)
        )

    def _post_import(self, payload, *, dry_run=False, headers=None):
        data = {"file": (io.BytesIO(json.dumps(payload).encode("utf-8")), "policy.json")}
        if dry_run:
            data["dry_run"] = "1"
        return self.client.post(
            _IMPORT_URL,
            data=data,
            headers=self._csrf_headers() if headers is None else headers,
            content_type="multipart/form-data",
        )

    def test_admin_panel_renders_policy_editor(self):
        self._login(str(self.base_admin_user.user_id))

        body = self.client.get("/admin/panels/icd11-browser").get_data(as_text=True)

        self.assertIn('id="icd11-browser-policy-form"', body)
        self.assertIn("/admin/api/icd11/mms/policy-export.xlsx?release=2026-01", body)
        self.assertIn('id="icd11-policy-import-btn"', body)
        self.assertIn('id="icd11-browser-columns"', body)
        self.assertIn('id="icd11-browser-path"', body)
        self.assertNotIn('id="icd11-browser-tree"', body)

    def test_children_api_reports_status_indicator(self):
        self._login(str(self.base_admin_user.user_id))

        chapter_rows = self.client.get(
            "/admin/api/icd11/mms/children",
            query_string={"parent_linearization_uri": "lin:chapter1"},
        ).get_json()["children"]
        child_rows = self.client.get(
            "/admin/api/icd11/mms/children",
            query_string={"parent_linearization_uri": "lin:1A00"},
        ).get_json()["children"]

        self.assertEqual(chapter_rows[0]["code"], "1A00")
        self.assertEqual(chapter_rows[0]["status_indicator"], "green")
        self.assertEqual(child_rows[0]["code"], "1A00.0")
        self.assertEqual(child_rows[0]["status_indicator"], "red")

    def test_children_api_yellow_when_only_a_child_is_selectable(self):
        cholera = self._row("lin:1A00")
        subtype = self._row("lin:1A00.0")
        cholera.is_coding_selectable = False
        subtype.is_coding_selectable = True
        db.session.flush()
        self._login(str(self.base_admin_user.user_id))

        rows = self.client.get(
            "/admin/api/icd11/mms/children",
            query_string={"parent_linearization_uri": "lin:chapter1"},
        ).get_json()["children"]

        self.assertEqual(rows[0]["code"], "1A00")
        self.assertEqual(rows[0]["status_indicator"], "yellow")

    def test_children_api_disabled_filter_keeps_ancestors_of_matches(self):
        self._login(str(self.base_admin_user.user_id))

        roots = self.client.get(
            "/admin/api/icd11/mms/children", query_string={"coding_filter": "disabled"}
        ).get_json()["children"]
        under_chapter = self.client.get(
            "/admin/api/icd11/mms/children",
            query_string={"parent_linearization_uri": "lin:chapter1", "coding_filter": "disabled"},
        ).get_json()["children"]

        self.assertEqual([row["linearization_uri"] for row in roots], ["lin:chapter1"])
        self.assertEqual(roots[0]["child_count"], 1)
        # 1A00 is selectable itself but stays reachable: its child 1A00.0 matches.
        self.assertEqual([row["code"] for row in under_chapter], ["1A00"])
        self.assertEqual(under_chapter[0]["child_count"], 1)

    def test_children_api_active_filter_counts_only_matching_children(self):
        self._login(str(self.base_admin_user.user_id))

        under_cholera = self.client.get(
            "/admin/api/icd11/mms/children",
            query_string={"parent_linearization_uri": "lin:chapter1", "coding_filter": "active"},
        ).get_json()["children"]

        self.assertEqual([row["code"] for row in under_cholera], ["1A00"])
        # 1A00.0 is not selectable, so 1A00 has no matching children.
        self.assertEqual(under_cholera[0]["child_count"], 0)

    def test_children_api_filter_with_no_match_returns_nothing(self):
        self._login(str(self.base_admin_user.user_id))

        response = self.client.get(
            "/admin/api/icd11/mms/children", query_string={"sex_filter": "male"}
        )

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.get_json()["children"], [])

    def test_children_api_rejects_unknown_filter_value(self):
        self._login(str(self.base_admin_user.user_id))

        response = self.client.get(
            "/admin/api/icd11/mms/children", query_string={"age_filter": "elderly"}
        )

        self.assertEqual(response.status_code, 400)

    def test_api_rejects_malformed_release(self):
        self._login(str(self.base_admin_user.user_id))

        response = self.client.get(
            "/admin/api/icd11/mms/policy-export", query_string={"release": "x\r\ny"}
        )

        self.assertEqual(response.status_code, 400)

    def test_policy_options_match_icd10_value_sets(self):
        self._login(str(self.base_admin_user.user_id))

        icd11 = self.client.get("/admin/api/icd11/mms/policy-options").get_json()
        icd10 = self.client.get("/admin/api/icd10/2019-2/policy-options").get_json()

        self.assertEqual(icd11["sex_selectable"], icd10["sex_selectable"])
        self.assertEqual(icd11["age_group_selectable"], icd10["age_group_selectable"])
        self.assertEqual(icd11["policy_status"], ["unreviewed", "reviewed"])

    def test_policy_patch_updates_category(self):
        self._login(str(self.base_admin_user.user_id))

        response = self.client.patch(
            _POLICY_URL,
            query_string={"linearization_uri": "lin:1A00.0"},
            json={
                "is_coding_selectable": True,
                "sex_selectable": "female",
                "age_group_selectable": "adult",
                "policy_status": "reviewed",
                "restriction_note": "maternal only",
            },
            headers=self._csrf_headers(),
        )

        self.assertEqual(response.status_code, 200, response.get_data(as_text=True))
        self.assertEqual(response.get_json()["status_indicator"], "green")
        row = self._row("lin:1A00.0")
        self.assertTrue(row.is_coding_selectable)
        self.assertEqual(row.sex_selectable, "female")
        self.assertEqual(row.age_group_selectable, "adult")
        self.assertEqual(row.policy_status, "reviewed")
        self.assertEqual(row.restriction_note, "maternal only")

    def test_policy_patch_without_status_keeps_existing_status(self):
        self._row("lin:1A00").policy_status = "reviewed"
        db.session.flush()
        self._login(str(self.base_admin_user.user_id))

        response = self.client.patch(
            _POLICY_URL,
            query_string={"linearization_uri": "lin:1A00"},
            json={"is_coding_selectable": False},
            headers=self._csrf_headers(),
        )

        self.assertEqual(response.status_code, 200)
        row = self._row("lin:1A00")
        self.assertFalse(row.is_coding_selectable)
        self.assertEqual(row.policy_status, "reviewed")

    def test_policy_patch_rejects_invalid_values(self):
        self._login(str(self.base_admin_user.user_id))

        for body in (
            {"is_coding_selectable": "yes"},
            {"sex_selectable": "other"},
            {"age_group_selectable": "elderly"},
            {"policy_status": "approved"},
            {"restriction_note": 5},
        ):
            with self.subTest(body=body):
                response = self.client.patch(
                    _POLICY_URL,
                    query_string={"linearization_uri": "lin:1A00"},
                    json=body,
                    headers=self._csrf_headers(),
                )
                self.assertEqual(response.status_code, 400)
        row = self._row("lin:1A00")
        self.assertTrue(row.is_coding_selectable)
        self.assertEqual(row.sex_selectable, "both")

    def test_policy_patch_rejects_structural_unknown_and_missing_nodes(self):
        self._login(str(self.base_admin_user.user_id))
        body = {"is_coding_selectable": True, "sex_selectable": "both", "age_group_selectable": "all"}

        chapter = self.client.patch(
            _POLICY_URL, query_string={"linearization_uri": "lin:chapter1"},
            json=body, headers=self._csrf_headers(),
        )
        unknown = self.client.patch(
            _POLICY_URL, query_string={"linearization_uri": "lin:missing"},
            json=body, headers=self._csrf_headers(),
        )
        missing = self.client.patch(_POLICY_URL, json=body, headers=self._csrf_headers())

        self.assertEqual(chapter.status_code, 400)
        self.assertEqual(unknown.status_code, 404)
        self.assertEqual(missing.status_code, 400)
        self.assertIsNone(self._row("lin:chapter1").is_coding_selectable)

    def test_policy_patch_requires_csrf_token(self):
        self._login(str(self.base_admin_user.user_id))

        response = self.client.patch(
            _POLICY_URL,
            query_string={"linearization_uri": "lin:1A00"},
            json={"is_coding_selectable": False},
        )

        self.assertEqual(response.status_code, 400)
        self.assertTrue(self._row("lin:1A00").is_coding_selectable)

    def test_policy_writes_denied_for_project_pi(self):
        self._login(str(self.base_project_pi_user.user_id))

        patch = self.client.patch(
            _POLICY_URL,
            query_string={"linearization_uri": "lin:1A00"},
            json={"is_coding_selectable": False},
            headers=self._csrf_headers(),
        )
        imported = self._post_import({"items": []})
        exported = self.client.get("/admin/api/icd11/mms/policy-export")

        self.assertEqual(patch.status_code, 403)
        self.assertEqual(imported.status_code, 403)
        self.assertEqual(exported.status_code, 403)
        self.assertTrue(self._row("lin:1A00").is_coding_selectable)

    def test_policy_export_json_is_attachment_with_curated_rows(self):
        self._login(str(self.base_admin_user.user_id))

        response = self.client.get("/admin/api/icd11/mms/policy-export")

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.mimetype, "application/json")
        self.assertIn(
            'attachment; filename="icd11_mms_2026-01_policy_export.json"',
            response.headers["Content-Disposition"],
        )
        payload = response.get_json()
        self.assertEqual(payload["release"], "2026-01")
        self.assertEqual([item["code"] for item in payload["items"]], ["1A00"])

    def test_policy_export_json_round_trips_through_import(self):
        cholera = self._row("lin:1A00")
        cholera.policy_status = "reviewed"
        cholera.restriction_note = "confirmed outbreaks"
        db.session.flush()
        self._login(str(self.base_admin_user.user_id))
        exported = self.client.get("/admin/api/icd11/mms/policy-export").get_json()

        # Drift the row, then import the export back.
        self.client.patch(
            _POLICY_URL,
            query_string={"linearization_uri": "lin:1A00"},
            json={"is_coding_selectable": False, "policy_status": "unreviewed"},
            headers=self._csrf_headers(),
        )
        response = self._post_import(exported)

        self.assertEqual(response.status_code, 200)
        body = response.get_json()
        self.assertFalse(body["dry_run"])
        self.assertEqual(body["updated_items"], 1)
        self.assertEqual(body["failed_codes"], [])
        row = self._row("lin:1A00")
        self.assertTrue(row.is_coding_selectable)
        self.assertEqual(row.sex_selectable, "both")
        self.assertEqual(row.age_group_selectable, "all")
        self.assertEqual(row.policy_status, "reviewed")
        self.assertEqual(row.restriction_note, "confirmed outbreaks")
        # The re-exported 1A00 item equals the original: nothing is lost on the way.
        # (1A00.0 now exports too: the import reset it from unset to False.)
        again = self.client.get("/admin/api/icd11/mms/policy-export").get_json()
        self.assertEqual(
            next(item for item in again["items"] if item["code"] == "1A00"),
            exported["items"][0],
        )

    def test_policy_import_dry_run_reports_counts_and_changes_nothing(self):
        self._login(str(self.base_admin_user.user_id))
        payload = {
            "items": [
                {
                    "linearization_uri": "lin:1A00.0",
                    "is_coding_selectable": True,
                    "sex_selectable": "male",
                    "age_group_selectable": "child",
                    "policy_status": "reviewed",
                },
                {"linearization_uri": "lin:missing", "is_coding_selectable": True},
            ]
        }

        response = self._post_import(payload, dry_run=True)

        self.assertEqual(response.status_code, 200)
        body = response.get_json()
        self.assertTrue(body["dry_run"])
        self.assertEqual(body["total_items"], 2)
        self.assertEqual(body["updated_items"], 1)
        self.assertEqual(body["reset_items"], 1)
        self.assertEqual(
            body["failed_codes"], [{"code": "lin:missing", "reason": "unknown_or_non_editable_code"}]
        )
        cholera = self._row("lin:1A00")
        self.assertTrue(cholera.is_coding_selectable)
        self.assertEqual(cholera.sex_selectable, "both")
        subtype = self._row("lin:1A00.0")
        self.assertIsNone(subtype.is_coding_selectable)
        self.assertEqual(subtype.policy_status, "unreviewed")

        applied = self._post_import(payload).get_json()

        self.assertFalse(applied["dry_run"])
        self.assertEqual(applied["updated_items"], 1)
        self.assertFalse(self._row("lin:1A00").is_coding_selectable)
        subtype = self._row("lin:1A00.0")
        self.assertTrue(subtype.is_coding_selectable)
        self.assertEqual(subtype.sex_selectable, "male")
        self.assertEqual(subtype.policy_status, "reviewed")

    def test_policy_import_rejects_bad_files(self):
        self._login(str(self.base_admin_user.user_id))

        missing_file = self.client.post(
            _IMPORT_URL, data={}, headers=self._csrf_headers(), content_type="multipart/form-data"
        )
        not_json = self.client.post(
            _IMPORT_URL,
            data={"file": (io.BytesIO(b"not json"), "policy.json")},
            headers=self._csrf_headers(),
            content_type="multipart/form-data",
        )
        no_items = self._post_import({"rows": []})
        bad_uri = self._post_import({"items": [{"linearization_uri": 7}]})

        self.assertEqual(missing_file.status_code, 400)
        self.assertEqual(not_json.status_code, 400)
        self.assertEqual(no_items.status_code, 400)
        self.assertEqual(bad_uri.status_code, 400)
        self.assertTrue(self._row("lin:1A00").is_coding_selectable)

    def test_policy_export_xlsx_lists_categories_with_policy(self):
        self._login(str(self.base_admin_user.user_id))

        response = self.client.get("/admin/api/icd11/mms/policy-export.xlsx")

        self.assertEqual(response.status_code, 200)
        self.assertEqual(
            response.mimetype,
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        self.assertIn(
            'attachment; filename="icd11_mms_2026-01_policy_export.xlsx"',
            response.headers["Content-Disposition"],
        )
        workbook = load_workbook(BytesIO(response.data), read_only=True, data_only=True)
        sheet = workbook["ICD11 Policy"]
        rows = list(sheet.iter_rows(values_only=True))
        headers = rows[0]
        by_code = {row[0]: dict(zip(headers, row)) for row in rows[1:]}
        self.assertEqual(sorted(by_code), ["1A00", "1A00.0"])
        self.assertEqual(by_code["1A00"]["Coding Allowed"], "Yes")
        self.assertEqual(by_code["1A00"]["Sex Selectable"], "both")
        self.assertEqual(by_code["1A00"]["Age Selectable"], "all")
        self.assertEqual(by_code["1A00.0"]["Coding Allowed"], "No")
        self.assertEqual(by_code["1A00.0"]["Linearization URI"], "lin:1A00.0")
