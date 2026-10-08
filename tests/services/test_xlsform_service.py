"""The project's ODK XLSForm download (digitva-aek).

Service (``app/services/xlsform_service.py``), route
(``app/routes/admin_xlsform.py``) and CLI (``app/commands/xlsform.py``).
Policy: docs/policy/va-form-project-configuration.md, "The ODK form is a
project output".
"""
import io
import json
from datetime import UTC, datetime
from pathlib import Path
from unittest.mock import patch

import sqlalchemy as sa
from openpyxl import load_workbook
from pyxform.xls2xform import convert

from app import db
from app.models import (
    MapProjectSiteOdk,
    VaAccessRoles,
    VaAccessScopeTypes,
    VaProjectMaster,
    VaProjectSites,
    VaSiteMaster,
    VaStatuses,
    VaUserAccessGrants,
)
from app.models.mas_instrument_locales import (
    LIFECYCLE_APPROVED,
    LIFECYCLE_DRAFT,
    LIFECYCLE_IN_REVIEW,
    MapInstrumentTranslations,
    MasInstrumentLocales,
)
from app.models.mas_languages import MasLanguages
from app.models.mas_organization import MasOrgUnit
from app.services import organization_service as org
from app.services import served_form_service
from app.services import xlsform_service as xf
from tests.base import BaseTestCase

REPO = Path(__file__).resolve().parents[2]
ND01 = REPO / "docs/kb/WHO_VA_2022_Docs/ND01_ICMRVA_WHOVA2022.xlsx"


def _sheet(data: bytes, name: str) -> list[dict]:
    """A downloaded sheet as header-keyed rows (blank cells absent)."""
    workbook = load_workbook(io.BytesIO(data), read_only=True)
    try:
        rows = list(workbook[name].iter_rows(values_only=True))
    finally:
        workbook.close()
    header = rows[0]
    return [{h: v for h, v in zip(header, row) if v is not None} for row in rows[1:]]


def _headers(data: bytes, name: str) -> list[str]:
    workbook = load_workbook(io.BytesIO(data), read_only=True)
    try:
        return [c for c in next(workbook[name].iter_rows(values_only=True)) if c]
    finally:
        workbook.close()


class XlsFormBase(BaseTestCase):
    FULL = "XLSF01"  # social autopsy, death summary, narration, geography; no medical records
    LEAN = "XLSF02"  # medical records only
    FORM_ID = "XLSF01_FORM"

    @classmethod
    def _project(cls, project_id, **fields):
        now = datetime.now(UTC)
        project = VaProjectMaster(
            project_id=project_id,
            project_code=project_id,
            project_name=f"Xlsform {project_id}",
            project_nickname=project_id,
            project_status=VaStatuses.active,
            project_registered_at=now,
            project_updated_at=now,
            **fields,
        )
        db.session.add(project)
        db.session.flush()
        return project

    @classmethod
    def _locale(cls, code, name, *, active, state):
        db.session.add(
            MasInstrumentLocales(
                instrument_code="WHO_2022_VA", locale_code=code, language_name=name,
                version=3, is_active=active, lifecycle_state=state, updated_at=datetime.now(UTC),
            )
        )
        db.session.flush()

    @classmethod
    def _text(cls, code, kind, key, field, text, source="imported"):
        db.session.add(
            MapInstrumentTranslations(
                instrument_code="WHO_2022_VA", locale_code=code, item_kind=kind,
                item_key=key, field=field, text=text, source=source,
            )
        )

    @classmethod
    def setUpClass(cls):
        super().setUpClass()
        for code, name in (("hindi", "Hindi"), ("english", "English")):
            if db.session.get(MasLanguages, code) is None:
                db.session.add(MasLanguages(language_code=code, language_name=name, is_active=True))
        cls._project(
            cls.FULL,
            project_structure_mode="organization",
            web_intake_medical_records_enabled=False,
            web_intake_narration_languages=["hindi", "english"],
        )
        cls._project(
            cls.LEAN,
            social_autopsy_enabled=False,
            web_intake_death_summary_enabled=False,
            web_intake_intake_note="",
        )

        # Locales: only hi is approved and active. kn is in review, ta a draft,
        # mr approved but switched off. Each has a Id10010 label so a leak shows.
        cls._locale("hi", "Hindi", active=True, state=LIFECYCLE_APPROVED)
        cls._locale("kn", "Kannada", active=False, state=LIFECYCLE_IN_REVIEW)
        cls._locale("ta", "Tamil", active=False, state=LIFECYCLE_DRAFT)
        cls._locale("mr", "Marathi", active=False, state=LIFECYCLE_APPROVED)
        for code in ("hi", "kn", "ta", "mr"):
            cls._text(code, "question", "Id10010", "label", f"LABEL-{code}")
        cls._text("hi", "question", "sa01", "label", "HI-SA01", source="edited")
        cls._text("hi", "choice", "sas01/1", "label", "HI-SAS01-1")
        cls._text("hi", "choice", "YES_NO/yes", "label", "HI-YES")  # a WHO choice, from the cached reference
        cls._text("hi", "question", "Id10017", "label", "MACHINE-DRAFT", source="machine")

        # Organization tree and the ODK form mapping.
        district = org.create_level(cls.FULL, level_code="district", level_name="District", depth=1)
        phc = org.create_level(cls.FULL, level_code="phc", level_name="PHC", depth=2)
        d01 = org.create_unit(cls.FULL, org_level_id=district.org_level_id, unit_code="D01", unit_name="District One")
        org.create_unit(
            cls.FULL, org_level_id=phc.org_level_id, unit_code="P01", unit_name="PHC One",
            parent_org_unit_id=d01.org_unit_id,
        )
        now = datetime.now(UTC)
        db.session.add(
            VaSiteMaster(site_id="XS01", site_name="Xlsform Site", site_abbr="XS01",
                         site_status=VaStatuses.active, site_registered_at=now, site_updated_at=now)
        )
        db.session.flush()
        db.session.add(
            VaProjectSites(project_id=cls.FULL, site_id="XS01", project_site_status=VaStatuses.active,
                           project_site_registered_at=now, project_site_updated_at=now)
        )
        db.session.add(
            MapProjectSiteOdk(project_id=cls.FULL, site_id="XS01", odk_project_id=1, odk_form_id=cls.FORM_ID)
        )
        db.session.add(
            VaProjectSites(project_id=cls.LEAN, site_id="XS01", project_site_status=VaStatuses.active,
                           project_site_registered_at=now, project_site_updated_at=now)
        )

        cls.pi_user = cls._get_or_make_user("xlsf.pi@test.local", "XlsfPi12345")
        db.session.add(
            VaUserAccessGrants(
                user_id=cls.pi_user.user_id, role=VaAccessRoles.project_pi,
                scope_type=VaAccessScopeTypes.project, project_id=cls.FULL,
                grant_status=VaStatuses.active, notes="xlsform test PI",
            )
        )
        db.session.commit()

    def _project_row(self, project_id):
        return db.session.get(VaProjectMaster, project_id)

    def _build(self, project_id, **kwargs):
        form, data = xf.build_project_xlsform(self._project_row(project_id), **kwargs)
        return form, data

    def _survey_names(self, data):
        return {row["name"] for row in _sheet(data, "survey") if "name" in row}


class XlsFormContentTests(XlsFormBase):
    def test_the_vendored_reference_is_the_kb_copy(self):
        kb = REPO / "docs/kb/WHO_VA_2022_Docs/2022whova_xls_form_for_odk_multilingual.xlsx"
        self.assertEqual(xf.REFERENCE_WORKBOOK.read_bytes(), kb.read_bytes())

    def test_survey_is_who_plus_exactly_the_enabled_extensions(self):
        form, data = self._build(self.FULL)
        names = self._survey_names(data)
        # Present first: the WHO base and every enabled extension.
        self.assertIn("Id10010", names)  # WHO
        self.assertIn("sa01", names)  # social_autopsy
        self.assertIn("ds_available", names)  # death_summary
        self.assertIn("narr_language", names)  # narration_language
        self.assertIn("consent_mode", names)  # digitva_core
        self.assertIn("dob_precision", names)  # doris_support_whova_2022
        self.assertIn("org_district_code", names)  # geography
        self.assertIn("intake_note", names)  # intake_screen
        self.assertEqual(
            form.extensions,
            ["digitva_core", "doris_support_whova_2022", "social_autopsy", "intake_screen",
             "death_summary", "geography", "narration_language"],
        )
        # Absent: medical records is switched off, ABHA is not configured.
        self.assertNotIn("md_available", names)
        self.assertNotIn("md_im1", names)
        self.assertNotIn("abha_number", names)

    def test_an_extension_that_is_off_is_absent_and_the_one_that_is_on_is_present(self):
        _form, data = self._build(self.LEAN)
        names = self._survey_names(data)
        self.assertIn("md_available", names)
        self.assertIn("md_im30", names)
        self.assertIn("consent_mode", names)
        for gone in ("sa01", "sa_tu13", "ds_available", "narr_language", "intake_note", "org_district_code"):
            self.assertNotIn(gone, names)

    def test_doris_changes_and_reference_overrides_apply(self):
        _form, data = self._build(self.FULL)
        rows = {r["name"]: r for r in _sheet(data, "survey") if "name" in r}
        self.assertEqual(rows["Id10308"]["required"], "yes")  # DORIS A9
        self.assertIn("Id10077_a", rows["Id10340"]["relevant"])  # DORIS A10
        self.assertEqual(rows["Id10366"]["constraint"], ". >= 100 and . <= 9999")  # DORIS A2
        self.assertEqual(rows["Id10304_a"]["relevant"], "selected(${Id10304},'yes')")  # digitva-13x
        self.assertEqual(rows["Id10230"]["agegroup"], "C_A")

    def test_added_rows_land_at_their_anchors(self):
        _form, data = self._build(self.FULL)
        names = [r["name"] for r in _sheet(data, "survey") if "name" in r]
        self.assertEqual(names[names.index("Id10013") + 1], "consent_mode")
        self.assertLess(names.index("Id10021"), names.index("dob_precision"))
        self.assertLess(names.index("Id10473"), names.index("custom_medical_certificate_upload"))
        self.assertLess(names.index("custom_medical_certificate_upload"), names.index("doris_autopsy_requested"))
        self.assertLess(names.index("intake_note"), names.index("Interviewer"))

    def test_only_approved_and_active_locales_get_columns(self):
        form, data = self._build(self.FULL)
        survey_headers = _headers(data, "survey")
        self.assertIn("label::English (en)", survey_headers)
        self.assertIn("label::Hindi (hi)", survey_headers)
        self.assertIn("hint::Hindi (hi)", survey_headers)
        self.assertIn("guidance_hint::Hindi (hi)", survey_headers)
        self.assertIn("constraint_message::Hindi (hi)", survey_headers)
        for other in ("Kannada", "Tamil", "Marathi", "French", "Arabic"):
            self.assertFalse([h for h in survey_headers if other in h], other)
        self.assertEqual([code for code, _name in form.languages], ["en", "hi"])

        rows = {r["name"]: r for r in _sheet(data, "survey") if "name" in r}
        self.assertEqual(rows["Id10010"]["label::Hindi (hi)"], "LABEL-hi")
        self.assertEqual(rows["sa01"]["label::Hindi (hi)"], "HI-SA01")
        # A machine draft never reaches the form.
        self.assertNotIn("label::Hindi (hi)", rows["Id10017"])
        choice = next(r for r in _sheet(data, "choices") if r.get("list_name") == "sas01" and r["name"] == "1")
        self.assertEqual(choice["label::Hindi (hi)"], "HI-SAS01-1")

    def test_a_project_limiting_its_languages_drops_the_rest(self):
        # Build with Hindi first: the reference rows are cached per process, so a
        # translation written into them would leak into the next project's form.
        _form, with_hindi = self._build(self.FULL)
        self.assertIn("label::Hindi (hi)", _headers(with_hindi, "choices"))
        yes = next(r for r in _sheet(with_hindi, "choices") if r.get("list_name") == "YES_NO" and r["name"] == "yes")
        self.assertEqual(yes["label::Hindi (hi)"], "HI-YES")
        project = self._project_row(self.FULL)
        project.web_intake_available_locales = ["en"]
        db.session.flush()
        _form, data = self._build(self.FULL)
        for sheet in ("survey", "choices"):
            self.assertFalse([h for h in _headers(data, sheet) if "(hi)" in h or "Hindi" in h], sheet)

    def test_language_lists_replace_the_placeholders(self):
        _form, data = self._build(self.FULL)
        choices = _sheet(data, "choices")
        interview = {r["name"]: r["label::English (en)"] for r in choices if r.get("list_name") == "language"}
        self.assertEqual(interview, {"en": "English", "hi": "Hindi"})
        narration = {r["name"] for r in choices if r.get("list_name") == "narr_language"}
        self.assertEqual(narration, {"hindi", "english"})
        site = [r for r in choices if r.get("list_name") == "site"]
        self.assertEqual([(r["name"], r["label::English (en)"]) for r in site], [("XS01", "Xlsform Site")])

    def test_org_unit_choices_match_the_choices_csv_rows(self):
        _form, data = self._build(self.FULL)
        choices = _sheet(data, "choices")
        expected = org.export_odk_choices_rows(self.FULL)
        self.assertEqual({(r["list_name"], r["name"]) for r in expected}, {("org_district", "D01"), ("org_phc", "P01")})
        for row in expected:
            found = next(
                c for c in choices if c.get("list_name") == row["list_name"] and c["name"] == row["name"]
            )
            self.assertEqual(found["label::English (en)"], row["label"])
            self.assertEqual(found.get("parent_code", ""), row["parent_code"])
        survey = {r["name"]: r for r in _sheet(data, "survey") if "name" in r}
        self.assertEqual(survey["org_phc_code"]["choice_filter"], "parent_code=${org_district_code}")
        self.assertNotIn("choice_filter", survey["org_district_code"])

    def test_the_extension_rows_agree_with_the_web_definition(self):
        """Names, relevance and constraints of what each extension adds are the
        web form's own (the composed definition), not a second opinion."""
        form, data = self._build(self.FULL)
        rows = {r["name"]: r for r in _sheet(data, "survey") if "name" in r}
        web = {
            q["name"]: q
            for q in json.loads(served_form_service.served_definition(form.extensions).body)["questions"]
        }

        def squash(value):
            return " ".join((value or "").split())

        blocks = xf._project_blocks(self._project_row(self.FULL), form.extensions)
        added = {
            row["name"]
            for block in blocks
            if block.extension not in ("geography", "intake_screen")
            for row in block.survey
            if row.get("name") and not row["type"].startswith("begin group")
        }
        compared = 0
        for name in sorted(added & web.keys()):
            with self.subTest(question=name):
                self.assertEqual(squash(rows[name].get("relevant")), squash(web[name].get("relevant", {}).get("source")))
                self.assertEqual(squash(rows[name].get("constraint")), squash(web[name].get("constraint", {}).get("source")))
                compared += 1
        self.assertGreaterEqual(compared, 60)  # social autopsy, documents, DORIS, consent mode
        # Every question an enabled conditional extension contributes is on the form.
        tagged = {
            q["name"]
            for q in served_form_service.composed_definition()["questions"]
            if set(q.get("extensions", ())) & set(form.extensions)
        }
        self.assertTrue(tagged)
        self.assertTrue(tagged <= set(rows), sorted(tagged - set(rows)))
        # The WHO rows the DORIS extension changes carry the web form's cells too.
        for name in ("Id10340", "Id10366"):
            self.assertEqual(
                squash(rows[name].get("relevant")), squash(web[name].get("relevant", {}).get("source"))
            )
            self.assertEqual(
                squash(rows[name].get("constraint")), squash(web[name].get("constraint", {}).get("source"))
            )

    def test_every_web_deviation_is_applied_to_the_odk_form(self):
        """The deviations file is one list: what the web instrument carries
        (constraints, messages, relevance, choices, section) the ODK form has."""
        form, data = self._build(self.FULL)
        rows = {r["name"]: r for r in _sheet(data, "survey") if "name" in r}
        choices = _sheet(data, "choices")
        web = {
            q["name"]: q
            for q in json.loads(served_form_service.served_definition(form.extensions).body)["questions"]
        }
        deviations = xf._deviations()
        self.assertIn("Id10365", deviations["questions"])  # present before absent
        checked = set()
        for name, change in deviations["questions"].items():
            question, row = web[name], rows[name]
            with self.subTest(question=name):
                self.assertEqual(row.get("constraint"), (question.get("constraint") or {}).get("source"))
                self.assertEqual(
                    row.get("constraint_message::English (en)"), question["constraintMessage"].get("en")
                )
                self.assertEqual(row.get("relevant"), (question.get("relevant") or {}).get("source"))
                self.assertEqual(row.get("agegroup"), question.get("ageGroup"))
                if "choices" in change and name != "language":
                    odk = [(c["name"], c["label::English (en)"]) for c in choices if c.get("list_name") == question["listName"]]
                    self.assertEqual(odk, [(c["value"], c["label"]["en"]) for c in question["choices"]])
                checked.add(name)
        self.assertEqual(checked, set(deviations["questions"]))
        # Dropped, not merely unchanged (owner decision 2026-10-06: the one WHO check not kept).
        self.assertNotIn("constraint", rows["Id10365"])
        self.assertNotIn("constraint_message::English (en)", rows["Id10365"])
        # The nmh note sits first in the injuries group, as in the web section.
        names = [r["name"] for r in _sheet(data, "survey") if "name" in r]
        self.assertEqual(web["nmh"]["sectionPath"][-1], "injuries_accidents")
        self.assertEqual(names[names.index("injuries_accidents") + 1], "nmh")
        # The section relabel.
        self.assertEqual(rows["consented"]["label::English (en)"], "Interview completion")
        # The language question's placeholders never reach ODK; the list is the project's.
        self.assertEqual({c["name"] for c in choices if c.get("list_name") == "language"}, {"en", "hi"})

    def test_pyxform_converts_the_generated_form_without_error(self):
        for project_id in (self.FULL, self.LEAN):
            with self.subTest(project=project_id):
                _form, data = self._build(project_id)
                warnings = []
                result = convert(io.BytesIO(data), file_type=".xlsx", warnings=warnings)
                self.assertIn("<h:html", result.xform)
                self.assertFalse([w for w in warnings if "error" in w.lower()], warnings)

    def test_form_id_is_stable_and_version_is_new_each_time(self):
        first = datetime(2026, 10, 6, 9, 15, 1, tzinfo=UTC)
        second = datetime(2026, 10, 6, 9, 15, 2, tzinfo=UTC)
        form_a, data_a = self._build(self.FULL, now=first)
        form_b, data_b = self._build(self.FULL, now=second)
        self.assertEqual(form_a.form_id, self.FORM_ID)
        self.assertEqual(form_a.form_id, form_b.form_id)
        self.assertEqual((form_a.version, form_b.version), ("20261006091501", "20261006091502"))
        settings = _sheet(data_a, "settings")[0]
        self.assertEqual(settings["form_id"], self.FORM_ID)
        self.assertEqual(settings["version"], "20261006091501")
        self.assertEqual(settings["default_language"], "English (en)")
        self.assertEqual(_sheet(data_b, "settings")[0]["version"], "20261006091502")

    def test_a_project_with_no_odk_mapping_gets_a_deterministic_form_id(self):
        self.assertEqual(self._build(self.LEAN)[0].form_id, "XLSF02_WHOVA2022")

    def test_control_characters_are_stripped_from_text_cells(self):
        form = xf.XlsForm(
            form_id="X", version="1", title="t", languages=[("en", "English")],
            survey=[{"type": "text", "name": "q", "label": "Bell\x07 ring"}],
            choices=[{"list_name": "l", "name": "a", "label": "Unit\x0b One"}],
        )
        data = xf.to_workbook_bytes(form)  # raises IllegalCharacterError without the strip
        self.assertEqual(_sheet(data, "survey")[0]["label::English (en)"], "Bell ring")
        self.assertEqual(_sheet(data, "choices")[0]["label::English (en)"], "Unit One")

    def test_values_that_look_like_formulas_stay_text(self):
        form = xf.XlsForm(
            form_id="X", version="1", title="t", languages=[("en", "English")],
            survey=[{"type": "text", "name": "q", "label": "=1+1"}],
            choices=[{"list_name": "l", "name": "a", "label": "=SUM(1,1)"}],
        )
        workbook = load_workbook(io.BytesIO(xf.to_workbook_bytes(form)))
        self.assertEqual(workbook["choices"].cell(2, 3).data_type, "s")
        self.assertEqual(workbook["choices"].cell(2, 3).value, "=SUM(1,1)")
        self.assertEqual(workbook["survey"]["F2"].data_type, "s")


class XlsFormRefusalTests(XlsFormBase):
    def test_an_enabled_extension_with_no_spec_is_refused_by_name(self):
        real = xf._spec_blocks

        def without_social_autopsy(extension):
            return None if extension == "social_autopsy" else real(extension)

        with patch.object(xf, "_spec_blocks", without_social_autopsy):
            with self.assertRaises(xf.XlsFormError) as caught:
                self._build(self.FULL)
        self.assertEqual(caught.exception.status_code, 422)
        self.assertIn("social_autopsy", str(caught.exception))

    def test_an_extension_unknown_to_the_generator_is_refused_by_name(self):
        patched = ("WHO_2022_VA", ["digitva_core", "future_extension"])
        with patch(
            "app.routes.api.organization.project_instrument_and_extensions", return_value=patched
        ):
            with self.assertRaises(xf.XlsFormError) as caught:
                self._build(self.FULL)
        self.assertEqual(caught.exception.status_code, 422)
        self.assertIn("future_extension", str(caught.exception))

    def test_a_non_who_instrument_is_refused(self):
        with patch(
            "app.routes.api.organization.project_instrument_and_extensions",
            return_value=("PHMRC_2020", ["digitva_core"]),
        ):
            with self.assertRaises(xf.XlsFormError) as caught:
                self._build(self.FULL)
        self.assertEqual(caught.exception.status_code, 409)

    def test_too_many_org_units_are_refused_before_they_are_loaded(self):
        with patch.object(xf, "MAX_ORG_CHOICES", 1), patch.object(
            org, "export_odk_choices_rows", side_effect=AssertionError("loaded before the cap")
        ):
            with self.assertRaises(xf.XlsFormError) as caught:
                self._build(self.FULL)
        self.assertEqual(caught.exception.status_code, 422)
        self.assertIn("2 active units", str(caught.exception))

    def test_a_select_with_no_choices_is_refused_naming_the_list(self):
        project = self._project(
            "XLSF03", social_autopsy_enabled=False, web_intake_death_summary_enabled=False,
            web_intake_medical_records_enabled=False, web_intake_intake_note="",
        )  # no project site: the Site list would be empty
        with self.assertRaises(xf.XlsFormError) as caught:
            xf.build_project_xlsform(project)
        self.assertEqual(caught.exception.status_code, 422)
        self.assertIn("site", str(caught.exception))

    def test_a_missing_anchor_fails_loudly_instead_of_dropping_rows(self):
        block = xf._Block("x", "x", ("after", "NoSuchRow"), [{"type": "text", "name": "q"}], [], [])
        with self.assertRaises(xf.XlsFormError):
            xf._splice([{"type": "text", "name": "a"}], [block])

    def test_a_choice_list_that_diverges_from_who_is_refused(self):
        block = xf._Block("x", "x", None, [], [{"list_name": "YES_NO", "name": "maybe"}], [])
        with self.assertRaises(xf.XlsFormError):
            xf._merge_choices([{"list_name": "YES_NO", "name": "yes"}], [block], set())


class XlsFormOrgLevelTests(XlsFormBase):
    """Empty organization levels (digitva-8695): an optional one is left out of
    the form, required ones are all named in one error. Levels filter on the
    nearest required level above them (digitva-ch1u)."""

    def _org_project(self, project_id, levels, units, *, with_site=True):
        """A lean organization-mode project; ``levels`` are
        ``(code, name, depth, optional)``, ``units`` ``(code, level_code,
        parent_code)``, parents first."""
        project = self._project(
            project_id, project_structure_mode="organization", social_autopsy_enabled=False,
            web_intake_death_summary_enabled=False, web_intake_medical_records_enabled=False,
            web_intake_intake_note="",
        )
        if with_site:
            now = datetime.now(UTC)
            db.session.add(
                VaProjectSites(project_id=project_id, site_id="XS01", project_site_status=VaStatuses.active,
                               project_site_registered_at=now, project_site_updated_at=now)
            )
        made = {
            code: org.create_level(project_id, level_code=code, level_name=name, depth=depth, is_optional=optional)
            for code, name, depth, optional in levels
        }
        ids = {}
        for code, level_code, parent in units:
            ids[code] = org.create_unit(
                project_id, org_level_id=made[level_code].org_level_id, unit_code=code,
                unit_name=f"Unit {code}", parent_org_unit_id=ids[parent] if parent else None,
            ).org_unit_id
        db.session.flush()
        return project

    def _converts(self, data):
        warnings = []
        result = convert(io.BytesIO(data), file_type=".xlsx", warnings=warnings)
        self.assertIn("<h:html", result.xform)

    def test_every_required_empty_level_is_named_in_one_error_in_order(self):
        project = self._org_project(
            "XLSF04",
            [
                ("district", "District", 1, False),
                ("chc", "Community Health Centre", 2, False),
                ("phc", "PHC / UPHC / AAM-PHC", 3, False),
                ("subcentre", "Sub-centre / AAM-SHC", 4, False),
            ],
            [("D01", "district", None)],
            with_site=False,
        )
        with self.assertRaises(xf.XlsFormError) as caught:
            xf.build_project_xlsform(project)
        message = str(caught.exception)
        self.assertEqual(caught.exception.status_code, 422)
        self.assertIn(
            "which have none: Community Health Centre (org_chc_code), PHC / UPHC / AAM-PHC (org_phc_code), "
            "Sub-centre / AAM-SHC (org_subcentre_code). Add units at those levels",
            message,
        )
        self.assertIn("org_chc_code", message)
        self.assertNotIn("org_district_code", message)  # it has a unit
        # A non-organization list is named in its own sentence, in the same error.
        self.assertIn("uses the choice list site, which has no choices for this project (add its sites first).", message)

    def test_levels_filter_on_the_nearest_required_level_and_choices_on_that_ancestor(self):
        """digitva-ch1u: a CHC under the optional SDH and a CHC directly under
        the DH are both reachable from the district answer."""
        project = self._org_project(
            "XLSF08",
            [("district", "District", 1, False), ("taluka", "SDH", 2, True), ("chc", "CHC", 3, False)],
            [("D01", "district", None), ("T01", "taluka", "D01"), ("C01", "chc", "T01"), ("C02", "chc", "D01")],
        )
        _form, data = xf.build_project_xlsform(project)
        survey = {r["name"]: r for r in _sheet(data, "survey") if "name" in r}
        self.assertIn("org_taluka_code", survey)
        self.assertEqual(survey["org_taluka_code"]["choice_filter"], "parent_code=${org_district_code}")
        self.assertEqual(survey["org_chc_code"]["choice_filter"], "parent_code=${org_district_code}")
        self.assertNotIn("choice_filter", survey["org_district_code"])
        parents = {
            r["name"]: r.get("parent_code", "") for r in _sheet(data, "choices")
            if r.get("list_name") in ("org_district", "org_taluka", "org_chc")
        }
        self.assertEqual(parents, {"D01": "", "T01": "D01", "C01": "D01", "C02": "D01"})
        # The shared choices CSV keeps the direct parent for forms authored elsewhere.
        direct = {r["name"]: r["parent_code"] for r in org.export_odk_choices_rows("XLSF08")}
        self.assertEqual(direct["C01"], "T01")
        self.assertEqual(direct["C02"], "D01")
        self._converts(data)

    def test_a_unit_under_an_inactive_optional_parent_attaches_to_its_required_ancestor(self):
        project = self._org_project(
            "XLSF09",
            [("district", "District", 1, False), ("taluka", "SDH", 2, True), ("chc", "CHC", 3, False)],
            [("D01", "district", None), ("T01", "taluka", "D01"), ("T02", "taluka", "D01"), ("C01", "chc", "T01")],
        )
        db.session.scalar(sa.select(MasOrgUnit).where(MasOrgUnit.unit_code == "T01")).is_active = False
        db.session.flush()
        _form, data = xf.build_project_xlsform(project)
        parents = {
            r["name"]: r.get("parent_code", "") for r in _sheet(data, "choices")
            if r.get("list_name") in ("org_district", "org_taluka", "org_chc")
        }
        self.assertEqual(parents["C01"], "D01")
        self.assertIn("T02", parents)
        self.assertNotIn("T01", parents)
        self._converts(data)

    def test_an_optional_empty_middle_level_is_left_out(self):
        project = self._org_project(
            "XLSF05",
            [("district", "District", 1, False), ("taluka", "Taluka", 2, True), ("chc", "CHC", 3, False)],
            [("D01", "district", None), ("C01", "chc", "D01")],
        )
        _form, data = xf.build_project_xlsform(project)
        survey = {r["name"]: r for r in _sheet(data, "survey") if "name" in r}
        self.assertIn("org_chc_code", survey)
        self.assertIn("org_district_code", survey)
        self.assertNotIn("org_taluka_code", survey)
        self.assertEqual(survey["org_chc_code"]["choice_filter"], "parent_code=${org_district_code}")
        self._converts(data)

    def test_an_optional_empty_village_is_left_out_and_the_form_builds(self):
        project = self._org_project(
            "XLSF06",
            [("district", "District", 1, False), ("subcentre", "Sub-centre", 2, False), ("village", "Village", 3, True)],
            [("D01", "district", None), ("S01", "subcentre", "D01")],
        )
        _form, data = xf.build_project_xlsform(project)
        names = self._survey_names(data)
        self.assertIn("org_subcentre_code", names)
        self.assertNotIn("org_village_code", names)
        self._converts(data)

    def test_an_optional_level_with_units_stays_in_the_form(self):
        project = self._org_project(
            "XLSF07",
            [("district", "District", 1, False), ("subcentre", "Sub-centre", 2, False), ("village", "Village", 3, True)],
            [("D01", "district", None), ("S01", "subcentre", "D01"), ("V01", "village", "S01")],
        )
        _form, data = xf.build_project_xlsform(project)
        survey = {r["name"]: r for r in _sheet(data, "survey") if "name" in r}
        self.assertIn("org_village_code", survey)
        self.assertNotIn("required", survey["org_village_code"])
        self.assertEqual(survey["org_village_code"]["choice_filter"], "parent_code=${org_subcentre_code}")
        village = [r for r in _sheet(data, "choices") if r.get("list_name") == "org_village"]
        self.assertEqual([(r["name"], r["parent_code"]) for r in village], [("V01", "S01")])
        self._converts(data)


class XlsFormRouteTests(XlsFormBase):
    def _url(self, project_id=None, query=""):
        return f"/admin/api/projects/{project_id or self.FULL}/odk-xlsform.xlsx{query}"

    def test_admin_downloads_the_workbook(self):
        self._login(self.base_admin_id)
        response = self.client.get(self._url())
        self.assertEqual(response.status_code, 200)
        self.assertIn("spreadsheetml", response.mimetype)
        self.assertIn("attachment", response.headers["Content-Disposition"])
        self.assertIn(self.FORM_ID, response.headers["Content-Disposition"])
        self.assertEqual(response.headers["Cache-Control"], "no-store")
        workbook = load_workbook(io.BytesIO(response.data), read_only=True)
        self.assertEqual(workbook.sheetnames, ["survey", "choices", "settings"])
        workbook.close()
        self.assertIn("sa01", self._survey_names(response.data))

    def test_the_project_pi_downloads_it(self):
        self._login(str(self.pi_user.user_id))
        self.assertEqual(self.client.get(self._url()).status_code, 200)

    def test_two_downloads_share_form_id_and_differ_in_version(self):
        stamps = [datetime(2026, 10, 6, 10, 0, 0, tzinfo=UTC), datetime(2026, 10, 6, 10, 0, 7, tzinfo=UTC)]
        self._login(self.base_admin_id)
        with patch.object(xf, "datetime") as clock:
            clock.now.side_effect = stamps
            settings = [
                _sheet(self.client.get(self._url()).data, "settings")[0] for _ in stamps
            ]
        self.assertEqual(settings[0]["form_id"], settings[1]["form_id"])
        self.assertNotEqual(settings[0]["version"], settings[1]["version"])

    def test_the_pi_of_another_project_is_refused(self):
        self._login(self.base_project_pi_id)  # PI of BASE01, not XLSF01
        self.assertEqual(self.client.get(self._url()).status_code, 403)

    def test_a_plain_coder_is_refused(self):
        self._login(self.base_coder_id)
        self.assertEqual(self.client.get(self._url()).status_code, 403)

    def test_an_unknown_project_is_404_for_an_admin(self):
        self._login(self.base_admin_id)
        self.assertEqual(self.client.get(self._url("NOPE01")).status_code, 404)

    def test_an_extension_without_a_spec_is_a_422_naming_it(self):
        real = xf._spec_blocks
        self._login(self.base_admin_id)
        with patch.object(xf, "_spec_blocks", lambda e: None if e == "death_summary" else real(e)):
            response = self.client.get(self._url())
        self.assertEqual(response.status_code, 422)
        self.assertIn("death_summary", response.get_json()["error"])

    def test_several_mapped_forms_need_a_choice_and_the_choice_is_checked(self):
        db.session.add(
            MapProjectSiteOdk(project_id=self.FULL, site_id="XS01", odk_project_id=2, odk_form_id="XLSF01_OTHER")
        )
        db.session.flush()
        self._login(self.base_admin_id)
        response = self.client.get(self._url())
        self.assertEqual(response.status_code, 409)
        self.assertEqual(response.get_json()["forms"], [self.FORM_ID, "XLSF01_OTHER"])
        chosen = self.client.get(self._url(query="?form_id=XLSF01_OTHER"))
        self.assertEqual(chosen.status_code, 200)
        self.assertEqual(_sheet(chosen.data, "settings")[0]["form_id"], "XLSF01_OTHER")
        # An id that is not one of this project's mappings is never used.
        self.assertEqual(self.client.get(self._url(query="?form_id=SOMEONE_ELSES")).status_code, 404)


class XlsFormDiffTests(XlsFormBase):
    def test_diff_lists_rows_each_way_and_changes_without_merging(self):
        form, _data = self._build(self.FULL)
        report = xf.diff_against_workbook(form, ND01)
        # ND01 has its own site-local rows; the generated form has the layers ND01 lacks.
        self.assertIn("site_individual_id", report["survey_only_in_workbook"])
        self.assertIn("survey_district", report["survey_only_in_workbook"])
        self.assertIn("consent_mode", report["survey_only_in_generated"])
        self.assertIn("org_district_code", report["survey_only_in_generated"])
        self.assertIn("unique_id", {c["name"] for c in report["survey_changed"]})
        self.assertIn(("district", "088"), report["choices_only_in_workbook"])
        self.assertIn(("org_district", "D01"), report["choices_only_in_generated"])

    def test_cli_prints_the_report(self):
        runner = self.app.test_cli_runner()
        result = runner.invoke(
            args=["xlsform", "diff", "--workbook", str(ND01), "--project", self.FULL]
        )
        self.assertEqual(result.exit_code, 0, result.output)
        self.assertIn("Survey rows only in the workbook", result.output)
        self.assertIn("site_individual_id", result.output)
        self.assertIn("consent_mode", result.output)

    def test_cli_refuses_an_unknown_project(self):
        result = self.app.test_cli_runner().invoke(
            args=["xlsform", "diff", "--workbook", str(ND01), "--project", "NOPE01"]
        )
        self.assertNotEqual(result.exit_code, 0)
        self.assertIn("not found", result.output)
