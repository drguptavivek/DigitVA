"""HTTP contract of app/routes/admin_organization.py: authorization and thin dispatch."""
import io
from datetime import UTC, datetime
from unittest.mock import patch

import sqlalchemy as sa
from openpyxl import load_workbook

from app import db
from app.models import (
    VaAccessRoles,
    VaAccessScopeTypes,
    VaProjectMaster,
    VaStatuses,
    VaUserAccessGrants,
    VaUsers,
)
from app.models.mas_languages import MasLanguages
from app.models.mas_organization import MasOrgLevel, MasOrgUnitWorker
from app.services import organization_service as org
from tests.base import BaseTestCase


class AdminOrganizationApiTests(BaseTestCase):
    PROJECT = "ORGA01"
    SITES_PROJECT = "ORGS01"

    @classmethod
    def setUpClass(cls):
        super().setUpClass()
        now = datetime.now(UTC)
        db.session.add(
            VaProjectMaster(
                project_id=cls.PROJECT,
                project_code=cls.PROJECT,
                project_name="Org API Project",
                project_nickname="OrgApi",
                project_status=VaStatuses.active,
                project_registered_at=now,
                project_updated_at=now,
                project_structure_mode="organization",
            )
        )
        db.session.add(
            VaProjectMaster(
                project_id=cls.SITES_PROJECT,
                project_code=cls.SITES_PROJECT,
                project_name="Sites Project",
                project_nickname="Sites",
                project_status=VaStatuses.active,
                project_registered_at=now,
                project_updated_at=now,
            )
        )
        if db.session.get(MasLanguages, "en") is None:
            db.session.add(MasLanguages(language_code="en", language_name="English", is_active=True))
        db.session.commit()

    def _url(self, suffix=""):
        return f"/admin/api/organization/{self.PROJECT}{suffix}"

    def _import_users(self, text, dry_run="1"):
        return self.client.post(
            self._url("/project-users/import"),
            data={"file": (io.BytesIO(text.encode("utf-8")), "users.csv"), "dry_run": dry_run},
            content_type="multipart/form-data", headers=self._csrf_headers(),
        )

    def test_blank_csv_templates(self):
        self._login(str(self.base_admin_id))
        units = self.client.get(self._url("/templates/units.csv"))
        users = self.client.get(self._url("/templates/project-users.csv"))
        self.assertEqual(units.status_code, 200)
        self.assertEqual(users.status_code, 200)
        self.assertTrue(units.get_data(as_text=True).startswith("unit_code,unit_name,level_code,parent_code"))
        self.assertNotIn(",path,", units.get_data(as_text=True))
        self.assertEqual(users.get_data(as_text=True).strip(),
                         "email,name,role,org_unit_code,cadre_code,language_codes,phone")

    def test_blank_xlsx_templates(self):
        self._login(str(self.base_admin_id))
        for suffix, first in (("/templates/units.xlsx", "unit_code"),
                              ("/templates/project-users.xlsx", "email")):
            response = self.client.get(self._url(suffix))
            self.assertEqual(response.status_code, 200)
            workbook = load_workbook(io.BytesIO(response.data), read_only=True)
            self.assertEqual(workbook.active.cell(1, 1).value, first)
            workbook.close()

    def test_project_users_xlsx_has_project_aware_dependent_lists(self):
        self._login(str(self.base_admin_id))
        org.seed_default_organization(self.PROJECT)
        levels = {level.level_code: level for level in org.list_levels(self.PROJECT)}
        district = org.create_unit(
            self.PROJECT,
            org_level_id=levels["district"].org_level_id,
            unit_code="9DIST",
            unit_name="District 9",
        )
        org.create_unit(
            self.PROJECT,
            org_level_id=levels["chc"].org_level_id,
            parent_org_unit_id=district.org_unit_id,
            unit_code="_CHC9",
            unit_name="CHC 9",
        )
        inactive = org.create_unit(
            self.PROJECT,
            org_level_id=levels["district"].org_level_id,
            unit_code="INACTIVE",
            unit_name="Inactive District",
        )
        org.set_unit_active(self.PROJECT, inactive.org_unit_id, False)
        response = self.client.get(self._url("/templates/project-users.xlsx"))
        self.assertEqual(response.status_code, 200)
        workbook = load_workbook(io.BytesIO(response.data))
        self.assertEqual(workbook.sheetnames, ["project_users", "Instructions", "Reference"])
        self.assertEqual(tuple(workbook["project_users"].iter_rows(min_row=1, max_row=1, values_only=True))[0],
                         ("email", "name", "role", "org_unit_code", "cadre_code", "language_codes", "phone"))
        self.assertTrue(all(value is None for row in workbook["project_users"].iter_rows(min_row=2, values_only=True)
                            for value in row))
        self.assertEqual(workbook["project_users"].freeze_panes, "A2")
        self.assertEqual(workbook["Reference"].sheet_state, "hidden")
        self.assertEqual(workbook.defined_names["UnitCodes"].attr_text, "'Reference'!$B$2:$B$3")
        self.assertEqual(workbook.defined_names["UnitLevelKeys"].attr_text, "'Reference'!$C$2:$C$3")
        level_key_by_id = {
            level.org_level_id: f"L{index}"
            for index, level in enumerate(sorted(levels.values(), key=lambda level: level.depth), 1)
        }
        reference = workbook["Reference"]
        reference_units = {
            reference.cell(row, 2).value: reference.cell(row, 3).value
            for row in range(2, 4)
        }
        self.assertEqual(reference_units, {
            "9DIST": level_key_by_id[district.org_level_id],
            "_CHC9": level_key_by_id[levels["chc"].org_level_id],
        })
        for unit_code, expected_cadre in (("9DIST", "MO"), ("_CHC9", "SMO")):
            level_key = reference_units[unit_code]
            cadre_column = next(cell.column for cell in reference[1] if cell.value == level_key)
            cadres_at_level = {
                reference.cell(row, cadre_column).value
                for row in range(2, reference.max_row + 1)
                if reference.cell(row, cadre_column).value
            }
            self.assertIn(expected_cadre, cadres_at_level)
            self.assertIn(f"Cadres_{level_key}", workbook.defined_names)
        validations = {str(validation.sqref): validation for validation in workbook["project_users"].data_validations.dataValidation}
        self.assertIn("C2:C1001", validations)
        self.assertIn("D2:D1001", validations)
        self.assertIn("E2:E1001", validations)
        self.assertNotIn("Reference!", validations["E2:E1001"].formula1)
        self.assertIn("UnitLevelKeys", validations["E2:E1001"].formula1)
        self.assertIn("UnitCodes", validations["E2:E1001"].formula1)
        self.assertIn("MATCH($D2", validations["E2:E1001"].formula1)
        self.assertTrue(all(validation.showErrorMessage for validation in validations.values()))
        self.assertTrue(all(cell.font.name == "Arial" for row in workbook["project_users"].iter_rows(min_row=1, max_row=1) for cell in row))
        self.assertIn("en;hi", workbook["Instructions"]["B6"].value)
        reference_codes = [workbook["Reference"].cell(row, 2).value for row in range(2, 4)]
        self.assertEqual(reference_codes, ["9DIST", "_CHC9"])
        for column, value in enumerate(
            ["template-user@example.org", "Template User", "reviewer", "9DIST", "MO", "en", ""],
            1,
        ):
            workbook["project_users"].cell(2, column, value)
        content = io.BytesIO()
        workbook.save(content)
        preview = self.client.post(
            self._url("/project-users/import"),
            data={"file": (io.BytesIO(content.getvalue()), "users.xlsx")},
            content_type="multipart/form-data", headers=self._csrf_headers(),
        )
        self.assertEqual(preview.status_code, 200, preview.get_json())
        with patch("app.services.email_service.is_mail_configured", return_value=False):
            applied = self.client.post(
                self._url("/project-users/import"),
                data={"file": (io.BytesIO(content.getvalue()), "users.xlsx"), "dry_run": "0"},
                content_type="multipart/form-data", headers=self._csrf_headers(),
            )
        self.assertEqual(applied.status_code, 200, applied.get_json())
        self.assertEqual(applied.get_json()["created_users"], 1)
        workbook.close()

    def test_excel_csv_and_xlsx_user_import_preview(self):
        self._login(str(self.base_admin_id))
        csv_body = ("sep=;\r\n EMAIL ; NAME ; ROLE ; ORG_UNIT_CODE ; CADRE_CODE ; LANGUAGE_CODES ; PHONE ;\r\n"
                    "excel@example.org;Café;reviewer;;;en;;\r\n").encode("cp1252")
        response = self.client.post(
            self._url("/project-users/import"),
            data={"file": (io.BytesIO(csv_body), "users.csv")},
            content_type="multipart/form-data", headers=self._csrf_headers(),
        )
        self.assertEqual(response.status_code, 200, response.get_json())
        self.assertEqual(response.get_json()["rows"][0]["email"], "excel@example.org")

        template = self.client.get(self._url("/templates/project-users.xlsx"))
        workbook = load_workbook(io.BytesIO(template.data))
        workbook.active.append(["xlsx@example.org", "Workbook", "reviewer", "", "", "en", ""])
        content = io.BytesIO()
        workbook.save(content)
        response = self.client.post(
            self._url("/project-users/import"),
            data={"file": (io.BytesIO(content.getvalue()), "users.xlsx")},
            content_type="multipart/form-data", headers=self._csrf_headers(),
        )
        self.assertEqual(response.status_code, 200, response.get_json())
        self.assertEqual(response.get_json()["rows"][0]["email"], "xlsx@example.org")
        with patch("app.services.email_service.is_mail_configured", return_value=False):
            applied = self.client.post(
                self._url("/project-users/import"),
                data={"file": (io.BytesIO(content.getvalue()), "users.xlsx"), "dry_run": "0"},
                content_type="multipart/form-data", headers=self._csrf_headers(),
            )
        self.assertEqual(applied.status_code, 200, applied.get_json())
        self.assertEqual(applied.get_json()["created_users"], 1)
        self.assertIsNotNone(db.session.scalar(
            sa.select(VaUsers).where(VaUsers.email == "xlsx@example.org")
        ))

    def test_invalid_role_error_lists_grantable_roles(self):
        self._login(str(self.base_admin_id))
        response = self._import_users(
            "email,name,role,org_unit_code,cadre_code,language_codes,phone\n"
            "invalid-role@example.org,Name,bogus,,,en,\n"
        )
        self.assertEqual(response.status_code, 400)
        self.assertIn("role must be one of:", response.get_json()["error"])
        self.assertNotIn("VaAccessRoles", response.get_json()["error"])

    def test_project_pi_cannot_see_inactive_account_status_or_grant_pi(self):
        db.session.add(VaUserAccessGrants(
            user_id=self.base_project_pi_user.user_id,
            role=VaAccessRoles.project_pi,
            scope_type=VaAccessScopeTypes.project,
            project_id=self.PROJECT,
            grant_status=VaStatuses.active,
        ))
        inactive = VaUsers(email="inactive-import@example.org", name="Inactive Import",
                           user_status=VaStatuses.deactive, permission={},
                           landing_page="coder", pw_reset_t_and_c=False,
                           email_verified=True, vacode_language=["en"])
        inactive.set_password("UnusedPassword123!")
        db.session.add(inactive)
        db.session.commit()
        self._login(str(self.base_project_pi_id))
        header = "email,name,role,org_unit_code,cadre_code,language_codes,phone\n"
        unknown = self._import_users(header + "missing-import@example.org,,reviewer,,,,\n")
        disabled = self._import_users(header + "inactive-import@example.org,,reviewer,,,,\n")
        self.assertEqual(unknown.status_code, 400)
        self.assertEqual(disabled.status_code, 400)
        self.assertEqual(unknown.get_json()["error"].replace("missing-import@example.org", "EMAIL"),
                         disabled.get_json()["error"].replace("inactive-import@example.org", "EMAIL"))
        denied = self._import_users(header + "base.coder@test.local,,project_pi,,,,\n")
        self.assertEqual(denied.status_code, 400)
        self.assertIn("role must be one of:", denied.get_json()["error"])

    def test_blank_cadre_on_rerun_preserves_existing_grant_cadre(self):
        self._login(str(self.base_admin_id))
        org.seed_default_organization(self.PROJECT)
        levels = {level.level_code: level for level in org.list_levels(self.PROJECT)}
        district = org.create_unit(self.PROJECT, org_level_id=levels["district"].org_level_id,
                                   unit_code="D90", unit_name="District 90")
        chc = org.create_unit(self.PROJECT, org_level_id=levels["chc"].org_level_id,
                              unit_code="C90", unit_name="CHC 90",
                              parent_org_unit_id=district.org_unit_id)
        phc = org.create_unit(self.PROJECT, org_level_id=levels["phc"].org_level_id,
                              unit_code="P90", unit_name="PHC 90",
                              parent_org_unit_id=chc.org_unit_id)
        cadre = next(c for c in org.list_cadres(self.PROJECT) if c.cadre_code == "MO")
        grant = VaUserAccessGrants(
            user_id=self.base_coder_id, role=VaAccessRoles.reviewer,
            scope_type=VaAccessScopeTypes.org_unit, org_unit_id=phc.org_unit_id,
            cadre_id=cadre.cadre_id, grant_status=VaStatuses.active,
        )
        db.session.add(grant)
        db.session.commit()
        response = self._import_users(
            "email,name,role,org_unit_code,cadre_code,language_codes,phone\n"
            "base.coder@test.local,,reviewer,P90,,,\n", dry_run="0",
        )
        self.assertEqual(response.status_code, 200, response.get_json())
        self.assertEqual(response.get_json()["changed_grants"], 0)
        db.session.refresh(grant)
        self.assertEqual(grant.cadre_id, cadre.cadre_id)

    def test_other_project_pi_cannot_import_or_download_templates(self):
        self._login(str(self.base_project_pi_id))
        for suffix in ("/templates/units.csv", "/templates/units.xlsx",
                       "/templates/project-users.csv", "/templates/project-users.xlsx"):
            self.assertEqual(self.client.get(self._url(suffix)).status_code, 403)
        response = self._import_users(
            "email,name,role,org_unit_code,cadre_code,language_codes,phone\n"
            "someone@example.org,,reviewer,,,,\n"
        )
        self.assertEqual(response.status_code, 403)
        units = self.client.post(
            self._url("/import"),
            data={"file": (io.BytesIO(b"unit_code,unit_name,level_code\nP01,PHC,phc\n"),
                           "units.csv"), "sheet": "units", "dry_run": "1"},
            content_type="multipart/form-data", headers=self._csrf_headers(),
        )
        self.assertEqual(units.status_code, 403)

    def test_invalid_row_leaves_no_account_or_grant(self):
        self._login(str(self.base_admin_id))
        header = "email,name,role,org_unit_code,cadre_code,language_codes,phone\n"
        response = self._import_users(header +
            "valid@example.org,Valid,coder,,,en,\n" +
            "bad@example.org,Bad,admin,,,en,\n", dry_run="0")
        self.assertEqual(response.status_code, 400, response.get_json())
        self.assertIsNone(db.session.scalar(sa.select(VaUsers).where(VaUsers.email == "valid@example.org")))
        self.assertEqual(db.session.scalar(sa.select(sa.func.count()).select_from(VaUserAccessGrants)
                          .where(VaUserAccessGrants.project_id == self.PROJECT)), 0)

    def test_blank_line_keeps_physical_row_number(self):
        self._login(str(self.base_admin_id))
        response = self._import_users(
            "email,name,role,org_unit_code,cadre_code,language_codes,phone\n"
            "\n"
            "bad@example.org,Bad,admin,,,en,\n"
        )
        self.assertEqual(response.status_code, 400)
        self.assertIn("Row 3", response.get_json()["error"])

    def test_conflicting_new_user_profiles_reject_whole_file(self):
        self._login(str(self.base_admin_id))
        response = self._import_users(
            "email,name,role,org_unit_code,cadre_code,language_codes,phone\n"
            "conflict@example.org,First,reviewer,,,en,\n"
            "conflict@example.org,Second,data_manager,,,en,\n",
            dry_run="0",
        )
        self.assertEqual(response.status_code, 400)
        self.assertIn("profile differs", response.get_json()["error"])
        self.assertIsNone(db.session.scalar(sa.select(VaUsers).where(VaUsers.email == "conflict@example.org")))

    def test_admin_import_creates_one_account_and_rerun_retains_grant(self):
        self._login(str(self.base_admin_id))
        text = ("email,name,role,org_unit_code,cadre_code,language_codes,phone\n"
                "bulk@example.org,Bulk User,reviewer,,,en,\n")
        preview = self._import_users(text)
        self.assertEqual(preview.status_code, 200, preview.get_json())
        self.assertEqual(preview.get_json()["rows"][0]["scope"], "whole project")
        self.assertIsNone(db.session.scalar(sa.select(VaUsers).where(VaUsers.email == "bulk@example.org")))
        with patch("app.services.email_service.is_mail_configured", return_value=True), patch(
            "app.services.email_service.send_verification_email", return_value=True
        ) as verification, patch(
            "app.services.email_service.send_password_reset_email", return_value=True
        ) as password:
            first = self._import_users(text, dry_run="0")
        self.assertEqual(first.status_code, 200, first.get_json())
        self.assertEqual(first.get_json()["created_users"], 1)
        self.assertEqual(first.get_json()["invitations_queued"], 1)
        verification.assert_called_once()
        # digitva-kmoy: no reset link at creation; verifying emails the password.
        password.assert_not_called()
        second = self._import_users(text, dry_run="0")
        self.assertEqual(second.status_code, 200, second.get_json())
        self.assertEqual(second.get_json()["changed_grants"], 0)

    def test_disabled_email_reports_skipped_invitation(self):
        self._login(str(self.base_admin_id))
        text = ("email,name,role,org_unit_code,cadre_code,language_codes,phone\n"
                "skipped@example.org,Skipped,reviewer,,,en,\n")
        with patch("app.services.email_service.is_mail_configured", return_value=True), patch(
            "app.services.email_service._email_delivery_enabled", return_value=False
        ):
            response = self._import_users(text, dry_run="0")
        self.assertEqual(response.status_code, 200, response.get_json())
        self.assertEqual(response.get_json()["invitations_queued"], 0)
        self.assertTrue(any("skipped@example.org" in warning
                            for warning in response.get_json()["invite_warnings"]))

    def test_unconfigured_email_does_not_queue_invitations(self):
        self._login(str(self.base_admin_id))
        text = ("email,name,role,org_unit_code,cadre_code,language_codes,phone\n"
                "unconfigured@example.org,Unconfigured,reviewer,,,en,\n")
        with patch("app.services.email_service.is_mail_configured", return_value=False), patch(
            "app.services.email_service.send_verification_email"
        ) as verification, patch(
            "app.services.email_service.send_password_reset_email"
        ) as password:
            response = self._import_users(text, dry_run="0")
        self.assertEqual(response.status_code, 200, response.get_json())
        self.assertEqual(response.get_json()["invitations_queued"], 0)
        self.assertTrue(any("not configured" in warning
                            for warning in response.get_json()["invite_warnings"]))
        verification.assert_not_called()
        password.assert_not_called()

    def test_project_pi_can_attach_existing_user_but_cannot_create_account(self):
        db.session.add(VaUserAccessGrants(
            user_id=self.base_project_pi_user.user_id,
            role=VaAccessRoles.project_pi,
            scope_type=VaAccessScopeTypes.project,
            project_id=self.PROJECT,
            grant_status=VaStatuses.active,
        ))
        db.session.commit()
        self._login(str(self.base_project_pi_id))
        existing = ("email,name,role,org_unit_code,cadre_code,language_codes,phone\n"
                    "base.coder@test.local,,reviewer,,,,\n")
        response = self._import_users(existing, dry_run="0")
        self.assertEqual(response.status_code, 200, response.get_json())
        self.assertEqual(response.get_json()["created_users"], 0)
        new = ("email,name,role,org_unit_code,cadre_code,language_codes,phone\n"
               "new-pi@example.org,New,reviewer,,,en,\n")
        denied = self._import_users(new, dry_run="0")
        self.assertEqual(denied.status_code, 400, denied.get_json())
        self.assertIsNone(db.session.scalar(sa.select(VaUsers).where(VaUsers.email == "new-pi@example.org")))

    def test_panel_renders_for_admin(self):
        self._login(str(self.base_admin_id))
        response = self.client.get(f"/admin/panels/organization?project_id={self.PROJECT}")
        self.assertEqual(response.status_code, 200)
        self.assertIn(b'id="panel-organization"', response.data)
        html = response.get_data(as_text=True)
        self.assertIn('id="org-reference-model"', html)
        self.assertIn("<summary", html)  # collapsed without Bootstrap JS
        self.assertIn("District reference model", html)
        self.assertIn("CS Civil Surgeon", html)
        # The in-charge cadres suggest site_pi, shown as In-charge (held at a unit)
        self.assertIn('href="/help/user-roles#role-site_pi">In-charge</a>', html)
        self.assertNotIn('role-site_pi">site_pi</a>', html)
        self.assertNotIn('href="/help/user-roles#role-interview_supervisor"', html)
        self.assertIn('href="/help/user-roles#role-death_reporter"', html)
        self.assertIn('href="/help/user-roles">What each role can do</a>', html)
        self.assertIn("Populate district defaults", html)
        self.assertNotIn(">Seed template<", html)

    def test_coder_is_forbidden(self):
        self._login(str(self.base_coder_id))
        response = self.client.get(self._url())
        self.assertIn(response.status_code, (302, 403))

    def test_project_pi_of_other_project_is_forbidden(self):
        self._login(str(self.base_project_pi_id))
        response = self.client.get(self._url())
        self.assertEqual(response.status_code, 403)
        self.assertEqual(response.get_json()["error"], "You do not have access to that project.")

    def test_project_pi_manages_own_project(self):
        from app.models import VaAccessRoles, VaAccessScopeTypes, VaUserAccessGrants

        db.session.add(
            VaUserAccessGrants(
                user_id=self.base_project_pi_id,
                role=VaAccessRoles.project_pi,
                scope_type=VaAccessScopeTypes.project,
                project_id=self.PROJECT,
                grant_status=VaStatuses.active,
            )
        )
        db.session.commit()
        self._login(str(self.base_project_pi_id))
        headers = self._csrf_headers()
        self.assertEqual(self.client.post(self._url("/seed-template"), json={}, headers=headers).status_code, 200)
        level = self.client.get(self._url("/levels")).get_json()["levels"][0]
        created = self.client.post(
            self._url("/units"),
            json={"org_level_id": level["org_level_id"], "unit_code": "D01", "unit_name": "District One"},
            headers=headers,
        )
        self.assertEqual(created.status_code, 201, created.get_json())
        missing_flag = self.client.post(self._url(f"/units/{created.get_json()['unit']['org_unit_id']}/toggle"), json={}, headers=headers)
        self.assertEqual(missing_flag.status_code, 400)

    def test_admin_builds_tree_through_api(self):
        self._login(str(self.base_admin_id))
        headers = self._csrf_headers()

        seeded = self.client.post(self._url("/seed-template"), json={}, headers=headers)
        self.assertEqual(seeded.status_code, 200)
        self.assertEqual(seeded.get_json()["seeded"]["levels"], 7)

        levels = {lv["level_code"]: lv for lv in self.client.get(self._url("/levels")).get_json()["levels"]}
        created = self.client.post(
            self._url("/units"),
            json={"org_level_id": levels["district"]["org_level_id"], "unit_code": "d01", "unit_name": "District One", "google_maps_url": "https://maps.app.goo.gl/abc"},
            headers=headers,
        )
        self.assertEqual(created.status_code, 201, created.get_json())
        district = created.get_json()["unit"]
        self.assertEqual(district["unit_code"], "D01")

        bad = self.client.post(
            self._url("/units"),
            json={"org_level_id": levels["phc"]["org_level_id"], "parent_org_unit_id": district["org_unit_id"], "unit_code": "P01", "unit_name": "Skips CHC"},
            headers=headers,
        )
        self.assertEqual(bad.status_code, 400)
        self.assertIn("'chc'", bad.get_json()["error"])

        chc = self.client.post(
            self._url("/units"),
            json={"org_level_id": levels["chc"]["org_level_id"], "parent_org_unit_id": district["org_unit_id"], "unit_code": "C01", "unit_name": "CHC One"},
            headers=headers,
        ).get_json()["unit"]
        self.assertEqual(chc["path"], "D01.C01")
        self.assertEqual(chc["parent_code"], "D01")

        summary = self.client.get(self._url()).get_json()
        self.assertEqual(summary["units"][0]["children"][0]["unit_code"], "C01")
        self.assertEqual(len(summary["level_cadres"]), len(org.DEFAULT_LEVEL_CADRE_TEMPLATE))

        # Deactivating a parent unit with active children without a choice returns 409 with preview.
        refused_toggle = self.client.post(
            self._url(f"/units/{district['org_unit_id']}/toggle"), json={"is_active": False}, headers=headers
        )
        self.assertEqual(refused_toggle.status_code, 409)
        self.assertTrue(refused_toggle.get_json()["requires_child_action"])
        self.assertEqual([c["unit_code"] for c in refused_toggle.get_json()["active_children"]], ["C01"])

        toggled = self.client.post(
            self._url(f"/units/{district['org_unit_id']}/toggle"),
            json={"is_active": False, "cascade_deactivate_children": True},
            headers=headers,
        )
        self.assertEqual(toggled.get_json()["changed"], 2)
        self.assertEqual(self.client.get(self._url("/units")).get_json()["units"], [])
        self.assertEqual(len(self.client.get(self._url("/units?include_inactive=1")).get_json()["units"]), 2)

        perm = self.client.put(
            self._url("/level-cadres"),
            json={"org_level_id": levels["district"]["org_level_id"], "cadre_id": summary["cadres"][0]["cadre_id"], "can_fill_va_form": False, "can_code_va_form": True},
            headers=headers,
        )
        self.assertEqual(perm.status_code, 200)
        self.assertTrue(perm.get_json()["level_cadre"]["can_code_va_form"])

        reactivated = self.client.post(self._url(f"/units/{district['org_unit_id']}/toggle"), json={"is_active": True}, headers=headers)
        self.assertEqual(reactivated.status_code, 200)
        self.assertEqual(
            self.client.post(self._url(f"/units/{chc['org_unit_id']}/toggle"), json={"is_active": True}, headers=headers).status_code,
            200,
        )
        smo_cadre_id = next(c["cadre_id"] for c in summary["cadres"] if c["cadre_code"] == "SMO")
        worker = self.client.post(
            self._url("/workers"),
            json={"org_unit_id": chc["org_unit_id"], "cadre_id": smo_cadre_id,
                  "worker_name": "Dr Test", "email": "dr.test@example.org", "email_confirm": "dr.test@example.org", "language_codes": "en"},
            headers=headers,
        )
        self.assertEqual(worker.status_code, 201, worker.get_json())
        blocked = self.client.put(
            self._url("/level-cadres"),
            json={"org_level_id": levels["chc"]["org_level_id"], "cadre_id": smo_cadre_id, "can_fill_va_form": False, "can_code_va_form": True, "is_active": False},
            headers=headers,
        )
        self.assertEqual(blocked.status_code, 400)
        self.assertIn("first", blocked.get_json()["error"])

    def test_put_level_cadres_without_the_report_deaths_key_keeps_the_stored_value(self):
        self._login(str(self.base_admin_id))
        headers = self._csrf_headers()
        self.client.post(self._url("/seed-template"), json={}, headers=headers)
        levels = {lv["level_code"]: lv for lv in self.client.get(self._url("/levels")).get_json()["levels"]}
        cadres = {c["cadre_code"]: c for c in self.client.get(self._url()).get_json()["cadres"]}
        body = {"org_level_id": levels["subcentre"]["org_level_id"], "cadre_id": cadres["ANM"]["cadre_id"],
                "can_fill_va_form": False, "can_code_va_form": False}
        # Present first: the seeded ANM may report deaths at a sub-centre.
        seeded = {(r["level_code"], r["cadre_code"]): r for r in self.client.get(self._url("/level-cadres")).get_json()["level_cadres"]}
        self.assertTrue(seeded[("subcentre", "ANM")]["can_report_deaths"])
        kept = self.client.put(self._url("/level-cadres"), json=body, headers=headers)
        self.assertEqual(kept.status_code, 200, kept.get_json())
        self.assertTrue(kept.get_json()["level_cadre"]["can_report_deaths"])
        cleared = self.client.put(self._url("/level-cadres"), json={**body, "can_report_deaths": False}, headers=headers)
        self.assertFalse(cleared.get_json()["level_cadre"]["can_report_deaths"])
        again = self.client.put(self._url("/level-cadres"), json=body, headers=headers)
        self.assertFalse(again.get_json()["level_cadre"]["can_report_deaths"])  # absent keeps the cleared value

    def _seed_one_unit(self, headers):
        self.client.post(self._url("/seed-template"), json={}, headers=headers)
        levels = {lv["level_code"]: lv for lv in self.client.get(self._url("/levels")).get_json()["levels"]}
        level = levels["district"]
        created = self.client.post(
            self._url("/units"),
            json={"org_level_id": level["org_level_id"], "unit_code": "D01", "unit_name": "District One"},
            headers=headers,
        )
        return created.get_json()["unit"]["org_unit_id"]

    def test_va_presets_set_get_and_clear_round_trip(self):
        self._login(str(self.base_admin_id))
        headers = self._csrf_headers()
        unit_id = self._seed_one_unit(headers)

        empty = self.client.get(self._url(f"/units/{unit_id}/va-presets"))
        self.assertEqual(empty.status_code, 200)
        self.assertIsNone(empty.get_json()["va_presets"])
        self.assertEqual(empty.get_json()["inherited"], {})

        saved = self.client.put(
            self._url(f"/units/{unit_id}/va-presets"),
            json={"hiv_mortality": "high", "malaria_mortality": "low"},
            headers=headers,
        )
        self.assertEqual(saved.status_code, 200, saved.get_json())
        self.assertEqual(saved.get_json()["va_presets"]["hiv_mortality"], "high")
        self.assertEqual(saved.get_json()["va_presets"]["malaria_mortality"], "low")

        fetched = self.client.get(self._url(f"/units/{unit_id}/va-presets")).get_json()
        self.assertEqual(fetched["va_presets"]["hiv_mortality"], "high")

        cleared = self.client.delete(self._url(f"/units/{unit_id}/va-presets"), headers=headers)
        self.assertEqual(cleared.status_code, 200)
        self.assertTrue(cleared.get_json()["cleared"])
        self.assertIsNone(self.client.get(self._url(f"/units/{unit_id}/va-presets")).get_json()["va_presets"])

    def test_va_presets_reports_inherited_value_and_source_unit(self):
        self._login(str(self.base_admin_id))
        headers = self._csrf_headers()
        parent_id = self._seed_one_unit(headers)
        self.client.put(
            self._url(f"/units/{parent_id}/va-presets"), json={"hiv_mortality": "veryl"}, headers=headers,
        )
        levels = {lv["level_code"]: lv for lv in self.client.get(self._url("/levels")).get_json()["levels"]}
        child = self.client.post(
            self._url("/units"),
            json={
                "org_level_id": levels["chc"]["org_level_id"], "parent_org_unit_id": parent_id,
                "unit_code": "C01", "unit_name": "Child CHC",
            },
            headers=headers,
        ).get_json()["unit"]

        result = self.client.get(self._url(f"/units/{child['org_unit_id']}/va-presets")).get_json()

        self.assertIsNone(result["va_presets"])
        self.assertEqual(result["inherited"]["hiv_mortality"]["value"], "veryl")
        self.assertEqual(result["inherited"]["hiv_mortality"]["source_unit_name"], "District One")

    def test_va_presets_rejects_bad_value(self):
        self._login(str(self.base_admin_id))
        headers = self._csrf_headers()
        unit_id = self._seed_one_unit(headers)

        bad = self.client.put(
            self._url(f"/units/{unit_id}/va-presets"), json={"hiv_mortality": "bogus"}, headers=headers,
        )

        self.assertEqual(bad.status_code, 400)
        self.assertIsNone(org.get_unit_va_presets(self.PROJECT, unit_id))

    def test_va_presets_requires_csrf_token(self):
        self._login(str(self.base_admin_id))
        unit_id = self._seed_one_unit(self._csrf_headers())

        no_token = self.client.put(self._url(f"/units/{unit_id}/va-presets"), json={"hiv_mortality": "high"})

        self.assertEqual(no_token.status_code, 400)
        self.assertIn("CSRF", no_token.get_json()["error"])
        self.assertIsNone(org.get_unit_va_presets(self.PROJECT, unit_id))

    def test_va_presets_forbidden_for_coder(self):
        self._login(str(self.base_coder_id))

        response = self.client.get(self._url("/units/00000000-0000-0000-0000-000000000000/va-presets"))

        self.assertIn(response.status_code, (302, 403))

    def test_export_and_import_endpoints(self):
        self._login(str(self.base_admin_id))
        headers = self._csrf_headers()
        org.seed_default_organization(self.PROJECT)
        level = next(lv for lv in org.list_levels(self.PROJECT) if lv.level_code == "district")
        org.create_unit(self.PROJECT, org_level_id=level.org_level_id, unit_code="D01", unit_name="District One")
        db.session.commit()

        xlsx = self.client.get(self._url("/export.xlsx"))
        self.assertEqual(xlsx.status_code, 200)
        self.assertIn("spreadsheetml", xlsx.mimetype)
        csv_resp = self.client.get(self._url("/export/units.csv"))
        self.assertEqual(csv_resp.status_code, 200)
        self.assertIn("D01,District One,district", csv_resp.get_data(as_text=True))
        choices = self.client.get(self._url("/odk-choices.csv"))
        self.assertIn("org_district,D01,District One,", choices.get_data(as_text=True))
        self.assertEqual(self.client.get(self._url("/export/bogus.csv")).status_code, 400)

        dry = self.client.post(
            self._url("/import"),
            data={"file": (io.BytesIO(xlsx.data), "org.xlsx"), "dry_run": "1"},
            content_type="multipart/form-data",
            headers=headers,
        )
        self.assertEqual(dry.status_code, 200, dry.get_json())
        body = dry.get_json()
        self.assertTrue(body["dry_run"])
        self.assertFalse(body["plan"]["applied"])
        self.assertEqual(body["plan"]["counts"]["units"]["update"], 1)

        rejected = self.client.post(
            self._url("/import"),
            data={"file": (io.BytesIO(b"not a workbook"), "org.xlsx"), "dry_run": "1"},
            content_type="multipart/form-data",
            headers=headers,
        )
        self.assertEqual(rejected.status_code, 400)

    def test_upload_request_limits_return_json_413(self):
        self._login(str(self.base_admin_id))
        headers = self._csrf_headers()
        for suffix, size, filename in (
            ("/project-users/import", 1024 * 1024 + 64 * 1024, "users.csv"),
            ("/import", 5 * 1024 * 1024 + 64 * 1024, "units.csv"),
        ):
            with self.subTest(suffix=suffix):
                response = self.client.post(
                    self._url(suffix),
                    data={"file": (io.BytesIO(b"x" * (size + 1)), filename)},
                    content_type="multipart/form-data", headers=headers,
                )
                self.assertEqual(response.status_code, 413, response.get_json())
                self.assertIn("limit", response.get_json()["error"])

    def test_json_unit_text_is_safe_in_csv_and_xlsx_exports(self):
        self._login(str(self.base_admin_id))
        headers = self._csrf_headers()
        org.seed_default_organization(self.PROJECT)
        district = next(level for level in org.list_levels(self.PROJECT) if level.level_code == "district")
        created = self.client.post(
            self._url("/units"),
            json={"org_level_id": str(district.org_level_id), "unit_code": "D90",
                  "unit_name": "=HYPERLINK(1)", "latitude": "-12.5"},
            headers=headers,
        )
        self.assertEqual(created.status_code, 201, created.get_json())
        unit_id = created.get_json()["unit"]["org_unit_id"]
        updated = self.client.patch(
            self._url(f"/units/{unit_id}"), json={"remarks": "@SUM(1)"}, headers=headers,
        )
        self.assertEqual(updated.status_code, 200, updated.get_json())
        self.assertEqual(updated.get_json()["unit"]["unit_name"], "=HYPERLINK(1)")

        csv_rows = self.client.get(self._url("/export/units.csv")).get_data(as_text=True)
        self.assertIn("'=HYPERLINK(1)", csv_rows)
        self.assertIn("'@SUM(1)", csv_rows)
        self.assertIn("-12.500000", csv_rows)
        workbook = load_workbook(io.BytesIO(self.client.get(self._url("/export.xlsx")).data))
        try:
            headings = [cell.value for cell in workbook["units"][1]]
            exported = next(row for row in workbook["units"].iter_rows(min_row=2)
                            if row[0].value == "D90")
            self.assertEqual(exported[headings.index("unit_name")].value, "'=HYPERLINK(1)")
            self.assertEqual(exported[headings.index("unit_name")].data_type, "s")
            self.assertEqual(exported[headings.index("remarks")].value, "'@SUM(1)")
            self.assertEqual(exported[headings.index("latitude")].value, "-12.500000")
        finally:
            workbook.close()

    def _import_csv(self, headers, sheet, text, dry_run):
        return self.client.post(
            self._url("/import"),
            data={"file": (io.BytesIO(text.encode("utf-8-sig")), f"{sheet}.csv"), "sheet": sheet, "dry_run": "1" if dry_run else "0"},
            content_type="multipart/form-data",
            headers=headers,
        )

    def test_csv_import_adds_cadres_placements_and_workers(self):
        self._login(str(self.base_admin_id))
        headers = self._csrf_headers()
        org.seed_default_organization(self.PROJECT)
        levels = {lv.level_code: lv for lv in org.list_levels(self.PROJECT)}
        district = org.create_unit(self.PROJECT, org_level_id=levels["district"].org_level_id, unit_code="D01", unit_name="District One")
        org.create_unit(self.PROJECT, org_level_id=levels["chc"].org_level_id, parent_org_unit_id=district.org_unit_id, unit_code="C01", unit_name="CHC One")
        db.session.commit()

        cadres = "cadre_code,cadre_name\nLT,Lab Technician\n"
        dry = self._import_csv(headers, "cadres", cadres, dry_run=True)
        self.assertEqual(dry.status_code, 200, dry.get_json())
        self.assertEqual(dry.get_json()["plan"]["counts"]["cadres"]["create"], 1)
        self.assertNotIn("LT", {c.cadre_code for c in org.list_cadres(self.PROJECT)})
        self.assertEqual(self._import_csv(headers, "cadres", cadres, dry_run=False).status_code, 200)
        self.assertIn("LT", {c.cadre_code for c in org.list_cadres(self.PROJECT)})

        placements = "level_code,cadre_code,can_fill_va_form,can_code_va_form\nchc,LT,true,false\n"
        self.assertEqual(self._import_csv(headers, "level_cadres", placements, dry_run=False).status_code, 200)

        workers = "worker_code,worker_name,unit_code,cadre_code,phone,user_email,remarks,is_active\nW01,Lab Person,C01,LT,,,,\n"
        applied = self._import_csv(headers, "workers", workers, dry_run=False)
        self.assertEqual(applied.status_code, 200, applied.get_json())
        worker = {w["worker_code"]: w for w in org.list_workers(self.PROJECT, include_inactive=True)}["W01"]
        self.assertTrue(worker["is_active"])
        self.assertIsNone(worker["phone"])

        # A re-upload with a blank is_active keeps the worker active: blank
        # CSV cells mean "no value", as empty workbook cells do.
        again = self._import_csv(headers, "workers", workers.replace("Lab Person", "Lab Person Two"), dry_run=False)
        self.assertEqual(again.status_code, 200, again.get_json())
        worker = {w["worker_code"]: w for w in org.list_workers(self.PROJECT, include_inactive=True)}["W01"]
        self.assertEqual(worker["worker_name"], "Lab Person Two")
        self.assertTrue(worker["is_active"])

        # A worker whose cadre is not placed at the unit's level is refused.
        bad = self._import_csv(headers, "workers", "worker_code,worker_name,unit_code,cadre_code\nW02,Nobody,D01,LT\n", dry_run=True)
        self.assertEqual(bad.status_code, 400)
        self.assertTrue(bad.get_json()["plan"]["errors"])

        unknown_sheet = self._import_csv(headers, "bogus", cadres, dry_run=True)
        self.assertEqual(unknown_sheet.status_code, 400)

    def test_import_rejects_invalid_boolean_rows_without_partial_writes(self):
        self._login(str(self.base_admin_id))
        headers = self._csrf_headers()
        org.seed_default_organization(self.PROJECT)
        levels = {level.level_code: level for level in org.list_levels(self.PROJECT)}
        district = org.create_unit(
            self.PROJECT, org_level_id=levels["district"].org_level_id,
            unit_code="D90", unit_name="District 90",
        )
        org.create_unit(
            self.PROJECT, org_level_id=levels["chc"].org_level_id,
            unit_code="C90", unit_name="CHC 90", parent_org_unit_id=district.org_unit_id,
        )
        db.session.commit()
        original_level_cadres = org.list_level_cadres(self.PROJECT)
        cases = (
            ("levels", "level_code,level_name,depth,is_optional,is_active\n"
             "newlevel,New Level,8,false,true\ninvalid,Invalid Level,9,maybe,true\n",
             "is_optional", "NEWLEVEL"),
            ("levels", "level_code,level_name,depth,is_optional,is_active\n"
             "newlevel,New Level,8,false,true\ninvalid,Invalid Level,9,false,maybe\n",
             "is_active", "NEWLEVEL"),
            ("cadres", "cadre_code,cadre_name,is_active\nLT,Lab Technician,true\nXT,Other,falseish\n",
             "is_active", "LT"),
            ("level_cadres", "level_code,cadre_code,can_fill_va_form,can_code_va_form,is_active\n"
             "district,MO,true,false,true\ndistrict,SMO,maybe,false,true\n",
             "can_fill_va_form", None),
            ("level_cadres", "level_code,cadre_code,can_fill_va_form,can_code_va_form,is_active\n"
             "district,MO,true,false,true\ndistrict,SMO,false,maybe,true\n",
             "can_code_va_form", None),
            ("level_cadres", "level_code,cadre_code,can_fill_va_form,can_code_va_form,is_active\n"
             "district,MO,true,false,true\ndistrict,SMO,false,false,maybe\n",
             "is_active", None),
            ("workers", "worker_code,worker_name,unit_code,cadre_code,is_active\n"
             "W90,First,C90,MO,true\nW91,Second,C90,MO,maybe\n",
             "is_active", "W90"),
        )
        for sheet, content, field, created_code in cases:
            with self.subTest(sheet=sheet, field=field):
                for dry_run in (True, False):
                    response = self._import_csv(headers, sheet, content, dry_run=dry_run)
                    self.assertEqual(response.status_code, 400, response.get_json())
                    self.assertIn(f"{sheet} row 3: {field} must be true or false",
                                  response.get_json()["plan"]["errors"][0])
                if sheet == "levels":
                    self.assertNotIn(created_code, {level.level_code.upper() for level in
                                                    org.list_levels(self.PROJECT, include_inactive=True)})
                elif sheet == "cadres":
                    self.assertNotIn(created_code, {cadre.cadre_code for cadre in
                                                    org.list_cadres(self.PROJECT, include_inactive=True)})
                elif sheet == "level_cadres":
                    self.assertEqual(org.list_level_cadres(self.PROJECT), original_level_cadres)
                elif sheet == "workers":
                    self.assertNotIn(created_code, {worker["worker_code"] for worker in
                                                    org.list_workers(self.PROJECT, include_inactive=True)})

    def test_new_rows_honor_explicit_false_and_blank_preserves_it(self):
        self._login(str(self.base_admin_id))
        headers = self._csrf_headers()
        org.seed_default_organization(self.PROJECT)
        levels = {level.level_code: level for level in org.list_levels(self.PROJECT)}
        district = org.create_unit(
            self.PROJECT, org_level_id=levels["district"].org_level_id,
            unit_code="D90", unit_name="District 90",
        )
        org.create_unit(
            self.PROJECT, org_level_id=levels["chc"].org_level_id,
            unit_code="C90", unit_name="CHC 90", parent_org_unit_id=district.org_unit_id,
        )
        db.session.commit()
        uploads = (
            ("levels", "level_code,level_name,depth,is_active\nnewlevel,New Level,8,false\n"),
            ("cadres", "cadre_code,cadre_name,is_active\nLT,Lab Technician,false\n"),
            ("workers", "worker_code,worker_name,unit_code,cadre_code,is_active\n"
             "W90,Worker 90,C90,MO,false\n"),
        )
        for sheet, content in uploads:
            with self.subTest(sheet=sheet):
                preview = self._import_csv(headers, sheet, content, dry_run=True)
                self.assertEqual(preview.status_code, 200, preview.get_json())
                applied = self._import_csv(headers, sheet, content, dry_run=False)
                self.assertEqual(applied.status_code, 200, applied.get_json())
        level = next(level for level in org.list_levels(self.PROJECT, include_inactive=True)
                     if level.level_code == "newlevel")
        cadre = next(cadre for cadre in org.list_cadres(self.PROJECT, include_inactive=True)
                     if cadre.cadre_code == "LT")
        worker = next(worker for worker in org.list_workers(self.PROJECT, include_inactive=True)
                      if worker["worker_code"] == "W90")
        self.assertFalse(level.is_active)
        self.assertFalse(cadre.is_active)
        self.assertFalse(worker["is_active"])
        for sheet, content in (
            ("levels", "level_code,level_name,depth,is_active\nnewlevel,New Level,8,\n"),
            ("cadres", "cadre_code,cadre_name,is_active\nLT,Lab Technician,\n"),
            ("workers", "worker_code,worker_name,unit_code,cadre_code,is_active\n"
             "W90,Worker 90,C90,MO,\n"),
        ):
            response = self._import_csv(headers, sheet, content, dry_run=False)
            self.assertEqual(response.status_code, 200, response.get_json())
        db.session.expire_all()
        self.assertFalse(level.is_active)
        self.assertFalse(cadre.is_active)
        self.assertFalse(next(worker for worker in org.list_workers(self.PROJECT, include_inactive=True)
                              if worker["worker_code"] == "W90")["is_active"])

    def test_worker_code_is_optional(self):
        self._login(str(self.base_admin_id))
        headers = self._csrf_headers()
        org.seed_default_organization(self.PROJECT)
        levels = {lv.level_code: lv for lv in org.list_levels(self.PROJECT)}
        district = org.create_unit(self.PROJECT, org_level_id=levels["district"].org_level_id, unit_code="D01", unit_name="District One")
        chc = org.create_unit(self.PROJECT, org_level_id=levels["chc"].org_level_id, parent_org_unit_id=district.org_unit_id, unit_code="C01", unit_name="CHC One")
        mo = {c.cadre_code: c for c in org.list_cadres(self.PROJECT)}["MO"]
        db.session.commit()

        created = self.client.post(
            self._url("/workers"),
            json={"org_unit_id": str(chc.org_unit_id), "cadre_id": str(mo.cadre_id),
                  "worker_name": "Dr Kaur", "email": "dr.kaur@example.org", "email_confirm": "dr.kaur@example.org", "language_codes": "en"},
            headers=headers,
        )
        self.assertEqual(created.status_code, 201, created.get_json())
        self.assertEqual(created.get_json()["worker"]["worker_code"], "W00001")

        # Code-less rows match by unit and name: Dr Kaur is updated, Dr Rao is new,
        # and re-uploading the same file creates nothing more.
        rows = "worker_name,unit_code,cadre_code,phone\ndr kaur,C01,MO,111\nDr Rao,C01,MO,\n"
        for _ in range(2):
            self.assertEqual(self._import_csv(headers, "workers", rows, dry_run=False).status_code, 200)
        workers = {w["worker_code"]: w for w in org.list_workers(self.PROJECT, include_inactive=True)}
        self.assertEqual(set(workers), {"W00001", "W00002"})
        self.assertEqual(workers["W00001"]["phone"], "111")
        self.assertEqual(workers["W00002"]["worker_name"], "Dr Rao")

    def _default_worker_fixture(self, *, code="D45"):
        org.seed_default_organization(self.PROJECT)
        levels = {level.level_code: level for level in org.list_levels(self.PROJECT)}
        district = org.create_unit(
            self.PROJECT,
            org_level_id=levels["district"].org_level_id,
            unit_code=code,
            unit_name="Default District",
        )
        cadre = {item.cadre_code: item for item in org.list_cadres(self.PROJECT)}["CS"]
        db.session.commit()
        return district, cadre

    def test_worker_defaults_create_email_account_and_invitation(self):
        self._login(str(self.base_admin_id))
        headers = self._csrf_headers()
        district, cadre = self._default_worker_fixture()
        with patch("app.services.email_service.is_mail_configured", return_value=True), patch(
            "app.services.email_service.send_verification_email", return_value=True
        ) as verification:
            response = self.client.post(
                self._url("/workers"),
                json={
                    "org_unit_id": str(district.org_unit_id),
                    "cadre_id": str(cadre.cadre_id),
                    "worker_name": "Dr Defaults",
                    "email": "defaults@example.org", "email_confirm": "defaults@example.org",
                    "language_codes": "en",
                },
                headers=headers,
            )
        self.assertEqual(response.status_code, 201, response.get_json())
        self.assertEqual(response.headers.get("Cache-Control"), "no-store")
        onboarding = response.get_json()["onboarding"]
        self.assertEqual(onboarding["created_users"], 1)
        self.assertEqual(onboarding["changed_grants"], 1)
        self.assertEqual(onboarding["invitations_queued"], 1)
        verification.assert_called_once()
        user = db.session.scalar(sa.select(VaUsers).where(VaUsers.email == "defaults@example.org"))
        self.assertIsNotNone(user)
        self.assertEqual(response.get_json()["worker"]["user_id"], str(user.user_id))
        grant = db.session.scalar(sa.select(VaUserAccessGrants).where(
            VaUserAccessGrants.user_id == user.user_id,
            VaUserAccessGrants.org_unit_id == district.org_unit_id,
        ))
        self.assertEqual(grant.role, VaAccessRoles.site_pi)

    def test_worker_defaults_issue_mobile_code_and_no_store(self):
        self._login(str(self.base_admin_id))
        district, cadre = self._default_worker_fixture(code="D46")
        response = self.client.post(
            self._url("/workers"),
            json={
                "org_unit_id": str(district.org_unit_id),
                "cadre_id": str(cadre.cadre_id),
                "worker_name": "Mobile Defaults",
                "phone": "9876543210",
                "no_email": True,
                "phone_confirm": "9876543210",
                "year_of_birth": "1990",
                "sex": "female",
                "language_codes": "en",
            },
            headers=self._csrf_headers(),
        )
        self.assertEqual(response.status_code, 201, response.get_json())
        self.assertEqual(response.headers.get("Cache-Control"), "no-store")
        code = response.get_json()["onboarding"]["sign_in_codes"]
        self.assertEqual(len(code), 1)
        self.assertEqual(code[0]["mobile"], "******3210")
        user = db.session.scalar(sa.select(VaUsers).where(VaUsers.mobile_login == "9876543210"))
        self.assertIsNotNone(user)
        self.assertEqual((user.year_of_birth, user.sex), (1990, "female"))
        self.assertEqual(response.get_json()["worker"]["user_id"], str(user.user_id))

    def _assert_invalid_worker_account(self, **fields):
        self._login(str(self.base_admin_id))
        district, cadre = self._default_worker_fixture(code="DVALID")
        unit_id = district.org_unit_id
        before_users = db.session.scalar(sa.select(sa.func.count()).select_from(VaUsers))
        before_grants = db.session.scalar(sa.select(sa.func.count()).select_from(VaUserAccessGrants))
        response = self.client.post(
            self._url("/workers"),
            json={"org_unit_id": str(unit_id), "cadre_id": str(cadre.cadre_id),
                  "worker_name": "Rejected account", "email": "confirm@example.org",
                  "email_confirm": "confirm@example.org", "language_codes": "en", **fields},
            headers=self._csrf_headers(),
        )
        self.assertEqual(response.status_code, 400, response.get_json())
        self.assertEqual(db.session.scalar(sa.select(sa.func.count()).select_from(VaUsers)), before_users)
        self.assertEqual(db.session.scalar(sa.select(sa.func.count()).select_from(VaUserAccessGrants)), before_grants)
        self.assertEqual(db.session.scalar(sa.select(sa.func.count()).select_from(MasOrgUnitWorker).where(
            MasOrgUnitWorker.org_unit_id == unit_id)), 0)

    def test_worker_email_confirmation_required(self):
        self._assert_invalid_worker_account(email_confirm="")

    def test_worker_email_confirmation_mismatch(self):
        self._assert_invalid_worker_account(email_confirm="different@example.org")

    def test_worker_mobile_requires_explicit_no_email(self):
        self._assert_invalid_worker_account(email="", email_confirm="", phone="9876543210")

    def test_worker_mobile_confirmation_mismatch(self):
        self._assert_invalid_worker_account(email="", email_confirm="", no_email=True, phone="9876543210", phone_confirm="9876543211")

    def test_worker_mobile_rejects_nine_digits(self):
        self._assert_invalid_worker_account(phone="987654321")

    def test_worker_mobile_rejects_eleven_digits(self):
        self._assert_invalid_worker_account(phone="98765432101")

    def test_worker_mobile_rejects_non_ascii_digits(self):
        self._assert_invalid_worker_account(phone="९८७६५४३२१०")

    def test_worker_mobile_rejects_spaces(self):
        self._assert_invalid_worker_account(phone=" 9876543210")

    def test_worker_no_email_requires_mobile(self):
        self._assert_invalid_worker_account(email="", email_confirm="", no_email=True)

    def test_worker_no_email_rejects_email(self):
        self._assert_invalid_worker_account(no_email=True, phone="9876543210", phone_confirm="9876543210")

    def test_worker_no_email_rejects_stale_email_confirmation(self):
        self._assert_invalid_worker_account(
            email="", no_email=True, phone="9876543210", phone_confirm="9876543210"
        )

    def test_worker_invalid_profile_rolls_back_account_grant_and_worker(self):
        self._assert_invalid_worker_account(year_of_birth="1899")

    def test_worker_invalid_sex_rolls_back_account_grant_and_worker(self):
        self._assert_invalid_worker_account(sex="invalid")

    def test_worker_existing_account_preserves_password_and_grant_is_idempotent(self):
        self._login(str(self.base_admin_id))
        district, cadre = self._default_worker_fixture(code="D47")
        existing = db.session.get(VaUsers, self.base_coder_id)
        password = existing.password
        body = {
            "org_unit_id": str(district.org_unit_id),
            "cadre_id": str(cadre.cadre_id),
            "worker_name": "Existing Account",
            "user": str(existing.user_id),
            "year_of_birth": "1899",
            "sex": "invalid",
        }
        first = self.client.post(self._url("/workers"), json=body, headers=self._csrf_headers())
        self.assertEqual(first.status_code, 201, first.get_json())
        second = self.client.post(
            self._url("/workers"), json={**body, "worker_code": "W0047"}, headers=self._csrf_headers()
        )
        self.assertEqual(second.status_code, 201, second.get_json())
        self.assertEqual(second.get_json()["onboarding"]["created_users"], 0)
        self.assertEqual(second.get_json()["onboarding"]["changed_grants"], 0)
        db.session.refresh(existing)
        self.assertEqual(existing.password, password)
        self.assertEqual(db.session.scalar(sa.select(sa.func.count()).select_from(VaUserAccessGrants).where(
            VaUserAccessGrants.user_id == existing.user_id,
            VaUserAccessGrants.org_unit_id == district.org_unit_id,
            VaUserAccessGrants.role == VaAccessRoles.site_pi,
        )), 1)

    def test_project_pi_cannot_create_worker_account_from_defaults(self):
        district, cadre = self._default_worker_fixture(code="D48")
        db.session.add(VaUserAccessGrants(
            user_id=self.base_project_pi_user.user_id,
            role=VaAccessRoles.project_pi,
            scope_type=VaAccessScopeTypes.project,
            project_id=self.PROJECT,
            grant_status=VaStatuses.active,
        ))
        db.session.commit()
        self._login(str(self.base_project_pi_id))
        response = self.client.post(
            self._url("/workers"),
            json={
                "org_unit_id": str(district.org_unit_id),
                "cadre_id": str(cadre.cadre_id),
                "worker_name": "PI Cannot Create",
                "email": "pi-new@example.org", "email_confirm": "pi-new@example.org",
                "language_codes": "en",
            },
            headers=self._csrf_headers(),
        )
        self.assertEqual(response.status_code, 400, response.get_json())
        self.assertIsNone(db.session.scalar(sa.select(VaUsers).where(VaUsers.email == "pi-new@example.org")))
        self.assertFalse(db.session.scalar(sa.select(sa.exists().where(
            sa.and_(MasOrgUnitWorker.project_id == self.PROJECT,
                    MasOrgUnitWorker.worker_name == "PI Cannot Create")
        ))))

    def test_worker_defaults_reject_email_and_mobile_for_different_accounts(self):
        self._login(str(self.base_admin_id))
        district, cadre = self._default_worker_fixture(code="D50")
        existing = db.session.get(VaUsers, self.base_coder_id)
        existing.phone = "9876543210"
        existing.mobile_login = "9876543210"
        db.session.flush()
        response = self.client.post(
            self._url("/workers"),
            json={
                "org_unit_id": str(district.org_unit_id),
                "cadre_id": str(cadre.cadre_id),
                "worker_name": "Conflicting Identity",
                "email": "different@example.org",
                "phone": "9876543210",
                "language_codes": "en",
            },
            headers=self._csrf_headers(),
        )
        self.assertEqual(response.status_code, 400, response.get_json())
        self.assertIsNone(db.session.scalar(sa.select(VaUsers).where(VaUsers.email == "different@example.org")))
        self.assertFalse(db.session.scalar(sa.select(sa.exists().where(
            sa.and_(MasOrgUnitWorker.project_id == self.PROJECT,
                    MasOrgUnitWorker.worker_name == "Conflicting Identity")
        ))))

    def test_roster_only_worker_rejects_invalid_phone_on_submission(self):
        self._login(str(self.base_admin_id))
        district, cadre = self._default_worker_fixture(code="DNOEMAIL")
        unit_id = district.org_unit_id
        cadre_id = cadre.cadre_id
        org.upsert_level_cadre(
            self.PROJECT, org_level_id=district.org_level_id, cadre_id=cadre_id,
            can_fill_va_form=False, can_code_va_form=False, default_roles=[],
        )
        db.session.commit()
        response = self.client.post(
            self._url("/workers"),
            json={"org_unit_id": str(unit_id), "cadre_id": str(cadre_id),
                  "worker_name": "Invalid roster phone", "phone": "12345"},
            headers=self._csrf_headers(),
        )
        self.assertEqual(response.status_code, 400, response.get_json())
        self.assertIn("exactly 10 digits", response.get_json()["error"])
        self.assertEqual(db.session.scalar(sa.select(sa.func.count()).select_from(MasOrgUnitWorker).where(
            MasOrgUnitWorker.org_unit_id == unit_id)), 0)

    def test_worker_without_defaults_remains_roster_only(self):
        self._login(str(self.base_admin_id))
        district, cadre = self._default_worker_fixture(code="D49")
        level = db.session.get(MasOrgLevel, district.level.org_level_id)
        org.upsert_level_cadre(
            self.PROJECT,
            org_level_id=level.org_level_id,
            cadre_id=cadre.cadre_id,
            can_fill_va_form=False,
            can_code_va_form=False,
            default_roles=[],
        )
        db.session.commit()
        response = self.client.post(
            self._url("/workers"),
            json={
                "org_unit_id": str(district.org_unit_id),
                "cadre_id": str(cadre.cadre_id),
                "worker_name": "Roster Only",
                "email": "ignored@example.org",
            },
            headers=self._csrf_headers(),
        )
        self.assertEqual(response.status_code, 201, response.get_json())
        self.assertEqual(response.get_json()["onboarding"]["created_users"], 0)
        self.assertIsNone(db.session.scalar(sa.select(VaUsers).where(VaUsers.email == "ignored@example.org")))
        self.assertIsNone(response.get_json()["worker"]["user_id"])

    # -- project structure mode (docs/policy/organization-model.md) ----------

    def test_writes_refused_for_sites_project_reads_allowed(self):
        self._login(str(self.base_admin_id))
        headers = self._csrf_headers()
        base = f"/admin/api/organization/{self.SITES_PROJECT}"

        self.assertEqual(self.client.get(base).status_code, 200)
        self.assertEqual(self.client.get(f"{base}/levels").status_code, 200)

        refused = [
            self.client.post(f"{base}/seed-template", json={}, headers=headers),
            self.client.post(
                f"{base}/levels",
                json={"level_code": "district", "level_name": "District", "depth": 1},
                headers=headers,
            ),
            self.client.post(f"{base}/cadres", json={"cadre_code": "MO", "cadre_name": "MO"}, headers=headers),
            self.client.put(f"{base}/level-cadres", json={}, headers=headers),
            self.client.post(
                f"{base}/import",
                data={"file": (io.BytesIO(b"x"), "org.xlsx"), "dry_run": "1"},
                content_type="multipart/form-data",
                headers=headers,
            ),
        ]
        for response in refused:
            self.assertEqual(response.status_code, 409, response.get_json())
            self.assertIn("Organization", response.get_json()["error"])
        self.assertEqual(org.list_levels(self.SITES_PROJECT, include_inactive=True), [])
        self.assertEqual(org.list_cadres(self.SITES_PROJECT, include_inactive=True), [])

    def test_create_project_accepts_and_rejects_structure_mode(self):
        self._login(str(self.base_admin_id))
        headers = self._csrf_headers()
        payload = {"project_name": "Mode", "project_nickname": "Mode"}

        default = self.client.post("/admin/api/projects", json={**payload, "project_id": "MODE01"}, headers=headers)
        self.assertEqual(default.status_code, 201, default.get_json())
        self.assertEqual(default.get_json()["project"]["project_structure_mode"], "sites")

        explicit = self.client.post(
            "/admin/api/projects",
            json={**payload, "project_id": "MODE02", "project_structure_mode": "organization"},
            headers=headers,
        )
        self.assertEqual(explicit.status_code, 201, explicit.get_json())
        self.assertEqual(explicit.get_json()["project"]["project_structure_mode"], "organization")

        bogus = self.client.post(
            "/admin/api/projects",
            json={**payload, "project_id": "MODE03", "project_structure_mode": "tree"},
            headers=headers,
        )
        self.assertEqual(bogus.status_code, 400)
        self.assertIsNone(db.session.get(VaProjectMaster, "MODE03"))

        bad_update = self.client.put(
            "/admin/api/projects/MODE01", json={"project_structure_mode": ["sites"]}, headers=headers
        )
        self.assertEqual(bad_update.status_code, 400)

    def test_switch_back_to_sites_refused_while_units_active(self):
        self._login(str(self.base_admin_id))
        headers = self._csrf_headers()
        url = f"/admin/api/projects/{self.SITES_PROJECT}"

        to_org = self.client.put(url, json={"project_structure_mode": "organization"}, headers=headers)
        self.assertEqual(to_org.status_code, 200, to_org.get_json())
        self.assertEqual(to_org.get_json()["project"]["project_structure_mode"], "organization")

        org.seed_default_organization(self.SITES_PROJECT, include_cadres=False)
        level = org.list_levels(self.SITES_PROJECT)[0]
        unit = org.create_unit(self.SITES_PROJECT, org_level_id=level.org_level_id, unit_code="D01", unit_name="District One")
        db.session.commit()

        refused = self.client.put(url, json={"project_structure_mode": "sites"}, headers=headers)
        self.assertEqual(refused.status_code, 409)
        self.assertIn("active organization units", refused.get_json()["error"])
        db.session.expire_all()
        self.assertEqual(db.session.get(VaProjectMaster, self.SITES_PROJECT).project_structure_mode, "organization")
        self.assertIsNotNone(org.find_unit_by_code(self.SITES_PROJECT, "D01"))

        org.set_unit_active(self.SITES_PROJECT, unit.org_unit_id, False)
        db.session.commit()
        allowed = self.client.put(url, json={"project_structure_mode": "sites"}, headers=headers)
        self.assertEqual(allowed.status_code, 200, allowed.get_json())
        self.assertEqual(allowed.get_json()["project"]["project_structure_mode"], "sites")

    def test_user_lookup_active_and_channel_redaction(self):
        active_user = VaUsers(
            name="Dr Lookup Active",
            email="dr.lookup.active@example.org",
            phone="9876543210",
            mobile_login="9876543210",
            user_status=VaStatuses.active,
            permission={},
            landing_page="coder",
            pw_reset_t_and_c=False,
            email_verified=True,
            vacode_language=["en"],
        )
        active_user.set_password("UnusedPassword123!")
        inactive_user = VaUsers(
            name="Dr Lookup Inactive",
            email="dr.lookup.inactive@example.org",
            phone="9876543211",
            mobile_login="9876543211",
            user_status=VaStatuses.deactive,
            permission={},
            landing_page="coder",
            pw_reset_t_and_c=False,
            email_verified=True,
            vacode_language=["en"],
        )
        inactive_user.set_password("UnusedPassword123!")
        db.session.add_all([active_user, inactive_user])
        db.session.flush()
        db.session.add(
            VaUserAccessGrants(
                user_id=self.base_project_pi_user.user_id,
                role=VaAccessRoles.project_pi,
                scope_type=VaAccessScopeTypes.project,
                project_id=self.PROJECT,
                grant_status=VaStatuses.active,
            )
        )
        db.session.commit()

        # Platform admin sees full contact info.
        self._login(str(self.base_admin_id))
        res_admin = self.client.get(self._url("/users/lookup?q=dr.lookup.active@example.org"))
        self.assertEqual(res_admin.status_code, 200)
        body_admin = res_admin.get_json()
        self.assertTrue(body_admin["found"])
        self.assertEqual(body_admin["user"]["email"], "dr.lookup.active@example.org")
        self.assertEqual(body_admin["user"]["phone"], "9876543210")

        # Inactive user is not returned.
        res_inactive = self.client.get(self._url("/users/lookup?q=dr.lookup.inactive@example.org"))
        self.assertEqual(res_inactive.status_code, 200)
        self.assertFalse(res_inactive.get_json()["found"])

        # Non-admin Project PI has unqueried contact channel redacted.
        self._login(str(self.base_project_pi_id))
        res_pi_email = self.client.get(self._url("/users/lookup?q=dr.lookup.active@example.org"))
        self.assertEqual(res_pi_email.status_code, 200)
        user_by_email = res_pi_email.get_json()["user"]
        self.assertEqual(user_by_email["email"], "dr.lookup.active@example.org")
        self.assertIsNone(user_by_email["phone"])

        res_pi_phone = self.client.get(self._url("/users/lookup?q=9876543210"))
        self.assertEqual(res_pi_phone.status_code, 200)
        user_by_phone = res_pi_phone.get_json()["user"]
        self.assertEqual(user_by_phone["phone"], "9876543210")
        self.assertIsNone(user_by_phone["email"])

    def test_worker_multi_unit_coding_and_supervision_scope_preserves_site_pi(self):
        self._login(str(self.base_admin_id))
        headers = self._csrf_headers()
        org.seed_default_organization(self.PROJECT)
        levels = {lv.level_code: lv for lv in org.list_levels(self.PROJECT)}
        cadres = {c.cadre_code: c for c in org.list_cadres(self.PROJECT)}

        # Ensure MO at PHC can code VA forms and supervise interviews.
        org.upsert_level_cadre(
            self.PROJECT,
            org_level_id=levels["phc"].org_level_id,
            cadre_id=cadres["MO"].cadre_id,
            can_fill_va_form=False,
            can_code_va_form=True,
            can_supervise_interviews=True,
            default_roles=["coder", "interview_supervisor"],
        )
        dist = org.create_unit(
            self.PROJECT,
            org_level_id=levels["district"].org_level_id,
            unit_code="SCPD1",
            unit_name="Scope District 1",
        )
        chc = org.create_unit(
            self.PROJECT,
            org_level_id=levels["chc"].org_level_id,
            parent_org_unit_id=dist.org_unit_id,
            unit_code="SCPC1",
            unit_name="Scope CHC 1",
        )
        phc1 = org.create_unit(
            self.PROJECT,
            org_level_id=levels["phc"].org_level_id,
            parent_org_unit_id=chc.org_unit_id,
            unit_code="SCPP1",
            unit_name="Scope PHC 1",
        )
        phc2 = org.create_unit(
            self.PROJECT,
            org_level_id=levels["phc"].org_level_id,
            parent_org_unit_id=chc.org_unit_id,
            unit_code="SCPP2",
            unit_name="Scope PHC 2",
        )
        db.session.commit()

        created = self.client.post(
            self._url("/workers"),
            json={
                "org_unit_id": str(phc1.org_unit_id),
                "cadre_id": str(cadres["MO"].cadre_id),
                "worker_name": "Dr Cross Tree MO",
                "email": "dr.crosstree.mo@example.org",
                "email_confirm": "dr.crosstree.mo@example.org",
                "phone": "9123456780",
                "phone_confirm": "9123456780",
                "language_codes": "en",
                "coding_unit_ids": [str(phc1.org_unit_id), str(phc2.org_unit_id)],
                "supervision_unit_ids": [str(phc1.org_unit_id), str(phc2.org_unit_id)],
            },
            headers=headers,
        )
        self.assertEqual(created.status_code, 201, created.get_json())
        worker_payload = created.get_json()["worker"]
        self.assertEqual(sorted(worker_payload["coding_unit_codes"]), ["SCPP1", "SCPP2"])
        self.assertEqual(sorted(worker_payload["supervision_unit_codes"]), ["SCPP1", "SCPP2"])

        worker_id = worker_payload["worker_id"]
        user_id = worker_payload["user_id"]

        # Add an independent site_pi grant on PHC2 and ensure narrowing supervision_unit_ids preserves it.
        site_pi_grant = VaUserAccessGrants(
            user_id=user_id,
            role=VaAccessRoles.site_pi,
            scope_type=VaAccessScopeTypes.org_unit,
            org_unit_id=phc2.org_unit_id,
            grant_status=VaStatuses.active,
        )
        db.session.add(site_pi_grant)
        db.session.commit()

        patched = self.client.patch(
            self._url(f"/workers/{worker_id}"),
            json={
                "coding_unit_ids": [str(phc2.org_unit_id)],
                "supervision_unit_ids": [str(phc1.org_unit_id)],
            },
            headers=headers,
        )
        self.assertEqual(patched.status_code, 200, patched.get_json())
        updated_worker = patched.get_json()["worker"]
        self.assertEqual(updated_worker["coding_unit_codes"], ["SCPP2"])
        self.assertEqual(updated_worker["supervision_unit_codes"], ["SCPP1"])

        db.session.expire_all()
        refreshed_site_pi = db.session.get(VaUserAccessGrants, site_pi_grant.grant_id)
        self.assertEqual(refreshed_site_pi.grant_status, VaStatuses.active)

    def test_bulk_va_presets_and_inheritance(self):
        self._login(str(self.base_admin_id))
        headers = self._csrf_headers()
        org.seed_default_organization(self.PROJECT)
        levels = {lv.level_code: lv for lv in org.list_levels(self.PROJECT)}
        dist = org.create_unit(
            self.PROJECT,
            org_level_id=levels["district"].org_level_id,
            unit_code="PRSD1",
            unit_name="Preset District 1",
        )
        chc = org.create_unit(
            self.PROJECT,
            org_level_id=levels["chc"].org_level_id,
            parent_org_unit_id=dist.org_unit_id,
            unit_code="PRSC1",
            unit_name="Preset CHC 1",
        )
        phc1 = org.create_unit(
            self.PROJECT,
            org_level_id=levels["phc"].org_level_id,
            parent_org_unit_id=chc.org_unit_id,
            unit_code="PRSP1",
            unit_name="Preset PHC 1",
        )
        phc2 = org.create_unit(
            self.PROJECT,
            org_level_id=levels["phc"].org_level_id,
            parent_org_unit_id=chc.org_unit_id,
            unit_code="PRSP2",
            unit_name="Preset PHC 2",
        )
        db.session.commit()

        put_res = self.client.put(
            self._url("/va-presets"),
            json={
                "items": [
                    {"org_unit_id": str(dist.org_unit_id), "hiv_mortality": "high", "malaria_mortality": "low"},
                    {"org_unit_id": str(phc1.org_unit_id), "hiv_mortality": "veryl", "malaria_mortality": "high"},
                ]
            },
            headers=headers,
        )
        self.assertEqual(put_res.status_code, 200, put_res.get_json())
        by_code = {u["unit_code"]: u for u in put_res.get_json()["units"]}
        self.assertEqual(by_code["PRSD1"]["effective_hiv_mortality"], "high")
        self.assertEqual(by_code["PRSD1"]["effective_malaria_mortality"], "low")
        self.assertEqual(by_code["PRSP1"]["effective_hiv_mortality"], "veryl")
        self.assertEqual(by_code["PRSP1"]["effective_malaria_mortality"], "high")
        # PHC2 inherits from PRSD1 via PRSC1
        self.assertEqual(by_code["PRSP2"]["org_unit_id"], str(phc2.org_unit_id))
        self.assertIsNone(by_code["PRSP2"]["own_hiv_mortality"])
        self.assertEqual(by_code["PRSP2"]["effective_hiv_mortality"], "high")
        self.assertEqual(by_code["PRSP2"]["effective_malaria_mortality"], "low")

    def test_mentoring_overview_districts_and_scientist_grants(self):
        from app.models import MapMentorInstituteUser, MasMentorInstitute

        self._login(str(self.base_admin_id))
        headers = self._csrf_headers()
        org.seed_default_organization(self.PROJECT)
        levels = {lv.level_code: lv for lv in org.list_levels(self.PROJECT)}
        dist = org.create_unit(
            self.PROJECT,
            org_level_id=levels["district"].org_level_id,
            unit_code="MNTD1",
            unit_name="Mentoring District 1",
        )
        sci_user = VaUsers(
            name="Dr Mentor Scientist",
            email="dr.mentor.sci@example.org",
            user_status=VaStatuses.active,
            permission={},
            landing_page="coder",
            pw_reset_t_and_c=False,
            email_verified=True,
            vacode_language=["en"],
        )
        sci_user.set_password("UnusedPassword123!")
        inst = MasMentorInstitute(
            institute_code="AIIMSMNT",
            institute_name="AIIMS Mentoring Institute",
            is_active=True,
        )
        unattached_user = VaUsers(
            name="Dr Unattached Scientist",
            email="dr.unattached.sci@example.org",
            phone="9998887776",
            user_status=VaStatuses.active,
            permission={},
            landing_page="coder",
            pw_reset_t_and_c=False,
            email_verified=True,
            vacode_language=["en"],
        )
        unattached_user.set_password("UnusedPassword123!")
        unattached_inst = MasMentorInstitute(
            institute_code="UNATTACHMNT",
            institute_name="Unattached Mentoring Institute",
            is_active=True,
        )
        db.session.add_all([sci_user, inst, unattached_user, unattached_inst])
        db.session.flush()
        db.session.add_all([
            MapMentorInstituteUser(
                institute_id=inst.institute_id,
                user_id=sci_user.user_id,
                is_admin=False,
                is_active=True,
            ),
            MapMentorInstituteUser(
                institute_id=unattached_inst.institute_id,
                user_id=unattached_user.user_id,
                is_admin=False,
                is_active=True,
            ),
        ])
        db.session.commit()

        attach_res = self.client.post(
            self._url("/mentoring/districts"),
            json={"institute_code": "AIIMSMNT", "unit_code": "MNTD1", "action": "attach"},
            headers=headers,
        )
        self.assertEqual(attach_res.status_code, 200, attach_res.get_json())
        overview = attach_res.get_json()
        self.assertTrue(
            any(
                link["institute_code"] == "AIIMSMNT" and link["unit_code"] == "MNTD1"
                for link in overview["attached_links"]
            )
        )
        # Only attached institutes appear in `institutes` / `scientists`; unattached
        # institutes appear only in `all_institutes` without staff PII.
        self.assertEqual([i["institute_code"] for i in overview["institutes"]], ["AIIMSMNT"])
        self.assertNotIn("dr.unattached.sci@example.org", str(overview))
        self.assertNotIn("9998887776", str(overview))
        all_codes = {i["institute_code"]: i for i in overview["all_institutes"]}
        self.assertIn("UNATTACHMNT", all_codes)
        self.assertNotIn("scientists", all_codes["UNATTACHMNT"])

        grant_res = self.client.post(
            self._url("/mentoring/grants"),
            json={
                "user_id": str(sci_user.user_id),
                "role": "reviewer",
                "org_unit_ids": [str(dist.org_unit_id)],
            },
            headers=headers,
        )
        self.assertEqual(grant_res.status_code, 200, grant_res.get_json())
        sci_rows = [s for s in grant_res.get_json()["scientists"] if s["user_id"] == str(sci_user.user_id)]
        self.assertEqual(len(sci_rows), 1)
        self.assertEqual(sci_rows[0]["grants_by_role"].get("reviewer"), [str(dist.org_unit_id)])

        detach_res = self.client.post(
            self._url("/mentoring/districts"),
            json={"institute_code": "AIIMSMNT", "unit_code": "MNTD1", "action": "detach"},
            headers=headers,
        )
        self.assertEqual(detach_res.status_code, 200, detach_res.get_json())
        self.assertFalse(
            any(
                link["institute_code"] == "AIIMSMNT" and link["unit_code"] == "MNTD1"
                for link in detach_res.get_json()["attached_links"]
            )
        )

    def test_bulk_grid_enforces_home_cadre_capability_active_user_and_mentor_guards(self):
        import uuid

        from app.models import MasOrgUnitWorker

        self._login(str(self.base_admin_id))
        admin_user = db.session.get(VaUsers, uuid.UUID(str(self.base_admin_id)))
        db.session.get(VaProjectMaster, self.PROJECT).project_structure_mode = "organization"
        org.seed_default_organization(self.PROJECT)
        levels = {lv.level_code: lv for lv in org.list_levels(self.PROJECT)}
        cadres = {c.cadre_code: c for c in org.list_cadres(self.PROJECT)}
        dist = org.create_unit(
            self.PROJECT,
            org_level_id=levels["district"].org_level_id,
            unit_code="GRDD1",
            unit_name="Grid District 1",
        )
        chc = org.create_unit(
            self.PROJECT,
            org_level_id=levels["chc"].org_level_id,
            parent_org_unit_id=dist.org_unit_id,
            unit_code="GRDC1",
            unit_name="Grid CHC 1",
        )
        # Worker placed with cadre ANM (which cannot code VA forms at CHC).
        anm_worker = self._get_or_make_user("grid.anm.worker@test.local", "GridWorker123")
        db.session.add(MasOrgUnitWorker(
            project_id=self.PROJECT,
            org_unit_id=chc.org_unit_id,
            cadre_id=cadres["ANM"].cadre_id,
            user_id=anm_worker.user_id,
            worker_code="ANM001",
            worker_name="ANM Worker",
            is_active=True,
        ))
        # Deactivated worker.
        deactive_worker = self._get_or_make_user("grid.deactive.worker@test.local", "GridWorker123")
        deactive_worker.user_status = VaStatuses.deactive
        db.session.commit()

        # 1. Home cadre lacks can_code_va_form at CHC -> refused instead of substituting SMO/MO.
        with self.assertRaises(org.OrganizationError) as ctx:
            org.bulk_apply_project_users_grid(
                self.PROJECT,
                {
                    "user_ids": [str(anm_worker.user_id)],
                    "operation": "coder",
                    "scope_mode": "own",
                },
                actor=admin_user,
            )
        self.assertIn("ANM", str(ctx.exception))

        # 2. Deactivated user -> refused.
        with self.assertRaises(org.OrganizationError) as ctx:
            org.bulk_apply_project_users_grid(
                self.PROJECT,
                {
                    "user_ids": [str(deactive_worker.user_id)],
                    "operation": "collaborator",
                    "scope_mode": "project",
                },
                actor=admin_user,
            )
        self.assertIn("inactive", str(ctx.exception))

        # 3. Mentoring Institute member -> refused for project-scope grants in bulk grid.
        from app.services import mentor_institute_service as mentors

        mentors.create_institute("GRDMC", "Grid Medical College")
        mentor_user = self._get_or_make_user("grid.mentor.sci@test.local", "GridWorker123")
        mentors.add_member("GRDMC", mentor_user.email)
        db.session.commit()
        with self.assertRaises(org.OrganizationError) as ctx:
            org.bulk_apply_project_users_grid(
                self.PROJECT,
                {
                    "user_ids": [str(mentor_user.user_id)],
                    "operation": "reviewer",
                    "scope_mode": "project",
                },
                actor=admin_user,
            )
        self.assertIn("mentoring institute", str(ctx.exception))
