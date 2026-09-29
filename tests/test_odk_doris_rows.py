"""The ODK DORIS rows are generated from the web form's doris_support_whova_2022
extension (digitva-hln) and the committed copies are current.

tooling/who-va-2022/build-odk-doris-rows.mjs writes the rows out of the
extension's own blocks into vendor/who-va-2022/src/generated/
odk-doris-support-rows.json (its Node test pins that file byte for byte);
tooling/who-va-2022/build_odk_doris_rows.py renders it against ND01 into the
committed .md and .xlsx. Here: rendering the committed JSON reproduces both
files, and the rows still agree with the generated web artifacts.
"""

import importlib.util
import json
import unittest
from pathlib import Path

import openpyxl

REPO = Path(__file__).resolve().parents[1]
GENERATED = REPO / "vendor/who-va-2022/src/generated"


def _load_script():
    spec = importlib.util.spec_from_file_location(
        "build_odk_doris_rows", REPO / "tooling/who-va-2022/build_odk_doris_rows.py"
    )
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


def _trimmed(rows):
    """Rows as openpyxl reads them back: "" is an empty cell, none trail."""
    out = []
    for row in rows:
        row = [None if cell == "" else cell for cell in row]
        while row and row[-1] is None:
            row.pop()
        out.append(row)
    return out


class OdkDorisRowsTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.script = _load_script()
        cls.blocks = cls.script.load_blocks()
        nd01 = openpyxl.load_workbook(cls.script.ND01, read_only=True)
        cls.markdown, cls.sheets = cls.script.render(cls.blocks, nd01)
        server = json.loads((GENERATED / "who-va-2022.server-instrument.json").read_text())
        cls.web = {question["name"]: question for question in server["questions"]}
        reference = json.loads((GENERATED / "digitva-layers.reference.json").read_text())
        cls.labels = {
            entry["itemKey"]: entry["text"]
            for entry in reference["entries"]
            if "doris_support_whova_2022" in entry["extensions"] and entry["field"] == "label"
        }
        cls.added = [row for block in cls.blocks for row in block["survey"]]

    def test_committed_markdown_is_current(self):
        self.assertEqual(self.script.OUT_MD.read_text(encoding="utf-8"), self.markdown)

    def test_committed_workbook_is_current(self):
        committed = openpyxl.load_workbook(self.script.OUT_XLSX, read_only=True)
        self.assertEqual(committed.sheetnames, ["survey", "choices", "notes"])
        for title, rows in self.sheets.items():
            with self.subTest(title):
                self.assertEqual(_trimmed(committed[title].iter_rows(values_only=True)), _trimmed(rows))

    def test_every_block_of_annex_a_is_there(self):
        self.assertEqual([block["id"] for block in self.blocks], [f"A{n}" for n in range(1, 11)])
        agreed = [block["id"] for block in self.blocks if block["status"] == self.script.AGREED]
        self.assertEqual(agreed, ["A1", "A2", "A3"])
        others = {block["status"] for block in self.blocks if block["id"] not in agreed}
        self.assertEqual(others, {self.script.PROPOSED})

    def test_added_rows_match_the_web_questions(self):
        self.assertEqual(len(self.added), 19)
        for row in self.added:
            with self.subTest(row["name"]):
                web = self.web.get(row["name"])
                self.assertIsNotNone(web)
                self.assertEqual(row["relevant"] or None, web["relevant"])
                self.assertEqual(row["constraint"] or None, web["constraint"])
                self.assertEqual(row["label"], self.labels[row["name"]])
                kind = row["type"].split()[0]
                expected = {"select_one": "singleChoice", "acknowledge": "confirm"}.get(kind, kind)
                self.assertEqual(web["control"], expected)
        required = [row["name"] for row in self.added if row["required"] == "yes"]
        self.assertEqual(required, ["Id10366_confirm"])

    def test_changed_who_rows(self):
        changes = {change["name"]: change["cells"] for block in self.blocks for change in block["change"]}
        self.assertEqual(set(changes), {"Id10366", "Id10308", "Id10340"})
        self.assertEqual(changes["Id10308"], {"required": "yes"})
        self.assertEqual(changes["Id10366"]["constraint"], ". >= 100 and . <= 9999")
        self.assertIn("selected(${Id10308}, 'yes')", changes["Id10340"]["relevant"])
        # The replaced rows keep ND01's other cells: WHO's own label.
        survey = self.sheets["survey"]
        header = survey[0]
        by_name = {row[header.index("name")]: row for row in survey[1:]}
        self.assertTrue(by_name["Id10340"][header.index(self.script.LABEL)].startswith("(Id10340)"))

    def test_new_choices_match_the_web_lists_and_skip_nd01_lists(self):
        header = self.sheets["choices"][0]
        rows = [dict(zip(header, row)) for row in self.sheets["choices"][1:]]
        self.assertEqual(len(rows), 24)
        self.assertNotIn("YES_NO_DK_REF", {row["list_name"] for row in rows})
        for row in rows:
            with self.subTest(f"{row['list_name']}/{row['name']}"):
                self.assertEqual(self.labels[f"{row['list_name']}/{row['name']}"], row[self.script.LABEL])

    def test_a7_goes_after_the_group_end_and_the_order_is_bottom_up(self):
        self.assertIn("after the `end group` of `health_service_utilization`", self.markdown)
        self.assertIn("bottom-up (A8, A7, A4, A2, A5, A10, A9, A6, A3, A1)", self.markdown)


if __name__ == "__main__":
    unittest.main()
