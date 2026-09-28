"""CSV and workbook boundary checks shared by guided imports."""

import io
import unittest
from zipfile import ZIP_DEFLATED, ZipFile

from openpyxl import Workbook

from app.services.tabular_import_service import TabularImportError, parse_table

HEADERS = ("unit_code", "unit_name", "latitude")


def _parse(contents, filename="units.csv"):
    return parse_table(io.BytesIO(contents), filename, HEADERS,
                       max_bytes=1024 * 1024, max_rows=1000)


class TabularImportTests(unittest.TestCase):
    def test_excel_csv_variants(self):
        rows = _parse("sep=;\r\n Unit_Code ; Unit_Name ; Latitude ;\r\nP01;Café;-12.5;\r\n".encode("cp1252"))
        self.assertEqual(rows[0]["unit_name"], "Café")
        self.assertEqual(rows[0]["latitude"], "-12.5")

    def test_first_workbook_sheet(self):
        workbook = Workbook()
        workbook.active.append([" UNIT_CODE ", "UNIT_NAME", "latitude"])
        workbook.active.append(["P01", "PHC", -12.5])
        workbook.create_sheet("ignored").append(["wrong"])
        buffer = io.BytesIO()
        workbook.save(buffer)
        rows = _parse(buffer.getvalue(), "units.xlsx")
        self.assertEqual(rows[0]["unit_code"], "P01")
        self.assertEqual(rows[0]["latitude"], -12.5)

    def test_rejects_populated_trailing_column(self):
        with self.assertRaisesRegex(TabularImportError, "unexpected populated column"):
            _parse(b"unit_code,unit_name,latitude,\nP01,PHC,,surprise\n")

    def test_rejects_controls_formula_and_malformed_quoting(self):
        for contents in (
            b"unit_code,unit_name,latitude\nP01,A\x00B,\n",
            b"unit_code,unit_name,latitude\nP01,=HYPERLINK(1),\n",
            b"unit_code,unit_name,latitude\nP01,\"unclosed,\n",
        ):
            with self.subTest(contents=contents), self.assertRaises(TabularImportError):
                _parse(contents)

    def test_rejects_formula_in_workbook(self):
        workbook = Workbook()
        workbook.active.append(HEADERS)
        workbook.active.append(["P01", "=HYPERLINK(\"https://example.org\")", None])
        buffer = io.BytesIO()
        workbook.save(buffer)
        with self.assertRaisesRegex(TabularImportError, "formula-like"):
            _parse(buffer.getvalue(), "units.xlsx")

    def test_rejects_zip_valid_malformed_sheet_xml(self):
        workbook = Workbook()
        workbook.active.append(HEADERS)
        buffer = io.BytesIO()
        workbook.save(buffer)
        broken = io.BytesIO()
        with ZipFile(io.BytesIO(buffer.getvalue())) as source, ZipFile(broken, "w", ZIP_DEFLATED) as target:
            for entry in source.infolist():
                body = b"<worksheet><sheetData><row>" if entry.filename == "xl/worksheets/sheet1.xml" else source.read(entry)
                target.writestr(entry, body)
        with self.assertRaisesRegex(TabularImportError, "could not be read"):
            _parse(broken.getvalue(), "units.xlsx")

    def test_rejects_large_zip_member_and_worksheet_dimensions(self):
        workbook = Workbook()
        workbook.active.append(HEADERS)
        buffer = io.BytesIO()
        workbook.save(buffer)
        for replacement in (b" " * (51 * 1024 * 1024),
                            b'<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">'
                            b'<dimension ref="A1:XFD1048576"/></worksheet>'):
            with self.subTest(size=len(replacement)):
                crafted = io.BytesIO()
                with ZipFile(io.BytesIO(buffer.getvalue())) as source, ZipFile(crafted, "w", ZIP_DEFLATED) as target:
                    for entry in source.infolist():
                        target.writestr(entry, replacement if entry.filename == "xl/worksheets/sheet1.xml" else source.read(entry))
                with self.assertRaisesRegex(TabularImportError, "too large|dimensions"):
                    _parse(crafted.getvalue(), "units.xlsx")
