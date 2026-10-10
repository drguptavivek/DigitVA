"""Generate the Help field matrix and XLSX from source forms and composition.

Run with /Users/vivekgupta/.codex/.venv/bin/python after regenerating the
composed instrument. Reads source XLSForms without changing them.
"""

import hashlib
import html
import json
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[2]
web = json.loads((ROOT / "app/data/who-va-2022.composed.json").read_text())
sources = {
    "WHO2026": {"path": "vendor/who-va-2022/2022whova_xls_form_for_odk_multilingual.xlsx"},
    "ICMR ND01": {"path": "docs/kb/WHO_VA_2022_Docs/ND01_ICMRVA_WHOVA2022.xlsx"},
    **{
        "UNSW " + n: {"path": f"docs/kb/WHO_VA_2022_Docs/{n}_DS_WHOVA2022.xlsx"}
        for n in ["KA01", "KL01", "TR01"]
    },
}


def workbook_source(path):
    w = load_workbook(ROOT / path, read_only=True)
    out = {"path": path}
    for name in ["survey", "choices", "settings"]:
        sheet = w[name]
        it = sheet.iter_rows(values_only=True)
        heads = next(it)
        items = []
        for i, row in enumerate(it, 2):
            if any(v is not None for v in row):
                items.append(
                    {**{str(k): v for k, v in zip(heads, row) if k is not None}, "_row": i}
                )
        out[name] = items
    w.close()
    return out


sources = {label: workbook_source(source["path"]) for label, source in sources.items()}
cols = ["WHO2026", "DigitVA Web", "ICMR ND01", "UNSW KA01", "UNSW KL01", "UNSW TR01", "UNSW NC01"]
fields = {}
raws = {}
order = []


def add(label, name, row):
    if not name:
        return
    key = name
    fields.setdefault(key, {})[label] = row
    if key not in order:
        order.append(key)


for label, source in sources.items():
    for r in source["survey"]:
        if r.get("type") == "end group":
            continue
        name = r.get("name")
        add(label, name, r)
for q in web["questions"]:
    r = {
        "name": q["name"],
        "type": q.get("sourceType"),
        "required": "yes" if q.get("required") else None,
        "relevant": (q.get("relevant") or {}).get("source"),
        "constraint": (q.get("constraint") or {}).get("source"),
        "calculation": (q.get("calculation") or {}).get("source"),
        "appearance": q.get("appearance"),
        "read_only": q.get("readOnly"),
        "label::English (en)": q.get("label", {}).get("en"),
        "hint::English (en)": q.get("hint", {}).get("en"),
        "guidance_hint::English (en)": q.get("guidance", {}).get("en"),
        "default": q.get("default"),
        "_row": q.get("sourceRow"),
        "_sectionPath": " / ".join(q.get("sectionPath", [])),
        "constraint_message": q.get("constraintMessage", {}).get("en"),
        "_extensions": ", ".join(q.get("extensions", [])),
    }
    add("DigitVA Web", q["name"], r)
for sec in web["sections"]:
    add(
        "DigitVA Web",
        sec["name"],
        {
            "name": sec["name"],
            "type": "begin group",
            "relevant": (sec.get("relevant") or {}).get("source"),
            "label::English (en)": sec.get("label", {}).get("en"),
            "_row": sec.get("sourceRow"),
        },
    )
context = {
    "Site": "Project/site context",
    "unique_id": "Server-assigned case identifier",
    "survey_state": "Organization geography context",
    "survey_district": "Organization geography context",
    "survey_block": "Organization geography context",
    "site_individual_id": "Death-register context; source semantics differ",
    "org_<level_code>_code": "Organization ancestor unit codes",
}
for name in context:
    if name not in fields:
        fields[name] = {}
        order.append(name)


def purpose(name):
    if name in context:
        return "Identifiers and organization context"
    if name.startswith("doris_") or name in [
        "dob_precision",
        "dob_month_year",
        "dob_year",
        "Id10366_confirm",
    ]:
        return "DORIS support"
    if name.startswith(
        ("sa", "socialautopsy", "socioeconomic", "reachinghealthcare", "eventchronology")
    ):
        return "Social autopsy"
    if name.startswith(("md_", "ds_")) or name == "custom_medical_certificate_upload":
        return "Documents"
    if name in ["Id10476", "Id10476_audio", "imagenarr", "narr_language"]:
        return "Narrative"
    if name in ["abha_number", "abha_address"]:
        return "ABHA"
    if name.startswith("visit_") or name in ["interview_outcome", "consent_mode"]:
        return "DigitVA workflow"
    return "WHO / collection structure"


def cell(name, label):
    if label == "UNSW NC01":
        return "Not verified: workbook unavailable"
    r = fields[name].get(label)
    if not r:
        if label == "DigitVA Web" and name in context:
            return "Host supplied: " + context[name]
        return "Absent"
    ty = r.get("type") or "unknown"
    req = "required" if str(r.get("required", "")).lower() in ["yes", "true"] else "optional"
    rel = r.get("relevant") or "no question-level relevance"
    return f"{ty}; {req}; {rel}"


headers = ["Field", "Purpose"] + cols + ["DigitVA scope / notes"]
rows = []
for name in order:
    q = fields[name].get("DigitVA Web", {})
    note = q.get("_extensions") or ("Core/WHO; project options still apply" if q else "")
    if name == "Id10476":
        note = "DigitVA core override: always visible within consent gate; required multiline"
    if name == "Id10476_audio":
        note = "Optional audio; fallback-to-text hint removed in DigitVA"
    if name in context:
        note = context[name] + "; host context is separate from questionnaire fields"
    if name == "Id10304_a":
        note = "WHO2026 relevance differs from older project forms; inspect Rules sheet"
    rows.append([name, purpose(name)] + [cell(name, label) for label in cols] + [note])
wb = Workbook()
matrix = wb.active
matrix.title = "Field matrix"
matrix.append(headers)
for row in rows:
    matrix.append(row)
legend = wb.create_sheet("Read me")
legend.append(["Topic", "Explanation"])
for r in [
    [
        "Purpose",
        "Field-level comparison of WHO2026, actual DigitVA composed definition, local ICMR and mapped UNSW workbooks.",
    ],
    [
        "Blank relevance",
        "No question-level condition. Ancestor group relevance still applies; this is not unconditional access before consent.",
    ],
    [
        "DigitVA Web",
        "All-on composed artifact. Project-disabled extension fields may be omitted. Host context is marked separately.",
    ],
    [
        "Required",
        "Runtime client requiredness. Server relevance retention does not independently enforce all required answers.",
    ],
    [
        "NC01",
        "Live local mapping identifies UNSW NC01_DS_WHOVA2022, but no NC01 workbook is present locally; cells are unverified.",
    ],
    [
        "UNSW coverage",
        "Local mapping confirms KA01_DS_WHOVA2022, KL01_DS_WHOVA2022, TR01_DS_WHOVA2022. JIPMER is excluded from UNSW.",
    ],
    [
        "Dates",
        "WHO2026 is instrument V2.0 version2026081401. Project files have older settings versions; workbook comparison does not certify current Central deployment.",
    ],
    [
        "Sources",
        "Read-only original XLSX files. Source formulas stored as literal text; never recalculated or rewritten.",
    ],
    [
        "Details",
        "Rules contains exact per-source expressions/messages, including groups. Choices contains source lists and coded values.",
    ],
    [
        "Editing",
        "This is a generated audit table, not a questionnaire or a form to fill. Change source/composition and regenerate rather than editing comparisons.",
    ],
]:
    legend.append(r)
prov = wb.create_sheet("Sources")
prov.append(["Column", "File", "Version", "SHA256"])
for label, s in sources.items():
    p = ROOT / s["path"]
    settings = s.get("settings", [])
    version = settings[0].get("version") if settings else ""
    prov.append([label, s["path"], version, hashlib.sha256(p.read_bytes()).hexdigest()])
prov.append(
    [
        "DigitVA Web",
        "app/data/who-va-2022.composed.json",
        web["version"],
        hashlib.sha256((ROOT / "app/data/who-va-2022.composed.json").read_bytes()).hexdigest(),
    ]
)
prov.append(["UNSW NC01", "No local source workbook", None, None])
rules = wb.create_sheet("Rules")
attrs = [
    "type",
    "required",
    "relevant",
    "constraint",
    "calculation",
    "appearance",
    "read_only",
    "default",
    "choice_filter",
    "trigger",
    "agegroup",
    "label::English (en)",
    "hint::English (en)",
    "guidance_hint::English (en)",
]
rules.append(
    ["Source", "Field", "Excel/source row", "Section path"] + attrs + ["Constraint message"]
)
for name in order:
    for label in cols[:-1]:
        r = fields[name].get(label)
        if r:
            rules.append(
                [label, name, r.get("_row"), r.get("_sectionPath", "")]
                + [r.get(k) for k in attrs]
                + [r.get("constraint_message::English (en)", r.get("constraint_message"))]
            )
choices = wb.create_sheet("Choices")
choices.append(["Source", "List", "Value", "English label", "Row", "Parent state"])
for label, s in sources.items():
    for r in s["choices"]:
        if r.get("list_name") and r.get("name") is not None:
            choices.append(
                [
                    label,
                    r["list_name"],
                    str(r["name"]),
                    r.get("label::English (en)"),
                    r.get("_row"),
                    r.get("state"),
                ]
            )
seen = set()
for q in web["questions"]:
    for c in q.get("choices", []):
        key = (q.get("listName", q["name"]), str(c["value"]))
        if key not in seen:
            choices.append(
                ["DigitVA Web", *key, c.get("label", {}).get("en"), q.get("sourceRow"), None]
            )
            seen.add(key)
for sh in wb:
    sh.freeze_panes = "C2" if sh.title == "Field matrix" else "A2"
    sh.auto_filter.ref = sh.dimensions
    for c in sh[1]:
        c.font = Font(name="Arial", bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1E4C87")
    for row in sh.iter_rows(min_row=2):
        for c in row:
            if isinstance(c.value, str) and c.value.startswith("="):
                c.data_type = "s"
            c.font = Font(name="Arial", size=10)
            c.alignment = Alignment(vertical="top", wrap_text=True)
    for i in range(1, sh.max_column + 1):
        sh.column_dimensions[get_column_letter(i)].width = 28 if i < 3 else 42
    sh.column_dimensions["A"].width = 29
out = ROOT / "app/static/help/digitva-form-field-matrix.xlsx"
out.parent.mkdir(parents=True, exist_ok=True)
wb.save(out)
# Full field presence/rule table for Help, escaped before being included as static HTML.
parts = [
    '<p class="small text-muted">WHO instrument V2.0, revision 2026081401. Optional extension fields depend on project settings. Blank question relevance still inherits its group gates. NC01 source workbook is unavailable.</p>',
    f'<details><summary>Show the complete field matrix ({len(rows)} fields and groups)</summary><div class="table-responsive"><table class="table table-sm table-bordered align-middle small"><thead><tr>',
]
parts += ['<th scope="col">' + html.escape(h) + "</th>" for h in headers]
parts += ["</tr></thead><tbody>"]
for row in rows:
    parts.append("<tr>")
    parts += ["<td>" + html.escape(str(v)) + "</td>" for v in row]
    parts.append("</tr>")
parts += ["</tbody></table></div></details>"]
partial = ROOT / "app/templates/help/pages/_form-field-matrix.html"
partial.write_text("\n".join(line.rstrip() for line in "\n".join(parts).splitlines()) + "\n")
# Read back the deliverable; all sheets contain data and source expressions are literal strings.
check = load_workbook(out, read_only=True, data_only=False)
assert check["Field matrix"].max_row == len(rows) + 1
assert any(
    r[0] == "Id10476" and "no question-level relevance" in r[3]
    for r in check["Field matrix"].iter_rows(min_row=2, values_only=True)
)
assert all(c.data_type != "f" for sh in check for row in sh for c in row)
check.close()
print("Matrix fields", len(rows), "sheets", wb.sheetnames, "output", out)
